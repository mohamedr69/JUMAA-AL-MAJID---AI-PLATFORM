"""Project > Logs: the consolidated registers, read from the same records
the operational pages edit -- never a status of the register's own.

GET /projects/{id}/logs/systems                 the project's systems and the registers each takes part in
GET /projects/{id}/logs/drawings                the Drawings register: one row per logical shop drawing
GET /projects/{id}/logs/drawings/export.xlsx    the same, as a workbook (a revision-details sheet beside it)
GET /projects/{id}/logs/drawings/{drawing_id}   one drawing's record: revisions, evidence, activity (the Drawings page's own)

The Drawings register is `shop_drawings.log` per system, one after another
(`shop_drawings.register`): the Drawings page and this register cannot
disagree, because neither works anything out for itself. Filtering, search
and paging are the register's; folder opening and candidate decisions are
the Drawings page's endpoints, called from here as they are.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.deps import get_current_user
from app.models import Project, User
from app.routers.drawings import _catch_up, _drawing, shop_drawings_folder
from app.routers.projects import _get_project_or_404
from app.services import shop_drawings, system_rules

router = APIRouter(tags=["logs"])

ALL = "ALL"
PAGE_SIZES = (25, 50, 100)


def _systems(db: Session, project: Project, wanted: str | None, register: str) -> tuple[str, list[str], list[dict]]:
    """(the selection, the codes it covers, every system with its registers).
    ALL covers each system that takes part in `register`; a system that
    does not is refused, never quietly answered with another's rows."""
    systems = shop_drawings.register_systems(db, project)
    taking_part = [s["code"] for s in systems if s[register]]
    if not wanted or wanted.strip().upper() == ALL:
        return ALL, taking_part, systems
    code = system_rules.canonical(wanted) or wanted.strip().upper()
    if code not in taking_part:
        raise HTTPException(404, f"{wanted} has no {register.replace('_', ' ')} register on this project"
                                 + (f": the systems with one are {', '.join(taking_part)}" if taking_part else ""))
    return code, [code], systems


@router.get("/projects/{project_id}/logs/systems")
def logs_systems(project_id: int, _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """The project's systems, by their backend code and the name the project
    gives them, and which registers each takes part in: the fire-rated
    cables have material submittals and samples, never shop drawings."""
    project = _get_project_or_404(db, project_id)
    return {"systems": shop_drawings.register_systems(db, project),
            "drawings_in_scope": system_rules.drawings_in_scope(project)}


def _matches(row: dict, *, floor: str | None, status: str | None, revision: str | None, issues_only: bool,
             search: str | None) -> bool:
    if floor and floor not in row["floor_keys"]:
        return False
    if status and row["latest_status"] != status:
        return False
    if revision and (row["cells"].get(revision) or {}).get("status", "not_submitted") == "not_submitted" \
            and not (row["cells"].get(revision) or {}).get("candidate"):
        return False
    if issues_only and not row["hints"]:
        return False
    if search:
        words = [w for w in re.split(r"\s+", search.strip().upper()) if w]
        if not all(w in row["search"] for w in words):
            return False
    return True


def _summary(rows: list[dict]) -> dict:
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["latest_status"]] = counts.get(row["latest_status"], 0) + 1
    return {"total": len(rows), **{k: counts.get(k, 0) for k in shop_drawings.STATUS_LABELS},
            "review_items": sum(1 for row in rows if row["hints"]),
            "candidates": sum(len(row["candidates"]) for row in rows)}


@router.get("/projects/{project_id}/logs/drawings")
def drawings_register(project_id: int, system: str | None = None, floor: str | None = None, status: str | None = None,
                      revision: str | None = None, issues_only: bool = False, search: str | None = None,
                      page: int = 1, page_size: int = 25,
                      _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """The Drawings register: one row per logical shop drawing (and per
    floor still to draw), across every drawn system (ALL) or one. The
    same rows as the Drawings page's log for that system."""
    project = _get_project_or_404(db, project_id)
    _catch_up(db, project)
    selected, codes, systems = _systems(db, project, system, "drawings")
    built = shop_drawings.register(db, project, codes)
    rows = [row for row in built["rows"]
            if _matches(row, floor=floor, status=status, revision=revision, issues_only=issues_only, search=search)]
    size = page_size if page_size in PAGE_SIZES else PAGE_SIZES[0]
    pages = max(1, (len(rows) + size - 1) // size)
    current = min(max(page, 1), pages)
    start = (current - 1) * size
    floors = list({k: None for row in built["rows"] for k in row["floor_keys"]})
    names = {}
    for row in built["rows"]:
        if len(row["floor_keys"]) == 1:
            names.setdefault(row["floor_keys"][0], row["floor"])
    return {
        "project": {"id": project.id, "ep_number": project.ep_number, "name": project.project_name},
        "systems": systems, "selected_system": selected,
        "summary": _summary(built["rows"]), "filtered": len(rows),
        "available_revisions": built["revisions"],
        "floors": [{"key": k, "name": names.get(k, k)} for k in floors],
        "rows": [{k: v for k, v in row.items() if k != "search"} for row in rows[start:start + size]],
        "pagination": {"page": current, "page_size": size, "pages": pages, "total": len(rows)},
        "synced_at": project.documents_synced_at.isoformat() if project.documents_synced_at else None,
        "reconciled_at": project.drawings_reconciled_at.isoformat() if project.drawings_reconciled_at else None,
        "folder": shop_drawings_folder(selected) if selected != ALL and project.source_folder_path else None,
        "warnings": (["The project folder has not been synced yet: sync the documents to read the shop drawings."]
                     if project.source_folder_path and project.documents_synced_at is None else []),
    }


@router.get("/projects/{project_id}/logs/drawings/export.xlsx")
def export_drawings_register(project_id: int, system: str | None = None,
                             _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """The register as a workbook: the rows, and a Revision Details sheet
    with each official revision's submission, reply and file."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    project = _get_project_or_404(db, project_id)
    _catch_up(db, project)
    selected, codes, _systems_all = _systems(db, project, system, "drawings")
    built = shop_drawings.register(db, project, codes)
    label = "All" if selected == ALL else selected
    wb = Workbook()
    ws = wb.active
    ws.title = f"Drawings Register {label}"[:31]
    ws.append([f"EP-{project.ep_number} {project.project_name or ''} - Drawings Register - {label}"])
    ws.append([f"One row per shop drawing; statuses as the consultant answered each revision; last sync "
               f"{project.documents_synced_at.isoformat(timespec='minutes').replace('T', ' ') if project.documents_synced_at else 'never'}; "
               f"as of {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    revisions = built["revisions"]
    header = ["#", "System", "Floor", "Drawing Title", "Drawing Reference", *revisions, "Latest Revision", "Latest Status",
              "Issues / Hints", "Remarks"]
    ws.append(header)
    fills = {"approved": "C6EFCE", "approved_as_noted": "DDEBF7", "under_review": "FFEB9C", "not_approved": "FFC7CE",
             "not_submitted": "EDEDED", "reply_not_found": "F4B183"}
    first_revision_column = 6
    for i, row in enumerate(built["rows"], 1):
        floor = row["floor"] + (f" ({row['floor_secondary']})" if row.get("floor_secondary") else "")
        ws.append([i, row["system"], floor, row.get("title") or "-", row["reference"] or "Not submitted yet",
                   *(row["cells"][r]["label"] for r in revisions),
                   row["latest_revision"] or "-", shop_drawings.STATUS_LABELS.get(row["latest_status"], row["latest_status"]),
                   "; ".join(h["label"] for h in row["hints"]) or "-", row["remarks"] or "-"])
        for j, rev in enumerate(revisions):
            if row["cells"][rev].get("candidate"):
                continue
            ws.cell(row=ws.max_row, column=first_revision_column + j).fill = PatternFill(
                "solid", fgColor=fills.get(row["cells"][rev]["status"], "FFFFFF"))
    for c in ws[4]:
        c.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=13)
    for index, width in enumerate([5, 9, 30, 40, 40, *([18] * len(revisions)), 15, 22, 40, 50], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for r in ws.iter_rows(min_row=5):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    details = wb.create_sheet("Revision Details")
    details.append(["System", "Floor", "Drawing Reference", "Revision", "Status", "Submission Reference", "Submitted Date",
                    "Consultant Reply Reference", "Reply Date", "Drawing File", "Remarks"])
    for row in built["rows"]:
        if row["drawing_id"] is None:
            continue
        drawing = _drawing(db, project, row["drawing_id"])
        for r in sorted(drawing.revisions, key=lambda r: r.number):
            if not r.submitted:
                continue
            details.append([row["system"], row["floor"], row["reference"], r.revision,
                            shop_drawings.STATUS_LABELS.get(r.status, r.status), r.submission_reference or "-",
                            r.submitted_at.date().isoformat() if r.submitted_at else "-", r.reply_reference or "-",
                            r.reply_at.date().isoformat() if r.reply_at else "-", r.drawing_path or "-",
                            r.reply_text or r.note or "-"])
    for c in details[1]:
        c.font = Font(bold=True)
    for index, width in enumerate([9, 30, 40, 10, 22, 28, 14, 28, 14, 60, 50], 1):
        details.column_dimensions[get_column_letter(index)].width = width
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"EP-{project.ep_number} Drawings Register {label}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/projects/{project_id}/logs/drawings/{drawing_id}")
def drawing_record(project_id: int, drawing_id: int, _current_user: User = Depends(get_current_user),
                   db: Session = Depends(get_db)):
    """One drawing's record, as the Drawings page's details show it -- the
    same function, so the two cannot differ."""
    project = _get_project_or_404(db, project_id)
    return shop_drawings.detail(db, project, _drawing(db, project, drawing_id))
