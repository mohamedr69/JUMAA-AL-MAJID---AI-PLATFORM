"""The Drawings page (app/services/shop_drawings, drawing_issues, drawing_requirements).

GET  /projects/{id}/drawings/summary                    every system: floors, statuses, review items (the system cards)
GET  /projects/{id}/drawings/log?system=                one system's log: the building's floors x shop drawing revisions
GET  /projects/{id}/drawings/log/export.xlsx?system=    the same, as a workbook
GET  /projects/{id}/drawings/issues?system=             Review & Issues: the open findings (system checks and AI review)
POST /projects/{id}/drawings/issues/{issue}/resolve     an engineer settles one
GET  /projects/{id}/drawings/activity?system=           Activity / History
GET  /projects/{id}/drawings/sd/{drawing}               one shop drawing: overview, revisions, candidates, files, events
PATCH /projects/{id}/drawings/sd/{drawing}              correct its reference, floors or remarks
PUT  /projects/{id}/drawings/sd/{drawing}/revisions/{rev} set a revision's status (authoritative)
POST /projects/{id}/drawings/candidates/{cand}/confirm  a detected revision was submitted
POST /projects/{id}/drawings/candidates/{cand}/ignore   ... or was not
POST /projects/{id}/drawings/reconcile                  bring the records up to the document index now
POST /projects/{id}/drawings/open-folder                open a folder of the project in Explorer, on this PC
GET  /projects/{id}/drawings/required?system=           Actions Required: what the contractor must hand over
GET  /projects/{id}/drawings/required/export.xlsx?system=
POST /projects/{id}/drawings/required/request           the request email (records nothing)
POST /projects/{id}/drawings/required/request/sent      ... was sent: noted against each item

The router checks and answers; the records are written by the sync
(app.services.shop_drawings.reconcile) and read here. No request walks
the project folder or resolves an IFC drawing.
"""
from __future__ import annotations

import io
import os
import subprocess
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.database import get_db
from app.deps import get_current_user, require_role
from app.models import DrawingIssue, Project, ProjectShopDrawing, ShopDrawingCandidate, User
from app.routers.projects import CREATOR_ROLES, _get_project_or_404
from app.services import (activity, building_floors, drawing_issues, drawing_requirements, project_folders,
                          project_state, required_drawings, shop_drawings, system_rules)

router = APIRouter(tags=["drawings"])


def shop_drawings_folder(system: str | None) -> str:
    """Where a system's shop drawings are filed: each system has its own
    folder -- FA, ELS -- and the page opens the one on show."""
    name = project_folders.system_folder(system)
    return f"{project_folders.DRAWINGS}/SD/{name}" if name else f"{project_folders.DRAWINGS}/SD"


def _first_system(project) -> str:
    codes = list(system_rules.project_codes(project))
    return codes[0] if codes else "FAS"


# Kept for callers of the earlier names.
_systems_with_drawings = shop_drawings.systems_for
in_force_drawings = building_floors.in_force_as_log_input


def _system(db: Session, project: Project, system: str | None) -> tuple[str, list[str]]:
    """The system asked for, checked against the project's: no silent
    fallback from one system to another. Left out: the project's first."""
    # A project with no systems recorded yet (no DRF read) has the fire alarm's log, as before.
    systems = shop_drawings.project_systems(db, project) or [_first_system(project)]
    if system is None or not system.strip():
        return systems[0], systems
    code = system_rules.canonical(system)
    if code not in systems:
        raise HTTPException(404, f"{system} is not a system of this project" + (f": its systems are {', '.join(systems)}"
                                                                                  if systems else ""))
    return code, systems


def _catch_up(db: Session, project: Project) -> None:
    """The records were last brought up to the index before the last sync
    (or never): bring them up once, here. The sync does this itself when
    it runs; this is for projects synced before the records existed."""
    if shop_drawings.needs_reconcile(project):
        shop_drawings.reconcile(db, project, ai=False)


def _log(db: Session, project: Project, system: str | None = None) -> dict:
    _catch_up(db, project)
    wanted, systems = _system(db, project, system)
    out = shop_drawings.log(db, project, wanted)
    warnings: list[str] = []
    if project.source_folder_path and project.documents_synced_at is None:
        warnings.append("The project folder has not been synced yet: sync the documents to read the shop drawings.")
    from app.ifc.services import revisions

    in_force = revisions.in_force(db, project.id)
    out.update({
        "project": {"id": project.id, "ep_number": project.ep_number, "name": project.project_name},
        "system": wanted, "system_name": system_rules.CODE_NAMES.get(wanted, wanted), "systems": systems,
        "ifc": [{"id": d.id, "filename": d.filename, "revision": d.revision or "R0"} for d in in_force],
        "synced_at": project.documents_synced_at.isoformat() if project.documents_synced_at else None,
        "reconciled_at": project.drawings_reconciled_at.isoformat() if project.drawings_reconciled_at else None,
        "folder": shop_drawings_folder(wanted) if project.source_folder_path else None,
        "warnings": warnings,
    })
    return out


@router.get("/projects/{project_id}/drawings/summary")
def drawings_summary(project_id: int, _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Every system of the project with its floors, statuses and review
    items: the system cards, and the design manager's view."""
    project = _get_project_or_404(db, project_id)
    _catch_up(db, project)
    return {"project": {"id": project.id, "ep_number": project.ep_number, "name": project.project_name},
            "systems": shop_drawings.summary(db, project),
            "synced_at": project.documents_synced_at.isoformat() if project.documents_synced_at else None}


@router.get("/projects/{project_id}/drawings/log")
def drawings_log(project_id: int, system: str | None = None,
                 _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """One system's Drawings Log: each floor of the building with its shop
    drawing's status at every official revision, from the records."""
    return _log(db, _get_project_or_404(db, project_id), system)


@router.post("/projects/{project_id}/drawings/reconcile")
def reconcile(project_id: int, current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    """Bring the records up to the document index now (the sync does this
    itself when it runs)."""
    project = _get_project_or_404(db, project_id)
    return shop_drawings.reconcile(db, project, user=current_user, ai=False)


@router.get("/projects/{project_id}/drawings/log/export.xlsx")
def export_drawings_log(project_id: int, system: str | None = None,
                        _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    project = _get_project_or_404(db, project_id)
    log = _log(db, project, system)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Drawings Log {log['system']}"[:31]
    ws.append([f"EP-{project.ep_number} {project.project_name or ''} - Drawings Log - {log['system_name']} ({log['system']})"])
    ws.append([f"Project EP-{project.ep_number}; system {log['system']}; last sync "
               f"{log['synced_at'][:16].replace('T', ' ') if log['synced_at'] else 'never'}; "
               f"floors from {', '.join(f'{d['filename']} {d['revision']}' for d in log['ifc']) or 'no IFC drawing'}; "
               f"as of {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    header = ["#", "Floor", "Drawing Reference", "No. of floors", *log["revisions"], "Latest Revision", "Latest Status",
              "Issues / Hints", "Remarks"]
    ws.append(header)
    fills = {"approved": "C6EFCE", "approved_as_noted": "DDEBF7", "under_review": "FFEB9C", "not_approved": "FFC7CE",
             "not_submitted": "EDEDED", "reply_not_found": "F4B183"}
    first_revision_column = 5
    for i, row in enumerate(log["rows"], 1):
        hints = "; ".join(h["label"] for h in row.get("hints", []))
        ws.append([i, row["floor"], row["reference"] or "Not submitted yet", row["floors"],
                   *(row["cells"][r]["label"] for r in log["revisions"]),
                   row["latest_revision"] or "-", shop_drawings.STATUS_LABELS.get(row["latest_status"], row["latest_status"]),
                   hints or "-", row["remarks"] or "-"])
        for j, rev in enumerate(log["revisions"]):
            cell = ws.cell(row=ws.max_row, column=first_revision_column + j)
            status = row["cells"][rev]["status"]
            if row["cells"][rev].get("candidate"):
                continue     # a candidate is a hint, not a status: no colour
            cell.fill = PatternFill("solid", fgColor=fills.get(status, "FFFFFF"))
    for c in ws[4]:
        c.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=13)
    widths = [5, 34, 40, 12, *([18] * len(log["revisions"])), 15, 22, 40, 60]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for r in ws.iter_rows(min_row=5):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"EP-{project.ep_number} Drawings Log {log['system']}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# --- Review & Issues, Activity ----------------------------------------------------------------


@router.get("/projects/{project_id}/drawings/issues")
def issues(project_id: int, system: str | None = None, _current_user: User = Depends(get_current_user),
           db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    _catch_up(db, project)
    wanted, systems = _system(db, project, system)
    rows = drawing_issues.open_issues(db, project.id, wanted)
    return {"system": wanted, "systems": systems, "total": len(rows),
            "system_checks": [drawing_issues.out(i) for i in rows if i.source == drawing_issues.SYSTEM],
            "ai_review": [drawing_issues.out(i) for i in rows if i.source == drawing_issues.AI]}


# --- the building's floors: one row per physical floor ------------------------------------------


@router.get("/projects/{project_id}/drawings/floors")
def building_floors_view(project_id: int, _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """The Building Floor Registry: each floor's canonical identity, the
    project's own name for it, and the aliases the project holds -- read
    off its drawings' titles or confirmed by an engineer."""
    project = _get_project_or_404(db, project_id)
    _catch_up(db, project)
    rows = building_floors.registry(db, project.id, include_inactive=True)
    aliases = building_floors.alias_rows(db, project.id)
    return {
        "floors": [{"key": f.floor_key, "name": f.display_name, "secondary": f.secondary_name, "elevation": f.elevation,
                    "active": f.active, "merged_into": f.merged_into, "source": f.source, "ifc_sheet": f.ifc_sheet}
                   for f in rows],
        "aliases": [{"alias_key": a.alias_key, "canonical_key": a.canonical_key, "alias_label": a.alias_label,
                     "decision": a.decision, "source": a.source, "evidence": a.evidence or {},
                     "confirmed_by_id": a.confirmed_by_id, "updated_at": a.updated_at.isoformat() if a.updated_at else None}
                    for a in aliases],
    }


class FloorDecision(BaseModel):
    alias_key: str = Field(max_length=80)
    canonical_key: str = Field(max_length=80)
    # A merge that brings two drawings together on one floor goes ahead only when confirmed.
    confirm: bool = False


def _floor_decision(db: Session, project: Project, body: FloorDecision) -> tuple[str, str]:
    alias_key, canonical_key = body.alias_key.strip(), body.canonical_key.strip()
    if not alias_key or not canonical_key:
        raise HTTPException(422, "Both floors must be named")
    if alias_key == canonical_key:
        raise HTTPException(422, "A floor cannot be merged with itself")
    known = {f.floor_key for f in building_floors.registry(db, project.id, include_inactive=True)}
    if alias_key not in known:
        raise HTTPException(404, f"{alias_key} is not a floor of this building")
    # The floor merged into must be one the building has, or a level a
    # drawing could name ("L2"): never a key nothing will ever match.
    from app.services import drawing_log

    if canonical_key not in known and drawing_log._elevation(canonical_key) is None:
        raise HTTPException(404, f"{canonical_key} is not a floor of this building")
    return alias_key, canonical_key


@router.post("/projects/{project_id}/drawings/floors/merge")
def merge_floors(project_id: int, body: FloorDecision, current_user: User = Depends(require_role(*CREATOR_ROLES)),
                 db: Session = Depends(get_db)):
    """Two names, one physical floor: "1st Mechanical Floor" is L02 on this
    building. Kept as the project's alias and reused by every IFC revision
    and sync. Where both floors already carry shop drawings, answers 409
    `merge_review` with what a merge brings together, until confirmed;
    nothing of any drawing is overwritten either way."""
    project = _get_project_or_404(db, project_id)
    alias_key, canonical_key = _floor_decision(db, project, body)
    result = shop_drawings.merge_floors(db, project, alias_key, canonical_key, current_user, confirm=body.confirm)
    if result.get("requires_confirmation"):
        raise HTTPException(409, detail={
            "code": "merge_review",
            "message": f"Both {result['alias_label']} and {result['canonical_label']} already carry shop drawings: "
                       f"review what the merge brings together, then confirm.",
            **{k: v for k, v in result.items() if k not in ("merged", "requires_confirmation")}})
    activity.record(db, current_user, "drawings.floors_merged",
                    f"Merged the floor {result['alias_label']} as {result['canonical_label']}", project=project,
                    entity_type="building_floor", detail={"alias_key": alias_key, "canonical_key": canonical_key,
                                                          "drawings": result.get("drawings", 0)})
    return result


@router.post("/projects/{project_id}/drawings/floors/separate")
def keep_floors_separate(project_id: int, body: FloorDecision, current_user: User = Depends(require_role(*CREATOR_ROLES)),
                         db: Session = Depends(get_db)):
    """Two names, two floors: kept, so the question is not asked again."""
    project = _get_project_or_404(db, project_id)
    alias_key, canonical_key = _floor_decision(db, project, body)
    result = shop_drawings.keep_floors_separate(db, project, alias_key, canonical_key, current_user)
    activity.record(db, current_user, "drawings.floors_kept_separate",
                    f"Kept the floors {result['alias_label']} and {result['canonical_label']} separate", project=project,
                    entity_type="building_floor", detail={"alias_key": alias_key, "canonical_key": canonical_key})
    return result


class Resolve(BaseModel):
    resolution: str = Field(default="Reviewed", max_length=500)


@router.post("/projects/{project_id}/drawings/issues/{issue_id}/resolve")
def resolve_issue(project_id: int, issue_id: int, body: Resolve, current_user: User = Depends(require_role(*CREATOR_ROLES)),
                  db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    issue = db.get(DrawingIssue, issue_id)
    if issue is None or issue.project_id != project.id:
        raise HTTPException(404, "Issue not found")
    drawing_issues.resolve(db, issue, current_user, body.resolution.strip() or "Reviewed")
    shop_drawings.event(db, project, "issue.resolved", f"{drawing_issues.label_of(issue.kind)} resolved: {body.resolution}",
                        system=issue.system_code, floor_key=issue.floor_key, user=current_user,
                        detail={"issue_id": issue.id, "kind": issue.kind})
    project_state.record_change(db, project.id, "drawing", "updated", system=issue.system_code)
    db.commit()
    activity.record(db, current_user, "drawings.issue_resolved", f"Resolved: {issue.text[:120]}", project=project,
                    entity_type="drawing_issue", entity_id=issue.id, detail={"system": issue.system_code, "kind": issue.kind})
    return drawing_issues.out(issue)


@router.get("/projects/{project_id}/drawings/activity")
def drawings_activity(project_id: int, system: str | None = None, all_systems: bool = False,
                      _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    if all_systems:
        return {"system": None, "events": shop_drawings.events(db, project)}
    wanted, _systems = _system(db, project, system)
    return {"system": wanted, "events": shop_drawings.events(db, project, wanted)}


# --- one shop drawing --------------------------------------------------------------------------


def _drawing(db: Session, project: Project, drawing_id: int) -> ProjectShopDrawing:
    drawing = db.get(ProjectShopDrawing, drawing_id)
    if drawing is None or drawing.project_id != project.id:
        raise HTTPException(404, "Shop drawing not found")
    return drawing


@router.get("/projects/{project_id}/drawings/sd/{drawing_id}")
def drawing_detail(project_id: int, drawing_id: int, _current_user: User = Depends(get_current_user),
                   db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    return shop_drawings.detail(db, project, _drawing(db, project, drawing_id))


class DrawingPatch(BaseModel):
    drawing_reference: str | None = Field(default=None, max_length=160)
    floor_keys: list[str] | None = None
    remarks: str | None = Field(default=None, max_length=2000)


@router.patch("/projects/{project_id}/drawings/sd/{drawing_id}")
def patch_drawing(project_id: int, drawing_id: int, body: DrawingPatch,
                  current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    if body.drawing_reference is not None and not body.drawing_reference.strip():
        raise HTTPException(422, "A drawing reference cannot be blank")
    if body.floor_keys is not None:
        known = {f.floor_key for f in building_floors.registry(db, project.id, include_inactive=True)}
        unknown = [k for k in body.floor_keys if k not in known]
        if unknown:
            raise HTTPException(422, f"Not a floor of this building: {', '.join(unknown[:5])}")
    clash = (db.query(ProjectShopDrawing)
             .filter(ProjectShopDrawing.project_id == project.id, ProjectShopDrawing.system_code == drawing.system_code,
                     ProjectShopDrawing.drawing_reference == (body.drawing_reference or "").strip(),
                     ProjectShopDrawing.id != drawing.id).first()) if body.drawing_reference else None
    if clash is not None:
        raise HTTPException(409, f"{body.drawing_reference.strip()} is already another {drawing.system_code} drawing's reference")
    shop_drawings.set_drawing(db, project, drawing, current_user, drawing_reference=body.drawing_reference,
                              floor_keys=body.floor_keys, remarks=body.remarks)
    project_state.record_change(db, project.id, "drawing", "updated", entity_id=drawing.id, system=drawing.system_code)
    db.commit()
    activity.record(db, current_user, "drawings.corrected", f"Corrected the shop drawing {drawing.drawing_reference}",
                    project=project, entity_type="shop_drawing", entity_id=drawing.id, detail={"system": drawing.system_code})
    return shop_drawings.detail(db, project, drawing)


class RevisionStatus(BaseModel):
    status: str
    note: str = Field(default="", max_length=1000)
    submitted: bool = True


@router.put("/projects/{project_id}/drawings/sd/{drawing_id}/revisions/{revision}")
def set_revision_status(project_id: int, drawing_id: int, revision: str, body: RevisionStatus,
                        current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    """An engineer's word on a revision's status: it stands over what the
    folder says, and over what the AI says, from now on."""
    project = _get_project_or_404(db, project_id)
    drawing = _drawing(db, project, drawing_id)
    if body.status not in shop_drawings.OFFICIAL_STATUSES:
        raise HTTPException(422, f"status is one of {', '.join(shop_drawings.OFFICIAL_STATUSES)}")
    name = f"R{shop_drawings._rev(revision)}" if shop_drawings._rev(revision) >= 0 else None
    if name is None:
        raise HTTPException(422, f"'{revision}' is not a revision: R0, R1, R2 ...")
    shop_drawings.set_status(db, project, drawing, name, body.status, current_user, note=body.note.strip(),
                             submitted=body.submitted)
    project_state.record_change(db, project.id, "drawing", "updated", entity_id=drawing.id, system=drawing.system_code)
    db.commit()
    activity.record(db, current_user, "drawings.status_set",
                    f"Set {drawing.drawing_reference} {name} to {shop_drawings.STATUS_LABELS[body.status]}",
                    project=project, entity_type="shop_drawing", entity_id=drawing.id,
                    detail={"system": drawing.system_code, "revision": name, "status": body.status})
    return shop_drawings.detail(db, project, drawing)


def _candidate(db: Session, project: Project, candidate_id: int) -> ShopDrawingCandidate:
    candidate = db.get(ShopDrawingCandidate, candidate_id)
    if candidate is None or candidate.project_id != project.id:
        raise HTTPException(404, "Detected revision not found")
    return candidate


class Confirm(BaseModel):
    submission_reference: str | None = Field(default=None, max_length=160)


@router.post("/projects/{project_id}/drawings/candidates/{candidate_id}/confirm")
def confirm_candidate(project_id: int, candidate_id: int, body: Confirm | None = None,
                      current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    candidate = _candidate(db, project, candidate_id)
    if candidate.candidate_status != "available":
        raise HTTPException(409, f"This detected revision was already {candidate.candidate_status}")
    shop_drawings.confirm_candidate(db, project, candidate, current_user,
                                    submission_reference=(body.submission_reference or None) if body else None)
    drawing = candidate.drawing
    project_state.record_change(db, project.id, "drawing", "updated", entity_id=drawing.id, system=drawing.system_code)
    db.commit()
    activity.record(db, current_user, "drawings.revision_confirmed",
                    f"Confirmed {drawing.drawing_reference} {candidate.revision} as submitted", project=project,
                    entity_type="shop_drawing", entity_id=drawing.id, detail={"system": drawing.system_code, "revision": candidate.revision})
    return shop_drawings.detail(db, project, drawing)


class Ignore(BaseModel):
    reason: str = Field(default="", max_length=500)


@router.post("/projects/{project_id}/drawings/candidates/{candidate_id}/ignore")
def ignore_candidate(project_id: int, candidate_id: int, body: Ignore | None = None,
                     current_user: User = Depends(require_role(*CREATOR_ROLES)), db: Session = Depends(get_db)):
    project = _get_project_or_404(db, project_id)
    candidate = _candidate(db, project, candidate_id)
    if candidate.candidate_status != "available":
        raise HTTPException(409, f"This detected revision was already {candidate.candidate_status}")
    shop_drawings.ignore_candidate(db, project, candidate, current_user, (body.reason if body else "").strip())
    drawing = candidate.drawing
    project_state.record_change(db, project.id, "drawing", "updated", entity_id=drawing.id, system=drawing.system_code)
    db.commit()
    activity.record(db, current_user, "drawings.revision_ignored",
                    f"Ignored the detected {drawing.drawing_reference} {candidate.revision}", project=project,
                    entity_type="shop_drawing", entity_id=drawing.id, detail={"system": drawing.system_code, "revision": candidate.revision})
    return shop_drawings.detail(db, project, drawing)


# --- Open folder --------------------------------------------------------------------------------


class OpenFolder(BaseModel):
    # relative to the project folder: a file (shown selected) or a folder; none for the system's shop drawings folder
    path: str | None = None
    system: str | None = None


LOCAL = {"127.0.0.1", "::1", "localhost"}
_PROXY_HEADERS = ("x-forwarded-for", "forwarded", "x-real-ip", "x-forwarded-host")


def _local_desktop(request: Request) -> str | None:
    """Why this request may not open a window on this PC, or None when it
    may: the platform allows it, it came from a loopback address with no
    proxy in front of it, to a loopback host, from a page on a loopback
    origin -- a browser on this very PC, not one somewhere else."""
    if not get_settings().desktop_actions_enabled:
        return "Opening folders is switched off on this server (DESKTOP_ACTIONS_ENABLED)."
    host = request.client.host if request.client else ""
    if host not in LOCAL:
        return "Folders open on the PC the platform runs on: open the page there, or use the path shown."
    if any(request.headers.get(h) for h in _PROXY_HEADERS):
        return "The request came through a proxy: folders open only for a browser on the PC the platform runs on."
    served = (request.headers.get("host") or "").split(":")[0].strip("[]").lower()
    if served and served not in LOCAL:
        return "The page is not served on this PC's own address: use the path shown."
    origin = request.headers.get("origin") or request.headers.get("referer")
    if origin:
        origin_host = (urlsplit(origin).hostname or "").lower()
        if origin_host not in LOCAL:
            return "The page is open from another address: use the path shown."
    return None


@router.post("/projects/{project_id}/drawings/open-folder", status_code=204)
def open_folder(project_id: int, body: OpenFolder, request: Request,
                _current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Open the folder in Windows Explorer on the PC the platform runs on --
    so only for a page opened on that PC: from anywhere else it would open
    on someone else's screen. The system's own folder: never the fire
    alarm's for an emergency lighting page."""
    refused = _local_desktop(request)
    if refused:
        raise HTTPException(403, refused)
    project = _get_project_or_404(db, project_id)
    if not project.source_folder_path:
        raise HTTPException(404, "This project has no folder")
    if body.path is None:
        system, _systems = _system(db, project, body.system)
        relative = shop_drawings_folder(system)
    else:
        relative = body.path
    root = Path(project.source_folder_path).resolve()
    target = (root / relative).resolve()
    if not target.is_relative_to(root):
        raise HTTPException(403, "Outside the project folder")
    while not target.exists() and target != root:  # the nearest folder that is there
        target = target.parent
    if os.name != "nt":
        raise HTTPException(501, "Opening a folder needs Windows")
    args = ["explorer", f"/select,{target}"] if target.is_file() else ["explorer", str(target)]
    subprocess.Popen(args)  # noqa: S603 -- a path inside the project folder, checked above


def folder_for(db: Session, project: Project, system: str | None) -> str:
    """The relative folder Open folder would open, for the page to show."""
    code, _systems = _system(db, project, system)
    return shop_drawings_folder(code)


# --- Actions Required: the contractor's drawings, by folder -------------------------------------

REQUESTED = "drawings.requested"


def _required(db: Session, project, system: str) -> dict:
    return drawing_requirements.status(db, project, system)


@router.get("/projects/{project_id}/drawings/required")
def required(project_id: int, system: str | None = None, _current_user: User = Depends(get_current_user),
             db: Session = Depends(get_db)):
    """Each item the contractor must hand over before a system's shop
    drawings start: the material approval from the register (approved /
    not approved), the documents from the project folder (received / not
    received), and when each was requested."""
    project = _get_project_or_404(db, project_id)
    return _required(db, project, drawing_requirements.check_system(project, system))


@router.get("/projects/{project_id}/drawings/required/export.xlsx")
def export_required(project_id: int, system: str | None = None,
                    _current_user: User = Depends(get_current_user),
                    db: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    project = _get_project_or_404(db, project_id)
    data = _required(db, project, drawing_requirements.check_system(project, system))
    wb = Workbook()
    ws = wb.active
    ws.title = f"Actions Required {data['system']}"[:31]
    ws.append([f"EP-{project.ep_number} {project.project_name or ''} - Actions Required - {data['system']}"])
    ws.append([f"{data['readiness']}; a document is received when its folder in the project folder holds a file; "
               f"the approval is the material submittal register's; as of {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    ws.append(["#", "Document Category", "Description / Scope", "Format", "Status", "Requested Date", "Received / Approved Date",
               "Remarks", "Source"])
    n = 0
    for g in data["groups"]:
        ws.append(["", g["name"]])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        for i in g["items"]:
            n += 1
            ws.append([n, i["name"], i["purpose"], i["format"], i["status_label"],
                       (i["requested_at"] or "")[:10] or "-", (i["received_date"] or "")[:10] or "-", i["remarks"],
                       "Material submittal register" if i["kind"] == "approval" else i["folder"]])
            ws.cell(row=ws.max_row, column=5).fill = PatternFill("solid", fgColor="C6EFCE" if i["received"] else "FFC7CE")
    for c in ws[4]:
        c.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=13)
    for index, width in enumerate((5, 32, 50, 12, 14, 15, 18, 40, 40), 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for r in ws.iter_rows(min_row=5):
        for c in r:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"EP-{project.ep_number} Actions Required {data['system']}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


class RequestItems(BaseModel):
    keys: list[str]
    system: str | None = None


@router.post("/projects/{project_id}/drawings/required/request")
def request_items(project_id: int, body: RequestItems,
                  _current_user: User = Depends(require_role(*CREATOR_ROLES)),
                  db: Session = Depends(get_db)):
    """The email asking the contractor for the items. Generating it records
    nothing: the request is noted when it was sent (/request/sent)."""
    project = _get_project_or_404(db, project_id)
    system = drawing_requirements.check_system(project, body.system)
    keys = drawing_requirements.check_keys(system, body.keys)
    return {**required_drawings.request_text(project, keys), "system": system, "keys": keys}


@router.post("/projects/{project_id}/drawings/required/request/sent")
def request_sent(project_id: int, body: RequestItems,
                 current_user: User = Depends(require_role(*CREATOR_ROLES)),
                 db: Session = Depends(get_db)):
    """The request was sent: noted against each item, with who sent it."""
    project = _get_project_or_404(db, project_id)
    system = drawing_requirements.check_system(project, body.system)
    keys = drawing_requirements.check_keys(system, body.keys)
    drawing_requirements.mark_sent(db, project, system, keys, current_user)
    names = [required_drawings.BY_KEY[k].name for k in keys]
    shop_drawings.event(db, project, "requirement.requested", f"Requested from the contractor: {', '.join(names)}",
                        system=system, user=current_user, detail={"keys": keys})
    project_state.record_change(db, project.id, "drawing", "updated", system=system)
    db.commit()
    activity.record(db, current_user, REQUESTED, f"Requested from the contractor: {', '.join(names)}",
                    project=project, entity_type="project", entity_id=project.id, detail={"keys": keys, "system": system})
    return _required(db, project, system)
