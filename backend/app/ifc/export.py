"""The BOQ as per IFC drawings as an Excel workbook: Info, Fire Alarm,
Emergency Lighting, the floors tables, Pending Verification and
Occurrences. The original tool's export, unchanged."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def workbook(d, r: dict) -> io.BytesIO:
    """d: the drawing; r: it resolved (`resolve.resolved_drawing`), with
    every symbol on the floor plans answered."""
    groups = r["groups"]
    fi = r["floor_info"]
    single_floor = fi["mode"] == "single"

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(title: str, headers: list[str], rows: list[list], widths: list[int]):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        return ws

    boq = r["floor_boq"]
    blocks: dict[int, set] = {}
    for g in groups:
        if g["status"] == "verified" and g["device_type"]:
            blocks.setdefault(g["device_type"]["id"], set()).update(g["block_names"].keys())

    def block_names(dt: dict) -> str:
        return "; ".join(sorted(blocks.get(dt["id"], ())))[:500]

    def single_floor_rows(category: str) -> list[list]:
        rows = [[i, fi["floor_name"], b["device_type"]["name"], b["device_type"]["code"], b["qty"], b["device_type"]["unit"], block_names(b["device_type"])]
                for i, b in enumerate(boq[category]["building"], start=1)]
        if rows:
            rows.append(["", "", "TOTAL", "", boq[category]["qty"], "", ""])
        return rows

    floor_total, building = "Floor total", f"BUILDING TOTAL ({fi['floors']} floors)"

    def floor_wise_rows(category: str) -> list[list]:
        """Floor by floor in sheet order: each floor's devices, then its
        total; the building total per device underneath."""
        rows: list[list] = []
        n = 0
        for f in boq[category]["floors"]:
            for x in f["rows"]:
                n += 1
                dt = x["device_type"]
                rows.append([n, f["floor_name"], f["sheet"], dt["name"], dt["code"], x["per_floor"], f["multiplier"], x["qty"], dt["unit"], ""])
            if f["rows"]:
                rows.append(["", f["floor_name"], f["sheet"], floor_total, "", f["per_floor"], f["multiplier"], f["qty"], "", ""])
        if rows:
            rows.append([None] * 10)
            for b in boq[category]["building"]:
                dt = b["device_type"]
                rows.append(["", building, "", dt["name"], dt["code"], "", "", b["qty"], dt["unit"], block_names(dt)])
            rows.append(["", building, "", "TOTAL", "", "", "", boq[category]["qty"], "", ""])
        return rows

    wb.remove(wb.active)
    sub_fill = PatternFill("solid", fgColor="E8EEF5")
    total_fill = PatternFill("solid", fgColor="D9E2F3")
    for category, title in (("fire_alarm", "Fire Alarm"), ("emergency_light", "Emergency Lighting")):
        if single_floor:
            ws = sheet(title, ["S.No", "Floor", "Device", "Code", "Qty", "Unit", "Block names on drawing"],
                       single_floor_rows(category), [6, 28, 42, 10, 8, 8, 80])
            if ws.max_row > 1:
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
            continue
        rows = floor_wise_rows(category)
        ws = sheet(title, ["S.No", "Floor", "Sheet", "Device", "Code", "Qty per floor", "No. of floors", "Total qty", "Unit", "Block names on drawing"],
                   rows, [6, 30, 10, 42, 10, 13, 13, 11, 8, 80])
        last_floor_row = 1
        for i, row in enumerate(rows, start=2):
            if row[3] == floor_total:
                last_floor_row = i
                for c in ws[i]:
                    c.font = Font(bold=True)
                    c.fill = sub_fill
            elif row[1] == building:
                ws.cell(row=i, column=2).font = Font(bold=True)
                ws.cell(row=i, column=8).font = Font(bold=True)
                if row[3] == "TOTAL":
                    for c in ws[i]:
                        c.font = Font(bold=True)
                        c.fill = total_fill
        if last_floor_row > 1:
            ws.auto_filter.ref = f"A1:J{last_floor_row}"

    # floor by floor: each plan sheet's count, the floors it stands for, and the product
    plan_sheets = [s for s in r["sheets"] if s["kind"] == "plan"]
    other_sheets = [s for s in r["sheets"] if s["kind"] != "plan"]
    for category, title in (("fire_alarm", "Floors - Fire Alarm"), ("emergency_light", "Floors - Emergency")):
        codes: dict[str, str] = {}
        cell: dict[tuple[str, str], int] = {}
        for g in groups:
            if g["status"] != "verified" or g["device_type"]["category"] != category:
                continue
            code = g["device_type"]["code"]
            codes[code] = g["device_type"]["name"]
            for sh, n in g.get("by_sheet", {}).items():
                cell[(sh, code)] = cell.get((sh, code), 0) + n
        order = sorted(codes)
        headers = ["Sheet", "Floor (title block)", "Floors"] + order + ["Floor total"]
        rows = []
        for sh in plan_sheets:
            vals = [cell.get((sh["name"], c), 0) for c in order]
            rows.append([sh["name"], sh["floor_name"], sh["multiplier"]] + vals + [sum(vals)])
        build = [sum(cell.get((sh["name"], c), 0) * sh["multiplier"] for sh in plan_sheets) for c in order]
        rows.append(["", "BUILDING TOTAL (x floors)", ""] + build + [sum(build)])
        for sh in other_sheets:
            vals = [cell.get((sh["name"], c), 0) for c in order]
            if any(vals):
                rows.append([sh["name"], f"{sh['title']} - NOT COUNTED", ""] + vals + [sum(vals)])
        ws = sheet(title, headers, rows, [10, 40, 8] + [9] * len(order) + [12])
        ws.cell(row=len(plan_sheets) + 2, column=2).font = Font(bold=True)

    pending = [
        [g["label"] or "(no letters)", g["status"], g["count"], "; ".join(g["block_names"].keys())[:300],
         (g.get("suggestion") or {}).get("reason", "")]
        for g in groups if g["status"] in ("suggested", "unknown")
    ]
    sheet("Pending Verification", ["Letters", "Status", "Instances", "Block names", "Suggestion"], pending, [16, 12, 10, 70, 50])

    occ_rows = []
    for g in groups:
        if g["status"] != "verified":
            continue
        dt = g["device_type"]
        m = g.get("match") or {}
        how = {"library": f"Library match {m.get('score', 0):.0%}", "family": "Revit family match"}.get(m.get("kind"), "Verified symbol")
        for o in g["occurrences"]:
            occ_rows.append([dt["name"], dt["code"], how, o.get("sheet", ""), o["block_name"], g["label"], o["layer"], o["x"], o["y"], o["rotation"], o["space"],
                             "; ".join(f"{k}={v}" for k, v in (o.get("attrs") or {}).items())])
    sheet("Occurrences", ["Device", "Code", "Identified by", "Sheet", "Block", "Letters", "Layer", "X", "Y", "Rotation", "Placed in", "Attributes"], occ_rows,
          [36, 8, 20, 12, 50, 12, 18, 12, 12, 9, 30, 30])

    info = wb.create_sheet("Info", 0)
    info.append(["Drawing", d.filename])
    info.append(["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")])
    info.append(["Units", d.units])
    if single_floor:
        info.append(["Drawing covers", f"One floor: {fi['floor_name']}"])
    else:
        info.append(["Drawing covers", f"Multiple floors: {fi['plans']} floor plan(s) standing for {fi['floors']} floors"])
    if fi["excluded_sheets"]:
        info.append(["Schematic / detail sheets (excluded)", "; ".join(fi["excluded_sheets"])])
    info.append(["Architecture", "Found: only devices placed on the architecture are counted" if fi["architecture_found"] else "Not found in this drawing: every device on a floor plan is counted"])
    info.append(["Quantities", ("Listed floor by floor on the Fire Alarm and Emergency Lighting sheets: each floor plan's count x the floors it stands for, "
                                 "then the building total. " if not single_floor else "") +
                 "Riser, schematic and detail sheets, devices off the architecture and anything outside every sheet are not counted."])
    info.append(["Verified symbols not counted", r["totals"]["not_counted_instances"]])
    info.append(["Fire alarm devices (verified)", r["totals"]["fire_alarm"]])
    info.append(["Emergency lighting (verified)", r["totals"]["emergency_light"]])
    info.append(["Symbols pending verification", r["totals"]["suggested_symbols"] + r["totals"]["unknown_symbols"]])
    info.append(["Instances pending verification", r["totals"]["suggested_instances"] + r["totals"]["unknown_instances"]])
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 80
    for row in info.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
