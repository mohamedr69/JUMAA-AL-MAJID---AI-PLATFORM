"""The compliance statement workbook, as the company submits it.

The layout copies the company's own statements (EP-16481, EP-13705): a title
block naming the project, the product, the manufacturer and the local
supplier, then SL.NO | PROJECT SPECIFICATION - SECTION ... | <COMPANY>
COMPLIANCE | <CONTRACTOR> COMPLIANCE | REMARKS, with the specification's own
headings in bold and its outline indented. The contractor's column is left
empty for the contractor.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from app.core.config import get_settings
from app.models import ComplianceStatement, Project
from app.services.spec_finder import SYSTEMS

_THIN = Side(style="thin", color="999999")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_PART_FILL = PatternFill("solid", fgColor="D9E1F2")
_REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")


def build_workbook(project: Project, statement: ComplianceStatement, brands: set[str]) -> bytes:
    settings = get_settings()
    system = SYSTEMS.get(statement.system_code)
    system_name = system.name.upper() if system else statement.system_code
    section = ", ".join(statement.spec.get("section_numbers") or []) or "-"
    company = settings.company_name.upper()
    contractor = (project.contractor or "CONTRACTOR").upper()

    wb = Workbook()
    ws = wb.active
    ws.title = f"{statement.system_code} Compliance"[:31]
    widths = {"A": 10, "B": 78, "C": 20, "D": 20, "E": 34}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    title_lines = [
        (f"COMPLIANCE STATEMENT - {system_name} SYSTEM", Font(bold=True, size=14)),
        (f"PROJECT NAME : {(project.project_name or '').upper()}", Font(bold=True)),
        (f"EP NUMBER : EP-{project.ep_number}", Font(bold=False)),
        (f"CONSULTANT : {(project.consultant or '-').upper()}    MAIN CONTRACTOR : {contractor}", Font(bold=False)),
        (f"PRODUCT DESCRIPTION : {system_name} SYSTEM", Font(bold=False)),
        (f"MANUFACTURER : {', '.join(sorted(b.upper() for b in brands)) or '-'}", Font(bold=False)),
        (f"LOCAL SUPPLIER : M/s. {company}", Font(bold=False)),
    ]
    row = 1
    for text, font in title_lines:
        ws.cell(row=row, column=1, value=text).font = font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
    row += 1

    header_row = row
    headings = ["SL.NO", f"PROJECT SPECIFICATION - SECTION {section}", f"{company} COMPLIANCE",
                f"{contractor} COMPLIANCE", "REMARKS"]
    for column, text in enumerate(headings, start=1):
        cell = ws.cell(row=row, column=column, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BOX
    ws.row_dimensions[row].height = 32
    row += 1

    for item in statement.rows:
        label = item["label"] if item["level"] <= 1 else f"{item['label']}."
        text = item["text"] if item["level"] > 0 else f"{item['label']} - {item['text']}"
        values = [label if item["level"] > 0 else "", text, item.get("response", ""), "", item.get("remark", "")]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = _BOX
            cell.alignment = Alignment(
                vertical="top", wrap_text=True,
                horizontal="center" if column in (1, 3, 4) else "left",
                indent=min(max(item["level"] - 1, 0), 4) if column == 2 else 0,
            )
        if item.get("heading"):
            for column in range(1, 6):
                ws.cell(row=row, column=column).font = Font(bold=True)
            if item["level"] == 0:
                for column in range(1, 6):
                    ws.cell(row=row, column=column).fill = _PART_FILL
        elif item.get("state") == "review":
            ws.cell(row=row, column=3).fill = _REVIEW_FILL
        # Rough height so wrapped clauses are readable without resizing.
        ws.row_dimensions[row].height = max(15, 15 * (len(text) // 90 + 1))
        row += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.oddFooter.center.text = "Page &P of &N"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
