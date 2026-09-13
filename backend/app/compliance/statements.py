"""Reading a compliance statement: the company's past ones, or one submitted
for checking.

They are all the same table under different headings -- the specification's
clause, its number, and the answer ("Comply", "Noted", "Not applicable as per
the design drawing", "By others") -- laid out in whatever columns the engineer
of the day chose, in .xlsx, .xls or .docx. So the columns are found by what
they hold rather than by what they are called: the answer column is the one
full of answers, the clause is the longest text to its left, the number is
the short label before that, and the remark is what follows.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree

STATEMENT_PARSER_VERSION = "statement-1"
MAX_ROWS = 5000

# What an answer cell says. Kept broad: past statements use every variant.
RESPONSE_RE = re.compile(
    r"^\s*(complied|comply|complies|compliance|noted|note\b|not\s+applicable|n\s*/\s*a\b|na\b|by\s+others|"
    r"deviat|partially|agreed|accepted|acknowledged|yes\b|provided|refer|equivalent|approved\s+equal|"
    r"will\s+comply|shall\s+comply|clarification|not\s+compl|non[\s-]?compl|excluded|not\s+in\s+scope|"
    r"out\s+of\s+scope|to\s+be\s+confirmed|tbc\b)",
    re.IGNORECASE,
)
_LABEL_RE = re.compile(r"^\s*(PART\s*\d+|\d{1,2}(\.\d{1,2}){0,4}\.?|[A-Za-z][.)]?|\(?[ivx]{1,4}[.)]|[a-z]\))\s*$", re.IGNORECASE)
_COMPANY_RE = re.compile(r"al\s*arabia", re.IGNORECASE)

# The answers a statement is written in, and the canonical spelling of each.
CANONICAL = (
    ("Not applicable", re.compile(r"not\s+applicable|^\s*n\s*/?\s*a\b|excluded|not\s+in\s+scope|out\s+of\s+scope", re.I)),
    ("By others", re.compile(r"by\s+others", re.I)),
    ("Deviation", re.compile(r"deviat|not\s+compl|non[\s-]?compl", re.I)),
    ("Clarification required", re.compile(r"clarification|to\s+be\s+confirmed|tbc\b", re.I)),
    ("Complied with remark", re.compile(r"partially|equivalent|approved\s+equal", re.I)),
    ("Noted", re.compile(r"noted|note\b|acknowledged|agreed|accepted", re.I)),
    ("Comply", re.compile(r"compl|yes\b|provided|refer", re.I)),
)
RESPONSES = ("Comply", "Noted", "Complied with remark", "Not applicable", "By others", "Deviation", "Clarification required")


def canonical(response: str) -> str | None:
    for name, regex in CANONICAL:
        if regex.search(response or ""):
            return name
    return None


@dataclass
class StatementRow:
    label: str
    text: str
    response: str          # as written
    remark: str

    @property
    def answer(self) -> str | None:
        return canonical(self.response)


@dataclass
class Statement:
    rows: list[StatementRow]
    title_lines: list[str] = field(default_factory=list)   # what sits above the table
    response_heading: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def answered(self) -> list[StatementRow]:
        return [r for r in self.rows if r.answer]


def _clean(value) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return re.sub(r"\s+", " ", text).strip()


def interpret(table: list[list[str]]) -> Statement | None:
    """Find the columns of a compliance table by what they hold."""
    table = [row for row in table if any(cell for cell in row)][:MAX_ROWS]
    if len(table) < 5:
        return None
    width = max(len(row) for row in table)
    table = [row + [""] * (width - len(row)) for row in table]

    hits = [sum(1 for row in table if row[c] and len(row[c]) <= 120 and RESPONSE_RE.match(row[c])) for c in range(width)]
    best = max(hits)
    if best < 4:
        return None
    candidates = [c for c in range(width) if hits[c] >= max(4, best // 3)]
    # Several answer columns (the company's, then the contractor's): the
    # company's own when a heading names it, else the first.
    heading_rows = table[:15]
    response_col = next(
        (c for c in candidates if any(_COMPANY_RE.search(row[c]) for row in heading_rows)),
        candidates[0],
    )

    left = list(range(response_col))
    if not left:
        return None
    lengths = {c: sum(len(row[c]) for row in table if not RESPONSE_RE.match(row[c] or "x")) for c in left}
    text_col = max(lengths, key=lengths.get)
    label_col = None
    if text_col > 0:
        label_hits = {c: sum(1 for row in table if row[c] and _LABEL_RE.match(row[c])) for c in range(text_col)}
        if label_hits and max(label_hits.values()) >= 3:
            label_col = max(label_hits, key=label_hits.get)
    remark_cols = [c for c in range(response_col + 1, width) if c not in candidates]

    rows: list[StatementRow] = []
    title_lines: list[str] = []
    response_heading = None
    started = False
    for row in table:
        response = row[response_col]
        text = row[text_col]
        label = row[label_col] if label_col is not None else ""
        # Clause text that spills into a neighbouring merged cell.
        if not text:
            spill = [row[c] for c in range(response_col) if c not in (label_col,) and row[c]]
            text = max(spill, key=len) if spill else ""
        remark = " ".join(row[c] for c in remark_cols if row[c] and not RESPONSE_RE.match(row[c]))
        is_answer = bool(response and RESPONSE_RE.match(response) and len(response) <= 160)
        if not started and not is_answer:
            if re.search(r"complian", response or "", re.IGNORECASE):
                response_heading = response
                started = True
                continue
            title_lines += [cell for cell in row if cell]
            continue
        started = True
        if not text and not label:
            continue
        if text and re.fullmatch(r"(sl\.?\s*no\.?|clause|item|description|specification.*)", text, re.IGNORECASE):
            continue
        rows.append(StatementRow(label=label, text=text, response=response if is_answer else "", remark=remark if is_answer else remark))
    if len([r for r in rows if r.answer]) < 4:
        return None
    return Statement(rows=rows, title_lines=title_lines[:30], response_heading=response_heading)


def _xlsx_tables(content: bytes) -> list[list[list[str]]]:
    import openpyxl

    # Not read-only: one "Comply" merged down the rows of a list answers every
    # row of it, and only a full load knows the merged ranges.
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    tables = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS):
                rows.append([_clean(v) for v in row[:30]])
            for merged in sheet.merged_cells.ranges:
                value = rows[merged.min_row - 1][merged.min_col - 1] if merged.min_row <= len(rows) and merged.min_col <= 30 else ""
                for r in range(merged.min_row, min(merged.max_row, len(rows)) + 1):
                    for c in range(merged.min_col, min(merged.max_col, 30) + 1):
                        if r - 1 < len(rows) and c - 1 < len(rows[r - 1]) and not rows[r - 1][c - 1]:
                            rows[r - 1][c - 1] = value
            tables.append(rows)
    finally:
        workbook.close()
    return tables


def _xls_tables(content: bytes) -> list[list[list[str]]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover -- optional dependency
        raise ValueError("reading .xls needs the xlrd package") from exc
    book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    tables = []
    for sheet in book.sheets():
        rows = [[_clean(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, 30))] for r in range(min(sheet.nrows, MAX_ROWS))]
        for r0, r1, c0, c1 in sheet.merged_cells:
            if r0 >= len(rows) or c0 >= 30:
                continue
            value = rows[r0][c0]
            for r in range(r0, min(r1, len(rows))):
                for c in range(c0, min(c1, 30)):
                    if not rows[r][c]:
                        rows[r][c] = value
        tables.append(rows)
    return tables


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_tables(content: bytes) -> list[list[list[str]]]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    tables = []
    rows: list[list[str]] = []
    # Every table in the document read as one: statements split across pages
    # are often several tables with the same columns.
    for table in root.iter(f"{_W}tbl"):
        for tr in table.findall(f"{_W}tr"):
            cells = []
            for tc in tr.findall(f"{_W}tc"):
                paragraphs = ["".join(t.text or "" for t in p.iter(f"{_W}t")) for p in tc.iter(f"{_W}p")]
                cells.append(_clean(" ".join(p for p in paragraphs if p)))
            rows.append(cells)
    if rows:
        tables.append(rows)
    return tables


def read_statement(content: bytes, filename: str) -> Statement | None:
    """The compliance table in a workbook or document, or None when it holds
    none. The sheet with the most answers wins."""
    suffix = Path(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        tables = _xlsx_tables(content)
    elif suffix == ".xls":
        tables = _xls_tables(content)
    elif suffix == ".docx":
        tables = _docx_tables(content)
    else:
        raise ValueError(f"{suffix or 'this file'} is not a statement format that can be read (.xlsx, .xls, .docx)")
    found = [s for s in (interpret(t) for t in tables) if s is not None]
    return max(found, key=lambda s: len(s.answered)) if found else None
