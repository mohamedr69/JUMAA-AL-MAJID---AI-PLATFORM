"""The intake gate: is this document the project's, whole, and readable?

Every value the platform reads -- a DRF field, a BOQ line -- is only as good
as the document it came from, and the archive has documents that are not
what their place suggests: EP-31112's DRF sits under a different contractor's
folder than the project's own; EP-30058's folder holds "EP-30088 Design Sheet
ELS.pdf". Reading those faithfully produces confident, wrong numbers.

`check` looks at one document and reports findings, each with a severity:

  blocked  -- the document cannot be trusted until someone settles it
              (missing, unreadable, another project's, outside the project
              folder, fewer pages than it says it has)
  warning  -- worth a look, not disqualifying (page numbers out of order,
              the same content attached twice)

A blocked finding can be acknowledged by a design manager or admin with a
reason ("the file is misnamed; it is this project's sheet"): it then counts
as a warning, and who said so is kept.

Printed page numbers come from the text layer, or an OCR of the page's top
and bottom strips when a page has none -- which is how a scan of page 1 of 2
saved as a one-page PDF is caught.
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.timeutils import utc_now
from app.models import Project, ProjectDocument, User

INTAKE_VERSION = "intake-2026-09-15.1"

BLOCKED = "blocked"
WARNING = "warning"
OK = "ok"
UNCHECKED = "unchecked"

EP_RE = re.compile(r"(?<![A-Za-z0-9])EP[\s\-_#.]*(\d{4,6})(?!\d)", re.IGNORECASE)
PRINTED_PAGE_RE = re.compile(r"\bpage\s*[:#]?\s*(\d{1,3})\s*(?:of|/)\s*(\d{1,3})\b", re.IGNORECASE)
SUPPORTED = {".pdf": "application/pdf", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
             ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12"}
# A DRF scanned straight to an image is read as one (app.services.drf_extractor.DRF_EXTENSIONS).
IMAGES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".tif": "image/tiff", ".tiff": "image/tiff"}
IMAGE_MAGIC = ((b"\xff\xd8\xff", "image/jpeg"), (b"\x89PNG", "image/png"), (b"II*\x00", "image/tiff"), (b"MM\x00*", "image/tiff"))
# Pages beyond this are not OCR'd for printed numbers: the documents the
# gate is for are one to a few pages, and a 60-page upload is not scanned
# page by page on a click.
MAX_OCR_PAGES = 12
OCR_STRIP_FRACTION = 0.12
OCR_DPI = 200


@dataclass
class Finding:
    code: str
    severity: str
    message: str
    detail: dict = field(default_factory=dict)

    def as_dict(self) -> dict:
        return {"code": self.code, "severity": self.severity, "message": self.message, "detail": self.detail}


def sha256_of(path: Path) -> str | None:
    try:
        digest = hashlib.sha256()
        with open(path, "rb") as handle:
            for chunk in iter(lambda: handle.read(1 << 20), b""):
                digest.update(chunk)
        return digest.hexdigest()
    except OSError:
        return None


def ep_numbers(text: str) -> set[str]:
    return {match.group(1).lstrip("0") or "0" for match in EP_RE.finditer(text)}


def _is_within(path: Path, parent: Path) -> bool:
    try:
        return path.resolve().is_relative_to(parent.resolve())
    except (OSError, ValueError):
        return False


def _sniff(path: Path) -> tuple[str | None, str | None]:
    """(mime by content, a reason the content does not match its suffix)."""
    try:
        with open(path, "rb") as handle:
            head = handle.read(8)
    except OSError as exc:
        return None, str(exc)
    suffix = path.suffix.lower()
    if head.startswith(b"%PDF"):
        return SUPPORTED[".pdf"], None if suffix == ".pdf" else "the content is a PDF but the name says otherwise"
    if head.startswith(b"PK"):
        mime = SUPPORTED.get(suffix, "application/zip")
        return mime, None if suffix in (".xlsx", ".xlsm") else "the content is a zip/Office file but the name says otherwise"
    for magic, mime in IMAGE_MAGIC:
        if head.startswith(magic):
            return mime, None if IMAGES.get(suffix) == mime else "the content is an image of another type than the name says"
    return None, "the content is neither a PDF, an Excel workbook nor an image"


def printed_pages(document) -> dict:
    """What the pages say about themselves: every "Page x of N" found, per page."""
    numbers: list[int] = []
    totals: set[int] = set()
    per_page: dict[int, list[int]] = {}
    ocr_used = False
    for index in range(min(document.page_count, MAX_OCR_PAGES)):
        page = document[index]
        text = page.get_text() or ""
        if not text.strip():
            text = _ocr_strips(page)
            ocr_used = ocr_used or bool(text)
        found = [(int(a), int(b)) for a, b in PRINTED_PAGE_RE.findall(text) if 0 < int(a) <= int(b)]
        if found:
            per_page[index + 1] = [a for a, _ in found]
            numbers.extend(a for a, _ in found[:1])
            totals.update(b for _, b in found)
    return {"numbers": numbers, "declared_totals": sorted(totals), "pages_read": min(document.page_count, MAX_OCR_PAGES),
            "per_page": per_page, "ocr": ocr_used}


def _ocr_strips(page) -> str:
    try:
        import pymupdf
        import pytesseract
        from PIL import Image

        settings = get_settings()
        if settings.tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        pix = page.get_pixmap(dpi=OCR_DPI, colorspace=pymupdf.csGRAY)
        image = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        strip = int(image.height * OCR_STRIP_FRACTION)
        parts = [image.crop((0, 0, image.width, strip)), image.crop((0, image.height - strip, image.width, image.height))]
        return "\n".join(pytesseract.image_to_string(part, config="--psm 6") for part in parts)
    except Exception:  # noqa: BLE001 -- no OCR available: nothing printed can be read, which is not a finding
        return ""


def _page_findings(page_count: int, printed: dict) -> list[Finding]:
    findings: list[Finding] = []
    totals = printed.get("declared_totals") or []
    numbers = printed.get("numbers") or []
    if totals:
        declared = max(totals)
        if declared > page_count:
            findings.append(Finding(
                "INCOMPLETE_SOURCE", BLOCKED,
                f"The pages say there are {declared}, but the file has {page_count}: part of the document is missing.",
                {"declared": declared, "page_count": page_count, "printed_numbers": numbers}))
        elif len(totals) > 1:
            findings.append(Finding("PAGE_TOTALS_DISAGREE", WARNING,
                                    f"The pages give different totals ({', '.join(map(str, totals))}): "
                                    "two documents may be combined in one file.", {"totals": totals}))
        if numbers:
            expected = list(range(1, len(numbers) + 1))
            if sorted(numbers) != expected or numbers != sorted(numbers):
                findings.append(Finding("PAGE_SEQUENCE_GAP", WARNING,
                                        f"Printed page numbers run {', '.join(map(str, numbers))}: a page may be missing, "
                                        "repeated or out of order.", {"printed_numbers": numbers}))
    return findings


def _identity_findings(project: Project, path: Path, role: str, archive_root: Path | None) -> list[Finding]:
    findings: list[Finding] = []
    project_ep = project.ep_number.lstrip("0")

    in_name = ep_numbers(path.name)
    if in_name and project_ep not in in_name:
        findings.append(Finding(
            "EP_MISMATCH_FILENAME", BLOCKED,
            f"The file is named for EP-{', EP-'.join(sorted(in_name))}, not EP-{project.ep_number}: it may be another "
            "project's document.", {"filename_ep": sorted(in_name)}))

    settings = get_settings()
    uploads = Path(settings.uploads_root) / f"EP-{project.ep_number}"
    source = Path(project.source_folder_path) if project.source_folder_path else None
    if _is_within(path, uploads):
        return findings

    if source is not None:
        folder_eps = ep_numbers(source.name)
        if folder_eps and project_ep not in folder_eps:
            findings.append(Finding("SOURCE_FOLDER_EP_MISMATCH", BLOCKED,
                                    f"The project's folder ({source.name}) is named for another EP number.",
                                    {"folder_ep": sorted(folder_eps)}))
        if not _is_within(path, source):
            findings.append(Finding(
                "OUTSIDE_SOURCE_FOLDER", BLOCKED,
                f"The {role.replace('_', ' ')} is not inside the project's folder: the project folder is under "
                f"{_top(source, archive_root)} but the file is under {_top(path, archive_root)}.",
                {"source_folder": source.name}))
        else:
            # Folders between the project folder and the file that name
            # another EP number ("EP-30088 Old").
            between = path.resolve().relative_to(source.resolve()).parts[:-1]
            other = set().union(*(ep_numbers(part) for part in between)) - {project_ep} if between else set()
            if other:
                findings.append(Finding("EP_MISMATCH_FOLDER", WARNING,
                                        f"The file sits in a subfolder named for EP-{', EP-'.join(sorted(other))}.",
                                        {"folder_ep": sorted(other)}))
    elif archive_root is not None and not _is_within(path, archive_root):
        findings.append(Finding("OUTSIDE_ARCHIVE", BLOCKED, "The file is outside the project archive and the uploads folder."))
    return findings


def _top(path: Path, archive_root: Path | None) -> str:
    """The contractor-level folder a path sits under, for a readable message."""
    try:
        if archive_root is not None and _is_within(path, archive_root):
            parts = path.resolve().relative_to(archive_root.resolve()).parts
            return parts[0] if parts else path.name
    except (OSError, ValueError):
        pass
    parts = path.parts
    return parts[-3] if len(parts) >= 3 else str(path.parent)


def check(project: Project, role: str, path: Path, *, archive_root: Path | None = None) -> dict:
    """Everything the gate knows about one document, as ProjectDocument fields."""
    record: dict = {"filename": path.name, "findings": [], "sha256": None, "size": None, "mime": None,
                    "page_count": None, "printed_pages": None}
    findings: list[Finding] = []
    source = Path(project.source_folder_path) if project.source_folder_path else None
    try:
        record["relative_path"] = str(path.resolve().relative_to(source.resolve())) if source and _is_within(path, source) else path.name
    except (OSError, ValueError):
        record["relative_path"] = path.name

    findings.extend(_identity_findings(project, path, role, archive_root))

    if not path.is_file():
        findings.append(Finding("MISSING_FILE", BLOCKED, "The file is not there any more (moved, renamed or not synced)."))
    else:
        record["size"] = path.stat().st_size
        if record["size"] == 0:
            findings.append(Finding("EMPTY_FILE", BLOCKED, "The file is empty."))
        else:
            record["sha256"] = sha256_of(path)
            mime, mismatch = _sniff(path)
            record["mime"] = mime
            suffix = path.suffix.lower()
            readable = set(SUPPORTED) | (set(IMAGES) if role == "drf" else set())
            if suffix not in readable:
                findings.append(Finding("UNSUPPORTED_FORMAT", BLOCKED,
                                        f"{path.suffix or 'This file type'} is not a format the platform reads for a "
                                        f"{role.replace('_', ' ')} (PDF or Excel{', or an image' if role == 'drf' else ''})."))
            elif mismatch:
                findings.append(Finding("CONTENT_TYPE_MISMATCH", BLOCKED, f"The file's {mismatch}."))
            elif mime == SUPPORTED[".pdf"]:
                findings.extend(_pdf_findings(path, record))
            elif mime in IMAGES.values():
                findings.extend(_image_findings(path, record))
            else:
                findings.extend(_workbook_findings(path, record))

    record["findings"] = [f.as_dict() for f in findings]
    return record


def _pdf_findings(path: Path, record: dict) -> list[Finding]:
    try:
        import pymupdf

        with pymupdf.open(str(path)) as document:
            if document.needs_pass or document.is_encrypted:
                return [Finding("UNREADABLE", BLOCKED, "The PDF is password-protected.")]
            record["page_count"] = document.page_count
            if document.page_count == 0:
                return [Finding("UNREADABLE", BLOCKED, "The PDF has no pages.")]
            printed = printed_pages(document)
    except Exception as exc:  # noqa: BLE001 -- a malformed PDF raises anything
        return [Finding("UNREADABLE", BLOCKED, f"The PDF could not be opened: {exc}")]
    record["printed_pages"] = printed
    return _page_findings(record["page_count"], printed)


def _image_findings(path: Path, record: dict) -> list[Finding]:
    try:
        from PIL import Image

        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            record["page_count"] = getattr(image, "n_frames", 1)
            if image.width * image.height > 120_000_000:
                return [Finding("UNREADABLE", BLOCKED, "The image is too large to read safely.")]
    except Exception as exc:  # noqa: BLE001
        return [Finding("UNREADABLE", BLOCKED, f"The image could not be opened: {exc}")]
    return []


def _workbook_findings(path: Path, record: dict) -> list[Finding]:
    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad is not None:
                return [Finding("UNREADABLE", BLOCKED, f"The workbook is damaged ({bad}).")]
            # A zip bomb is refused before openpyxl expands it.
            if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                return [Finding("UNREADABLE", BLOCKED, "The workbook expands to more than 200 MB.")]
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        record["page_count"] = len(workbook.sheetnames)
        workbook.close()
    except Exception as exc:  # noqa: BLE001
        return [Finding("UNREADABLE", BLOCKED, f"The workbook could not be opened: {exc}")]
    return []


def status_of(findings: list[dict], acknowledged: list[dict]) -> str:
    acked = {a.get("code") for a in acknowledged or []}
    severities = [WARNING if f["severity"] == BLOCKED and f["code"] in acked else f["severity"] for f in findings]
    if BLOCKED in severities:
        return BLOCKED
    if WARNING in severities:
        return WARNING
    return OK


def project_documents(project: Project) -> list[tuple[str, str | None, Path]]:
    """(role, system code, path) of every document the project draws on."""
    documents: list[tuple[str, str | None, Path]] = []
    if project.drf_document_path:
        documents.append(("drf", None, Path(project.drf_document_path)))
    documents.extend(("design_sheet", sheet.system_code, Path(sheet.document_path)) for sheet in project.design_sheets)
    return documents


def run(db: Session, project: Project, ctx=None) -> list[ProjectDocument]:
    """Check every document on the project and store what was found.
    Documents no longer on the project are dropped from the record."""
    settings = get_settings()
    archive_root = Path(settings.projects_root) if settings.projects_root else None
    # Its own rows only: the document index (app.services.document_sync)
    # keeps every other file of the folder in the same table.
    existing = {(doc.role, doc.path): doc for doc in db.query(ProjectDocument)
                .filter(ProjectDocument.project_id == project.id, ProjectDocument.role.in_(("drf", "design_sheet")))}
    now = utc_now()
    kept: list[ProjectDocument] = []
    documents = project_documents(project)
    for index, (role, system_code, path) in enumerate(documents):
        if ctx is not None:
            ctx.progress(index, len(documents), f"Checking {path.name} ({index + 1} of {len(documents)})")
        record = check(project, role, path, archive_root=archive_root)
        row = existing.pop((role, str(path)), None)
        if row is None:
            row = ProjectDocument(project_id=project.id, role=role, path=str(path), first_seen_at=now, acknowledged=[])
            db.add(row)
        row.system_code = system_code
        for name, value in record.items():
            setattr(row, name, value)
        row.last_seen_at = now
        row.checked_at = now
        row.intake_version = INTAKE_VERSION
        kept.append(row)

    # Content attached twice, on this project or another.
    by_hash: dict[str, list[ProjectDocument]] = {}
    for row in kept:
        if row.sha256:
            by_hash.setdefault(row.sha256, []).append(row)
    for sha, rows in by_hash.items():
        if len(rows) > 1:
            for row in rows:
                others = [r.filename for r in rows if r is not row]
                row.findings = row.findings + [Finding("DUPLICATE_CONTENT", WARNING,
                                                       f"The same file content is attached again as {', '.join(others)}.",
                                                       {"others": others}).as_dict()]
    hashes = [row.sha256 for row in kept if row.sha256]
    if hashes:
        elsewhere = (db.query(ProjectDocument.sha256, Project.ep_number)
                     .join(Project, Project.id == ProjectDocument.project_id)
                     .filter(ProjectDocument.sha256.in_(hashes), ProjectDocument.project_id != project.id).all())
        for sha, ep in elsewhere:
            for row in by_hash.get(sha, []):
                row.findings = row.findings + [Finding("SAME_CONTENT_ON_ANOTHER_PROJECT", BLOCKED,
                                                       f"The same file is attached to EP-{ep}: it can belong to only one of them.",
                                                       {"ep_number": ep}).as_dict()]
    for row in kept:
        row.intake_status = status_of(row.findings, row.acknowledged)
    for stale in existing.values():
        db.delete(stale)
    db.commit()
    for row in kept:
        db.refresh(row)
    return kept


ACK_CODES_NEEDING_MANAGER = {"EP_MISMATCH_FILENAME", "OUTSIDE_SOURCE_FOLDER", "SOURCE_FOLDER_EP_MISMATCH",
                             "SAME_CONTENT_ON_ANOTHER_PROJECT", "INCOMPLETE_SOURCE", "OUTSIDE_ARCHIVE"}
NOT_ACKNOWLEDGEABLE = {"MISSING_FILE", "EMPTY_FILE", "UNREADABLE", "UNSUPPORTED_FORMAT", "CONTENT_TYPE_MISMATCH"}


class AcknowledgeRefused(Exception):
    pass


def acknowledge(db: Session, row: ProjectDocument, code: str, reason: str, user: User) -> ProjectDocument:
    """Accept a blocked finding, with the reason, as the named user. A file
    that is missing or unreadable cannot be acknowledged away: there is
    nothing to trust."""
    if not reason or len(reason.strip()) < 10:
        raise AcknowledgeRefused("Say why the finding is acceptable (at least a short sentence).")
    if code in NOT_ACKNOWLEDGEABLE:
        raise AcknowledgeRefused("A missing or unreadable file cannot be accepted: replace or re-attach it.")
    if not any(f["code"] == code for f in row.findings or []):
        raise AcknowledgeRefused("The document has no such finding.")
    entry = {"code": code, "reason": reason.strip()[:500], "by_id": user.id, "by_name": user.full_name,
             "at": utc_now().isoformat()}
    row.acknowledged = [a for a in (row.acknowledged or []) if a.get("code") != code] + [entry]
    row.intake_status = status_of(row.findings, row.acknowledged)
    db.commit()
    db.refresh(row)
    return row
