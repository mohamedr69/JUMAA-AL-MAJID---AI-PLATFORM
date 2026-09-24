"""The reply to a consultant's comments on a submittal, and its sheet.

A submittal returned "revise and resubmit" comes back with remarks. The
reply sheet answers them one by one -- the comment as the consultant
wrote it, whether we comply, and what we say about it -- and goes back
with the next revision.

It is a document the engineer writes and re-writes before it is sent, so
it is kept (`SubmittalReply`) rather than assembled fresh each time, and
exported to the workbook the company sends.

**The comments are not read out of the consultant's scan.** A returned
submittal is a scan of a numbered table, and reading that reliably is a
different job from reading a submittal form. What the platform has is the
reply's own text where a reading found one -- usually a line or two -- so
the sheet opens with that as its first comment and the engineer adds the
rest. Seeding it with something invented would be worse than an empty
sheet: the engineer would have to check every row before trusting any.
"""

from __future__ import annotations

import io
import html
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# The company's own name on the sheet, where the submittal does not say
# who it was from.
DEFAULT_SUPPLIER = "Al Arabia SSD"
HEADING = PatternFill("solid", fgColor="C6EFCE")
TITLE = PatternFill("solid", fgColor="FFFFFF")
EDGE = Side(style="thin", color="404040")
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
COLUMNS = ("SN", "Consultant Comments", "{supplier} Reply", "REMARKS")
WIDTHS = (6, 62, 22, 40)


def blank_row(number: int) -> dict:
    return {"sn": number, "comment": "", "reply": "Comply", "remark": ""}


def seed_rows(reply_text: str | None) -> list[dict]:
    """The sheet as it opens. The consultant's reply where the platform
    read one, split on the sentence breaks a short remark uses, so an
    answer covering two points opens as two rows to answer."""
    text = " ".join((reply_text or "").split())
    if not text:
        return [blank_row(1)]
    parts = [part.strip() for part in re.split(r"(?<=[.;])\s+", text) if part.strip()]
    return [{"sn": n, "comment": part, "reply": "Comply", "remark": ""}
            for n, part in enumerate(parts, start=1)] or [blank_row(1)]


def supplier_name(supplier: str | None) -> str:
    """The name in the reply column's heading. A submittal form writes it
    as the letter does -- "M/S. ALARABIA FOR SAFETY AND SECURITY LLC." --
    and the heading wants the company, not the salutation."""
    from app.routers.submittal import _maker

    name = _maker(supplier) or ""
    name = re.sub(r"\b(l\.?l\.?c|w\.?l\.?l|fze|est|co|company|limited|ltd)\b\.?", "", name, flags=re.IGNORECASE)
    name = " ".join(name.replace("&", " & ").split()).strip(" .,-")
    return name.title() if name else DEFAULT_SUPPLIER


def workbook(
    *,
    project_name: str,
    reference: str,
    revision: str,
    system_title: str,
    manufacturer: str | None,
    consultant: str | None,
    supplier: str | None,
    rows: list[dict],
) -> bytes:
    """The reply sheet as the company sends it: the heading block, then a
    row per comment."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Reply"
    heading = f"Reply to Consultant Comments on {system_title} submittal"
    supplier_label = supplier_name(supplier)

    def line(row: int, text: str) -> None:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True)
        cell.fill = TITLE
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center" if row == 1 else "left", vertical="center")

    line(1, heading)
    line(2, f"Project : {project_name or '—'}")
    line(3, f"Ref No : {reference} - {revision}")
    line(4, f"MANUFACTURER: {manufacturer or '—'}")
    line(5, f"Consultant : {consultant or '—'}")

    header_row = 6
    for index, title in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title.format(supplier=supplier_label))
        cell.font = Font(bold=True)
        cell.fill = HEADING
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for offset, row in enumerate(rows or [blank_row(1)], start=1):
        at = header_row + offset
        values = (row.get("sn") or offset, row.get("comment") or "", row.get("reply") or "", row.get("remark") or "")
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=at, column=index, value=value)
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="center" if index in (1, 3) else "left", vertical="center", wrap_text=True
            )
            if index == 3:
                cell.font = Font(bold=True, color="0000CC")
        # Tall enough for a wrapped comment without measuring the text.
        sheet.row_dimensions[at].height = max(30, 15 * (len(str(values[1])) // 55 + 1))

    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- the sheet as a PDF ----------------------------------------------------------
#
# What goes back to the consultant is a PDF: it is attached to the next
# revision and read, not edited. The workbook above stays for an engineer
# who wants to work on it in Excel; this is the one the Export button
# sends.

PAGE_CSS = """
body { font-family: sans-serif; font-size: 9pt; color: #111827; }
h1 { font-size: 13pt; text-align: center; margin: 0 0 10px 0; }
table { border-collapse: collapse; width: 100%; }
th, td { border: 1px solid #404040; padding: 4px 6px; vertical-align: top; }
th { background-color: #c6efce; font-size: 9pt; text-align: center; }
table.heading { margin-bottom: 10px; }
table.heading th { background-color: #ffffff; text-align: left; }
td.sn { text-align: center; }
td.reply { text-align: center; color: #1d4ed8; font-weight: bold; }
"""


def _cell(text) -> str:
    """User-written text on a page built from HTML: escaped, and with its
    own line breaks kept."""
    return html.escape(str(text or "")).replace("\n", "<br/>")


def _sheet_html(
    *,
    heading: str,
    facts: tuple,
    supplier_label: str,
    rows: list[dict],
    first_number: int,
    with_heading: bool,
) -> str:
    """One page of the sheet. The heading block is only on the first page;
    the column headings are on every one, because the story engine does
    not carry a `<thead>` over a page break of its own accord."""
    parts: list[str] = []
    if with_heading:
        parts += ["<h1>", html.escape(heading), "</h1>", '<table class="heading">']
        for label, value in facts:
            parts += ['<tr><th width="22%">', html.escape(label), "</th><td>",
                      _cell(value) or "&#8212;", "</td></tr>"]
        parts.append("</table>")
    # The story engine ignores a width in the stylesheet but honours one on
    # the cell, so the columns are set here rather than in PAGE_CSS.
    parts += ['<table><thead><tr><th width="6%">SN</th>',
              '<th width="47%">Consultant Comments</th><th width="17%">',
              html.escape(supplier_label), ' Reply</th><th width="30%">REMARKS</th>',
              "</tr></thead><tbody>"]
    for offset, row in enumerate(rows, start=first_number):
        parts += ['<tr><td class="sn">', _cell(row.get("sn") or offset),
                  '</td><td class="comment">', _cell(row.get("comment")),
                  '</td><td class="reply">', _cell(row.get("reply")),
                  '</td><td class="remark">', _cell(row.get("remark")), "</td></tr>"]
    parts.append("</tbody></table>")
    return "".join(parts)


def pdf(
    *,
    project_name: str,
    reference: str,
    revision: str,
    system_title: str,
    manufacturer: str | None,
    consultant: str | None,
    supplier: str | None,
    rows: list[dict],
) -> bytes:
    """The reply sheet as the PDF that goes back with the next revision.

    Laid out a page at a time -- as many comments as fit, then the next
    page -- so that a sheet of forty comments carries its column headings
    all the way through and no comment is cut in half by a page break.
    """
    import pymupdf

    rows = list(rows or [blank_row(1)])
    facts = (("Project", project_name), ("Ref No", f"{reference} - {revision}"),
             ("MANUFACTURER", manufacturer), ("Consultant", consultant))
    shared = {"heading": f"Reply to Consultant Comments on {system_title} submittal",
              "facts": facts, "supplier_label": supplier_name(supplier)}

    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (40, 40, -40, -40)

    def story_for(chunk: list[dict], start: int, with_heading: bool):
        return pymupdf.Story(html=_sheet_html(rows=chunk, first_number=start,
                                              with_heading=with_heading, **shared),
                             user_css=PAGE_CSS)

    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    remaining, number, with_heading = rows, 1, True
    while remaining:
        # The most comments that fit on this page, found by halving rather
        # than by guessing a row height: a comment can be a line or ten.
        low, high, best = 1, len(remaining), 0
        while low <= high:
            mid = (low + high) // 2
            if story_for(remaining[:mid], number, with_heading).place(where)[0]:
                high = mid - 1
            else:
                best, low = mid, mid + 1
        take = best or 1          # a comment too tall for one page still starts on one
        story = story_for(remaining[:take], number, with_heading)
        more = True
        while more:
            device = writer.begin_page(mediabox)
            more, _ = story.place(where)
            story.draw(device)
            writer.end_page()
        remaining, number, with_heading = remaining[take:], number + take, False
    writer.close()
    return buffer.getvalue()
