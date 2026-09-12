"""The material submittal register -> .xlsx.

A Register sheet (the submittals, as the page lists them) and a Materials
sheet (the BOQ's materials with the datasheet found for each), so the
covering letter and the package index can be built from one file.
"""

from collections.abc import Sequence
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas_design import MaterialItemOut, SubmittalOut

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEADER_FILL = PatternFill("solid", fgColor="E8EEF9")

STATUS_LABELS = {
    "not_submitted": "Not submitted",
    "under_review": "Under review",
    "approved": "Approved",
    "rejected": "Rejected",
}


def _head(ws: Worksheet, project, subtitle: str, stamp: str, columns: Sequence[tuple[str, int]]) -> None:
    ws["A1"] = f"EP-{project.ep_number}" + (f" — {project.project_name}" if project.project_name else "")
    ws["A1"].font = _TITLE
    ws["A2"] = subtitle
    ws["A2"].font = _BOLD
    ws["A3"] = stamp
    for col, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=5, column=col, value=title)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width


def submittal_register_workbook(
    project,
    submittals: Sequence[SubmittalOut],
    materials: Sequence[MaterialItemOut],
    *,
    exported_by: str,
    exported_at: datetime,
) -> bytes:
    stamp = f"Exported {exported_at:%d %b %Y %H:%M} UTC by {exported_by}"
    workbook = Workbook()

    register = workbook.active
    register.title = "Register"
    _head(
        register, project, "Material Submittals — Register", stamp,
        [("#", 5), ("Submittal title", 44), ("System", 10), ("Manufacturer", 22), ("Revision", 10),
         ("Status", 16), ("Materials", 11), ("With datasheet", 14), ("Last updated", 18), ("Note", 40)],
    )
    for i, item in enumerate(submittals, start=1):
        row = 5 + i
        values = [i, item.title, item.system_code, item.manufacturer, item.revision,
                  STATUS_LABELS.get(item.status, item.status), item.materials, item.materials_with_datasheet,
                  item.updated_at.strftime("%d %b %Y %H:%M"), item.note]
        for col, value in enumerate(values, start=1):
            register.cell(row=row, column=col, value=value)

    sheet = workbook.create_sheet("Materials")
    _head(
        sheet, project, "Material Submittals — Materials from the BOQ", stamp,
        [("System", 10), ("Part No.", 22), ("Description", 56), ("Qty", 9), ("Manufacturer", 20),
         ("Datasheet", 36), ("Document no.", 16)],
    )
    for i, item in enumerate(materials, start=1):
        row = 5 + i
        values = [item.system_code, item.part_no, item.description, item.quantity, item.manufacturer,
                  item.datasheet_filename or "Not in the library", item.document_no]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col, value=value)
            if col == 6 and not item.datasheet_filename:
                cell.font = Font(color="B45309")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
