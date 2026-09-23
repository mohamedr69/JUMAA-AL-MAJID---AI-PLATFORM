"""The company's record of every job and the engineer designing it.

A workbook in the archive, kept by the design manager outside the
platform: one row per system, so a job with a fire alarm and an emergency
light monitoring system is two rows under one EP number. **Read, never
written.** It is the manager's register, and the platform's own projects
are a handful of the jobs in it -- six of eleven hundred today -- so this
reports what the register says rather than trying to reconcile the two.

Read on demand and kept under the file's timestamp and size, the same way
a datasheet is (`datasheet_library`): the whole workbook parses in about
a fifth of a second, and the register is edited daily, so a copy in the
database would only be a staler answer to the same question.

A project's rows do not always agree, and that is information rather than
a fault: sixty-four jobs are split between two engineers and ninety-two
carry two statuses at once. Both are kept as lists -- a manager reviewing
the work is looking for exactly those.
"""

from __future__ import annotations

import os
import threading
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import get_settings
from app.services import brands

# The columns the register is read from, by position: its headers are
# stable, and a name that drifts ("SE/ E #") is worse to key on.
DRF_NO, EP_NUMBER, JOB_NO, SCOPE, PROJECT_NAME = 0, 1, 2, 3, 4
CLIENT, CONSULTANT, CONTRACTOR, STARTED, SYSTEM = 5, 6, 7, 8, 9
PRODUCT, ACTION_DATE, MATERIAL_STATUS, STATUS, ENGINEER = 10, 11, 12, 13, 14
LAST_COLUMN = 21
# Rows whose EP number is one of these say "no job number", not a job.
NOT_AN_EP = {"", "N/A", "NA", "-", "TBA", "TBC"}


# --- divisions ---------------------------------------------------------------
#
# Which division an engineer designs for, as the design manager gave it on
# 23 Sep 2026. The division is the **engineer's**, not the job's: the
# register's System column runs to eighty-odd spellings and a fire alarm
# engineer picks up the occasional CCTV job, so grouping by system put the
# same person in both tables. One engineer, one division.
#
# `resigned` is kept rather than dropped: their jobs are still in the
# register and still have to be accounted for, but a manager reviewing who
# is carrying what wants the people who are here.
FIRE_ALARM, ELV, UNASSIGNED = "fire_alarm", "elv", "unassigned"
DIVISION_NAMES = {
    FIRE_ALARM: "Fire Alarm Division",
    ELV: "ELV Division",
    UNASSIGNED: "Not assigned to a division",
}


@dataclass(frozen=True)
class Engineer:
    name: str
    division: str
    resigned: bool = False


ENGINEERS: tuple[Engineer, ...] = (
    Engineer("Ameer", FIRE_ALARM),
    Engineer("Siraj", FIRE_ALARM),
    Engineer("Kuthup", FIRE_ALARM),
    Engineer("Sameer", FIRE_ALARM),
    Engineer("Aref", FIRE_ALARM),
    Engineer("Asad", FIRE_ALARM),
    Engineer("Ramadan", FIRE_ALARM),
    Engineer("Sohail", ELV),
    Engineer("Aswar", ELV),
    Engineer("Sameer TP", ELV),
    # No longer with the company; their jobs remain in the register.
    Engineer("Wahab", ELV, resigned=True),
    Engineer("Yasser", FIRE_ALARM, resigned=True),
    Engineer("Saleem", FIRE_ALARM, resigned=True),
    Engineer("Junner", ELV, resigned=True),
    Engineer("Ashick", ELV, resigned=True),
    Engineer("Shameem", ELV, resigned=True),
    Engineer("Afrose", FIRE_ALARM, resigned=True),
    Engineer("Thasneef", FIRE_ALARM, resigned=True),
    Engineer("Junner / Wahab", ELV, resigned=True),
    Engineer("Skariah", ELV, resigned=True),
)
_BY_NAME = {e.name.casefold(): e for e in ENGINEERS}

# Not engineers. "Abudhabi Office" and "Alex" are an office and a log
# entry; the manager asked for them to be left out rather than shown with
# a job each.
NOT_AN_ENGINEER = {"abudhabi office", "alex"}

# The statuses worth a tile, biggest first. Every status is counted in the
# job total; these are the ones a manager reads off at a glance. They do
# not sum to the total -- "HO to Maintenance" and "Cancelled" are real and
# have no tile -- so the page says the total separately.
TILE_STATUSES = ("On Going", "Testing & Commissioning", "Completed", "On Hold")


def engineer_of(name: str) -> Engineer | None:
    """The engineer a name in the register means, or None when it is not
    one -- an office, a log entry, or a name nobody has placed yet.

    Matched without case, which folds the register's "wahab" into "Wahab":
    one person, entered two ways.
    """
    text = (name or "").strip()
    if not text or text.casefold() in NOT_AN_ENGINEER:
        return None
    return _BY_NAME.get(text.casefold())


class RegisterUnavailable(RuntimeError):
    """The register is not configured, or not reachable from here."""


@dataclass
class RegisterProject:
    """One job as the register holds it, its rows put together."""

    ep_number: str
    project_name: str | None
    client: str | None
    consultant: str | None
    contractor: str | None
    scope: str | None
    # A job split between two engineers is two names here, not a choice
    # between them: it is what a manager is reviewing the register for.
    engineers: list[str] = field(default_factory=list)
    statuses: list[str] = field(default_factory=list)
    systems: list[str] = field(default_factory=list)
    products: list[str] = field(default_factory=list)
    started: datetime | None = None
    # The newest action across the job's rows: what "latest" orders by.
    last_action: datetime | None = None
    rows: int = 0


def register_path() -> Path:
    """The workbook, or a refusal saying which of the two things is wrong."""
    settings = get_settings()
    configured = (settings.project_register or "").strip()
    if not configured:
        raise RegisterUnavailable("The project register is not configured (PROJECT_REGISTER is unset)")
    path = Path(os.path.expandvars(configured)).expanduser()
    if not path.is_absolute():
        if not settings.projects_root:
            raise RegisterUnavailable("The project archive is not configured (PROJECTS_ROOT is unset)")
        path = Path(settings.projects_root) / path
    if not path.is_file():
        raise RegisterUnavailable(f"The project register was not found at {path}")
    return path


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _when(value) -> datetime | None:
    return value if isinstance(value, datetime) else None


def _read(path: Path) -> list[RegisterProject]:
    settings = get_settings()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[settings.project_register_sheet]
        rows = list(sheet.iter_rows(min_row=2, max_col=LAST_COLUMN, values_only=True))
    finally:
        workbook.close()

    found: dict[str, RegisterProject] = {}
    for row in rows:
        ep = _text(row[EP_NUMBER])
        if ep.upper() in NOT_AN_EP:
            continue
        project = found.get(ep)
        if project is None:
            project = RegisterProject(
                ep_number=ep,
                project_name=_text(row[PROJECT_NAME]) or None,
                client=_text(row[CLIENT]) or None,
                consultant=_text(row[CONSULTANT]) or None,
                contractor=_text(row[CONTRACTOR]) or None,
                scope=_text(row[SCOPE]) or None,
            )
            found[ep] = project
        project.rows += 1
        for value, into in (
            (_text(row[ENGINEER]), project.engineers),
            (_text(row[STATUS]), project.statuses),
            (_text(row[SYSTEM]), project.systems),
            # The register spells a brand several ways ("Hikvision",
            # "HIKVISION"); one spelling here, as everywhere else.
            (brands.normalise(_text(row[PRODUCT])) or "", project.products),
        ):
            if value and value not in into:
                into.append(value)
        started, action = _when(row[STARTED]), _when(row[ACTION_DATE])
        if started and (project.started is None or started < project.started):
            project.started = started
        if action and (project.last_action is None or action > project.last_action):
            project.last_action = action
    return list(found.values())


_lock = threading.Lock()
# The parsed register under the file's timestamp and size, so an edit is
# picked up and an unchanged file is not read twice.
_cache: tuple[tuple[float, int], list[RegisterProject]] | None = None


def projects(*, newest_first: bool = True) -> list[RegisterProject]:
    """Every job in the register, newest action first.

    A job whose rows carry no date sorts last rather than first: an
    unparsed date is not the beginning of time, and the register holds one
    row dated 1966.
    """
    global _cache
    path = register_path()
    stat = path.stat()
    stamp = (stat.st_mtime, stat.st_size)
    with _lock:
        if _cache is None or _cache[0] != stamp:
            _cache = (stamp, _read(path))
        found = list(_cache[1])
    if newest_first:
        found.sort(key=lambda p: (p.last_action is not None, p.last_action or datetime.min), reverse=True)
    return found


@dataclass
class EngineerLoad:
    """One engineer's work."""

    name: str
    resigned: bool = False
    total: int = 0
    by_status: dict[str, int] = field(default_factory=dict)
    latest: RegisterProject | None = None


@dataclass
class Division:
    key: str
    label: str
    total: int = 0
    by_status: dict[str, int] = field(default_factory=dict)
    engineers: list[EngineerLoad] = field(default_factory=list)
    latest: RegisterProject | None = None
    # Of `total`, how many are carried by someone who has left. A manager
    # looking at a division's load needs to know how much of it has no one
    # on it any more.
    resigned_total: int = 0


def _bucket(project: RegisterProject) -> str:
    """The one status to count a job under. A job whose rows disagree is
    counted under the liveliest of them -- work in progress is what a
    manager is looking at -- and the disagreement itself is still shown on
    the job's own row."""
    for status in TILE_STATUSES:
        if status in project.statuses:
            return status
    return project.statuses[0] if project.statuses else "N/A"


def divisions() -> list[Division]:
    """The work by division, and within it by engineer, newest first.

    A division is the engineers in it, so a job counts once per engineer
    carrying it: the sixty-odd jobs shared by two engineers count for
    both, and if those two are in different divisions the job appears in
    both. A name the roster does not know is kept under "not assigned"
    rather than dropped -- a new engineer must be visible, not silent.
    """
    found: dict[str, Division] = {
        key: Division(key=key, label=DIVISION_NAMES[key]) for key in (FIRE_ALARM, ELV, UNASSIGNED)
    }
    loads: dict[tuple[str, str], EngineerLoad] = {}

    for project in projects():  # newest first, so the first seen is the latest
        status = _bucket(project)
        for raw in project.engineers:
            name = (raw or "").strip()
            if not name or name.casefold() in NOT_AN_ENGINEER:
                continue
            known = engineer_of(name)
            key = known.division if known else UNASSIGNED
            # "wahab" and "Wahab" are one person; the roster's spelling wins.
            label = known.name if known else name
            resigned = bool(known and known.resigned)

            division = found[key]
            division.total += 1
            division.by_status[status] = division.by_status.get(status, 0) + 1
            if resigned:
                division.resigned_total += 1
            if division.latest is None:
                division.latest = project

            load = loads.get((key, label))
            if load is None:
                load = loads[(key, label)] = EngineerLoad(name=label, resigned=resigned)
                division.engineers.append(load)
            load.total += 1
            load.by_status[status] = load.by_status.get(status, 0) + 1
            if load.latest is None:
                load.latest = project

    for division in found.values():
        # Present first, then by how much each is carrying.
        division.engineers.sort(key=lambda e: (e.resigned, -e.total, e.name))
    return [d for d in (found[FIRE_ALARM], found[ELV], found[UNASSIGNED]) if d.total]


def statuses() -> list[str]:
    """Every status the register holds, the ones with a tile first and the
    rest after, so a filter lists them in the order a manager reads them."""
    seen: set[str] = set()
    for project in projects(newest_first=False):
        seen.update(project.statuses)
    tiled = [s for s in TILE_STATUSES if s in seen]
    return tiled + sorted(seen - set(tiled))


def engineers() -> list[str]:
    """The engineers the register names, for filtering by one of them."""
    seen: set[str] = set()
    for project in projects(newest_first=False):
        seen.update(project.engineers)
    return sorted(seen)


def updated_at() -> datetime | None:
    """When the register was last edited, so the page can say how fresh
    what it is showing is."""
    try:
        return datetime.fromtimestamp(register_path().stat().st_mtime)
    except RegisterUnavailable:
        return None


def revision() -> str | None:
    """What the register is right now, as one short string.

    The file's timestamp and size -- the same pair the cache is kept
    under. An open page asks for this on a timer and reads the register
    again only when it changes: a `stat` costs nothing, where parsing two
    thousand rows on a timer would cost something every time for nothing.

    None when the register cannot be reached, which a caller must tell
    apart from "unchanged".
    """
    try:
        stat = register_path().stat()
    except RegisterUnavailable:
        return None
    return f"{stat.st_mtime:.6f}-{stat.st_size}"
