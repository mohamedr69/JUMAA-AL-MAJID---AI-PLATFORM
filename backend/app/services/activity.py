"""Each user's record on the platform: what they did, and what is theirs.

`record` is called by the routes that change something -- and by opening a
project -- after their own work has succeeded, so a refused or failed
request leaves no event. `account` reads a user's record back as one
document: the account itself, the projects they created, are assigned to or
opened, the submittals, BOQ revisions and compliance statements they made,
and their activity, newest first. `account_workbook` is the same as an Excel
file, one sheet per part.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session

from app.core.timeutils import utc_now
from app.models import (
    ActivityEvent,
    ComplianceAudit,
    ComplianceStatement,
    Project,
    ProjectBoqRevision,
    ProjectSubmittal,
    ProjectSubmittalEvent,
    User,
)

# Opening a project is recorded once per window, not per request: the page
# asks for the project on every section it shows, and React StrictMode sends
# each request twice in development.
OPEN_WINDOW = timedelta(minutes=30)


def project_label(project: Project) -> str:
    return f"EP-{project.ep_number}" + (f" {project.project_name}" if project.project_name else "")


def record(
    db: Session,
    user: User,
    action: str,
    summary: str,
    *,
    project: Project | None = None,
    entity_type: str | None = None,
    entity_id: int | None = None,
    detail: dict | None = None,
    commit: bool = True,
) -> ActivityEvent:
    """Add one event for `user`. With `commit=False` it goes out with the
    caller's own commit -- for a deletion, whose project is gone after it."""
    event = ActivityEvent(
        user_id=user.id,
        at=utc_now(),
        action=action,
        summary=summary,
        project_id=project.id if project is not None else None,
        project_label=project_label(project) if project is not None else None,
        entity_type=entity_type,
        entity_id=entity_id,
        detail=detail,
    )
    db.add(event)
    if commit:
        db.commit()
    return event


def record_open(db: Session, user: User, project: Project) -> None:
    recent = (
        db.query(ActivityEvent.id)
        .filter(
            ActivityEvent.user_id == user.id,
            ActivityEvent.project_id == project.id,
            ActivityEvent.action == "project.opened",
            ActivityEvent.at >= utc_now() - OPEN_WINDOW,
        )
        .first()
    )
    if recent is None:
        record(db, user, "project.opened", f"Opened {project_label(project)}", project=project,
               entity_type="project", entity_id=project.id)


def events(
    db: Session,
    user_id: int,
    *,
    limit: int = 100,
    offset: int = 0,
    project_id: int | None = None,
    action: str | None = None,
) -> tuple[list[ActivityEvent], int]:
    query = db.query(ActivityEvent).filter(ActivityEvent.user_id == user_id)
    if project_id is not None:
        query = query.filter(ActivityEvent.project_id == project_id)
    if action:
        # "boq" matches every boq.* action; "boq.saved" only that one.
        query = query.filter(or_(ActivityEvent.action == action, ActivityEvent.action.like(f"{action}.%")))
    total = query.count()
    rows = query.order_by(ActivityEvent.at.desc(), ActivityEvent.id.desc()).offset(offset).limit(limit).all()
    return rows, total


@dataclass
class AccountProject:
    id: int | None
    label: str
    ep_number: str | None
    project_name: str | None
    status: str | None
    created: bool = False
    assigned: bool = False
    opened_count: int = 0
    changes: int = 0
    last_activity_at: datetime | None = None
    deleted: bool = False


@dataclass
class Account:
    user: User
    counts: dict[str, int]
    projects: list[AccountProject]
    submittals: list[dict]
    boq_revisions: list[dict]
    compliance_statements: list[dict]
    activity: list[ActivityEvent]
    activity_total: int
    actions: list[str] = field(default_factory=list)


def account(db: Session, user: User, *, activity_limit: int = 100) -> Account:
    uid = user.id

    # --- projects: created, assigned, or touched in the activity log ---------
    by_id: dict[int, AccountProject] = {}

    def entry(project: Project) -> AccountProject:
        if project.id not in by_id:
            by_id[project.id] = AccountProject(
                id=project.id, label=project_label(project), ep_number=project.ep_number,
                project_name=project.project_name, status=project.status.value,
            )
        return by_id[project.id]

    owned = db.query(Project).filter(or_(Project.created_by_id == uid, Project.design_engineer_id == uid)).all()
    for project in owned:
        item = entry(project)
        item.created = project.created_by_id == uid
        item.assigned = project.design_engineer_id == uid
        item.last_activity_at = project.updated_at

    per_project = (
        db.query(
            ActivityEvent.project_id,
            func.max(ActivityEvent.project_label),
            func.sum(case((ActivityEvent.action == "project.opened", 1), else_=0)),
            func.sum(case((ActivityEvent.action == "project.opened", 0), else_=1)),
            func.max(ActivityEvent.at),
        )
        .filter(ActivityEvent.user_id == uid, ActivityEvent.project_id.isnot(None))
        .group_by(ActivityEvent.project_id)
        .all()
    )
    live = {p.id: p for p in db.query(Project).filter(Project.id.in_([row[0] for row in per_project])).all()}
    for project_id, label, opened, changes, last_at in per_project:
        if project_id in live:
            item = entry(live[project_id])
        else:
            item = by_id.setdefault(project_id, AccountProject(
                id=project_id, label=label or f"Project {project_id}", ep_number=None,
                project_name=None, status=None, deleted=True,
            ))
        item.opened_count = int(opened or 0)
        item.changes = int(changes or 0)
        if last_at and (item.last_activity_at is None or last_at > item.last_activity_at):
            item.last_activity_at = last_at
    projects = sorted(by_id.values(), key=lambda p: p.last_activity_at or datetime.min, reverse=True)

    # --- submittals: created by the user, or changed by them -----------------
    touched = {row[0] for row in db.query(ProjectSubmittalEvent.submittal_id).filter(ProjectSubmittalEvent.by_id == uid)}
    submittal_rows = (
        db.query(ProjectSubmittal)
        .filter(or_(ProjectSubmittal.created_by_id == uid, ProjectSubmittal.id.in_(list(touched))))
        .order_by(ProjectSubmittal.updated_at.desc())
        .all()
    )
    submittals = [
        {
            "id": s.id, "project_id": s.project_id, "project_label": project_label(s.project),
            "title": s.title, "reference": s.reference, "system_code": s.system_code,
            "manufacturer": s.manufacturer, "revision": s.revision, "status": s.status.value,
            "reply_code": s.reply_code, "created_by_user": s.created_by_id == uid, "updated_at": s.updated_at,
            "changes_by_user": sum(1 for e in s.events if e.by_id == uid),
        }
        for s in submittal_rows
    ]

    # --- BOQ revisions issued -------------------------------------------------
    boq_revisions = [
        {
            "project_id": r.project_id, "project_label": project_label(r.project), "number": r.number,
            "label": r.label, "note": r.note, "lines": len(r.items or []), "issued_at": r.issued_at,
        }
        for r in db.query(ProjectBoqRevision).filter(ProjectBoqRevision.issued_by_id == uid)
        .order_by(ProjectBoqRevision.issued_at.desc()).all()
    ]

    # --- compliance statements created or approved ---------------------------
    audit_counts = dict(
        db.query(ComplianceAudit.statement_id, func.count(ComplianceAudit.id))
        .filter(ComplianceAudit.user_id == uid)
        .group_by(ComplianceAudit.statement_id)
        .all()
    )
    statements = (
        db.query(ComplianceStatement, Project)
        .join(Project, Project.id == ComplianceStatement.project_id)
        .filter(or_(
            ComplianceStatement.created_by_id == uid,
            ComplianceStatement.approved_by_id == uid,
            ComplianceStatement.id.in_(list(audit_counts)),
        ))
        .order_by(ComplianceStatement.updated_at.desc())
        .all()
    )
    compliance_statements = [
        {
            "id": s.id, "project_id": p.id, "project_label": project_label(p), "kind": s.kind,
            "system_code": s.system_code, "clauses": len(s.rows or []),
            "created_by_user": s.created_by_id == uid, "approved_by_user": s.approved_by_id == uid,
            "approved_at": s.approved_at, "clause_changes_by_user": int(audit_counts.get(s.id, 0)),
            "updated_at": s.updated_at,
        }
        for s, p in statements
    ]

    action_counts = dict(
        db.query(ActivityEvent.action, func.count(ActivityEvent.id))
        .filter(ActivityEvent.user_id == uid)
        .group_by(ActivityEvent.action)
        .all()
    )
    counts = {
        "projects": len(projects),
        "projects_created": sum(1 for p in projects if p.created),
        "projects_assigned": sum(1 for p in projects if p.assigned),
        "projects_opened": sum(1 for p in projects if p.opened_count),
        "logins": action_counts.get("auth.login", 0),
        "boq_saves": action_counts.get("boq.saved", 0),
        "boq_revisions_issued": len(boq_revisions),
        "submittals": len(submittals),
        "submittals_created": sum(1 for s in submittals if s["created_by_user"]),
        "compliance_statements": len(compliance_statements),
        "compliance_approvals": sum(1 for s in compliance_statements if s["approved_by_user"]),
        "compliance_clause_changes": sum(audit_counts.values()),
        "activity_events": sum(action_counts.values()),
    }
    activity, total = events(db, uid, limit=activity_limit)
    return Account(
        user=user, counts=counts, projects=projects, submittals=submittals, boq_revisions=boq_revisions,
        compliance_statements=compliance_statements, activity=activity, activity_total=total,
        actions=sorted(action_counts),
    )


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(["" if value is None else value for value in row])
    for index, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[index - 1] or "")) for r in rows[:500]])
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(width + 2, 10), 60)
    sheet.freeze_panes = "A2"


# The workbook holds the newest events up to this many, so a long history
# cannot make one export hold the whole table in memory; the Account sheet
# says when the list was cut.
EXPORT_EVENT_LIMIT = 50_000


def account_workbook(db: Session, user: User) -> bytes:
    """The user's whole record as a workbook -- every activity event up to
    EXPORT_EVENT_LIMIT, not just the latest page of them."""
    data = account(db, user, activity_limit=1)
    all_events, total = events(db, user.id, limit=EXPORT_EVENT_LIMIT)
    workbook = Workbook()
    workbook.remove(workbook.active)

    _sheet(workbook, "Account", ["Field", "Value"], [
        ["Name", user.full_name], ["Email", user.email], ["Role", user.role.value.replace("_", " ").title()],
        ["Active", "Yes" if user.is_active else "No"], ["Account created", user.created_at],
        ["Last sign-in", user.last_login_at], ["Exported", utc_now()],
        *([["Activity sheet", f"newest {len(all_events):,} of {total:,} events"]] if total > len(all_events) else []),
        *[[name.replace("_", " ").capitalize(), value] for name, value in data.counts.items()],
    ])
    _sheet(workbook, "Projects",
           ["Project", "EP number", "Name", "Status", "Created by user", "Assigned to user", "Times opened",
            "Changes", "Last activity", "Deleted"],
           [[p.label, p.ep_number, p.project_name, p.status, "Yes" if p.created else "", "Yes" if p.assigned else "",
             p.opened_count, p.changes, p.last_activity_at, "Yes" if p.deleted else ""] for p in data.projects])
    _sheet(workbook, "Activity", ["When (UTC)", "Action", "Project", "Summary", "Detail"],
           [[e.at, e.action, e.project_label, e.summary, _detail_text(e.detail)] for e in all_events])
    _sheet(workbook, "Submittals",
           ["Project", "Title", "Reference", "System", "Manufacturer", "Revision", "Status", "Reply code",
            "Created by user", "Changes by user", "Updated"],
           [[s["project_label"], s["title"], s["reference"], s["system_code"], s["manufacturer"], s["revision"],
             s["status"].replace("_", " "), s["reply_code"], "Yes" if s["created_by_user"] else "",
             s["changes_by_user"], s["updated_at"]] for s in data.submittals])
    _sheet(workbook, "BOQ revisions", ["Project", "Revision", "Lines", "Note", "Issued (UTC)"],
           [[r["project_label"], r["label"], r["lines"], r["note"], r["issued_at"]] for r in data.boq_revisions])
    _sheet(workbook, "Compliance",
           ["Project", "System", "Kind", "Clauses", "Created by user", "Approved by user", "Approved (UTC)",
            "Clause changes by user", "Updated"],
           [[s["project_label"], s["system_code"], s["kind"], s["clauses"], "Yes" if s["created_by_user"] else "",
             "Yes" if s["approved_by_user"] else "", s["approved_at"], s["clause_changes_by_user"], s["updated_at"]]
            for s in data.compliance_statements])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _detail_text(detail: dict | None) -> str:
    if not detail:
        return ""
    return "; ".join(f"{key}: {value}" for key, value in detail.items())
