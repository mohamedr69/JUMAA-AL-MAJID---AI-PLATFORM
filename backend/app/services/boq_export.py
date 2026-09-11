"""Bill of Quantities -> .xlsx.

A Summary sheet (project details and per-system totals), then one sheet per
system in the same order as the BOQ page's tabs. Quantities that are numbers
are written as numbers so the engineer can work with them in Excel; the
sheets' own "Lot" and similar are kept as text and left out of the totals,
which say so.
"""

import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

UNASSIGNED = "Unassigned"

# (header, column width in characters)
COLUMNS: list[tuple[str, int]] = [
    ("#", 5),
    ("Group", 30),
    ("Manufacturer", 16),
    ("Model / Part No.", 26),
    ("Description", 60),
    ("Qty", 9),
    ("Unit", 8),
    ("Unit Price", 13),
    ("Total Price", 14),
    ("Remarks", 30),
]
QTY_COL = 6
TOTAL_PRICE_COL = 9
HEADER_ROW = 5

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEADER_FILL = PatternFill("solid", fgColor="E8EEF9")
_RULE = Border(top=Side(style="thin", color="999999"))
_PRICE_FORMAT = "#,##0.00"


class BoqLine(Protocol):
    system_code: str | None
    group_heading: str | None
    manufacturer: str | None
    catalog_no: str | None
    description: str
    quantity: str | None
    unit: str | None
    unit_price: Decimal | None
    total_price: Decimal | None
    remarks: str | None


class ProjectLike(Protocol):
    ep_number: str
    project_name: str | None
    client: str | None
    consultant: str | None
    contractor: str | None
    location: str | None
    scope_of_work: str | None


def parse_quantity(text: str | None) -> Decimal | None:
    """The number a quantity cell holds, or None for "Lot" and the like.
    Thousands separators and stray spaces are not a reason to refuse it."""
    cleaned = re.sub(r"[,\s]", "", text or "")
    return Decimal(cleaned) if re.fullmatch(r"\d+(\.\d+)?", cleaned) else None


@dataclass
class SystemTotals:
    system: str
    lines: int = 0
    units: Decimal = Decimal(0)
    # Lines quoted as a word rather than a number, by word: {"Lot": 2}.
    by_word: dict[str, int] = field(default_factory=dict)
    total_price: Decimal | None = None


def totals_for(system: str, lines: Iterable[BoqLine]) -> SystemTotals:
    totals = SystemTotals(system)
    for line in lines:
        totals.lines += 1
        quantity = parse_quantity(line.quantity)
        if quantity is not None:
            totals.units += quantity
        elif line.quantity and line.quantity.strip():
            word = line.quantity.strip()
            totals.by_word[word] = totals.by_word.get(word, 0) + 1
        if line.total_price is not None:
            totals.total_price = (totals.total_price or Decimal(0)) + line.total_price
    return totals


def group_by_system(
    lines: Sequence[BoqLine], sheet_codes: Sequence[str | None]
) -> list[tuple[str, list[BoqLine]]]:
    """Systems in the BOQ page's tab order: the project's Design Sheets first,
    then any other system a line carries, then lines with none."""
    order: list[str] = []
    for code in [*sheet_codes, *(line.system_code for line in lines)]:
        if code and code not in order:
            order.append(code)
    groups = [(code, [line for line in lines if line.system_code == code]) for code in order]
    unassigned = [line for line in lines if not line.system_code]
    if unassigned:
        groups.append((UNASSIGNED, unassigned))
    return [(system, members) for system, members in groups if members]


def _number(value: Decimal) -> int | Decimal:
    return int(value) if value == value.to_integral_value() else value


def _title_block(ws: Worksheet, project: ProjectLike, subtitle: str, stamp: str) -> None:
    ws["A1"] = f"EP-{project.ep_number}" + (f" — {project.project_name}" if project.project_name else "")
    ws["A1"].font = _TITLE
    ws["A2"] = subtitle
    ws["A2"].font = _BOLD
    ws["A3"] = stamp


def _write_system_sheet(
    ws: Worksheet, project: ProjectLike, system: str, lines: list[BoqLine], label: str, stamp: str
) -> None:
    _title_block(ws, project, f"{label} — {system}", stamp)

    for col, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

    row = HEADER_ROW
    for number, line in enumerate(lines, start=1):
        row += 1
        quantity = parse_quantity(line.quantity)
        values = [
            number,
            line.group_heading,
            line.manufacturer,
            line.catalog_no,
            line.description,
            _number(quantity) if quantity is not None else line.quantity,
            line.unit,
            line.unit_price,
            line.total_price,
            line.remarks,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col in (2, 5, 10))
        ws.cell(row=row, column=8).number_format = _PRICE_FORMAT
        ws.cell(row=row, column=9).number_format = _PRICE_FORMAT

    # Live formulas, so the totals follow any edit made in Excel afterwards.
    first, last = HEADER_ROW + 1, row
    total_row = row + 1
    qty_col = get_column_letter(QTY_COL)
    price_col = get_column_letter(TOTAL_PRICE_COL)
    ws.cell(row=total_row, column=5, value="Total")
    ws.cell(row=total_row, column=QTY_COL, value=f"=SUM({qty_col}{first}:{qty_col}{last})")
    price = ws.cell(row=total_row, column=TOTAL_PRICE_COL, value=f"=SUM({price_col}{first}:{price_col}{last})")
    price.number_format = _PRICE_FORMAT
    for col in range(1, len(COLUMNS) + 1):
        ws.cell(row=total_row, column=col).font = _BOLD
        ws.cell(row=total_row, column=col).border = _RULE

    words = totals_for(system, lines).by_word
    if words:
        listed = ", ".join(f"{count} × {word}" for word, count in words.items())
        ws.cell(row=total_row + 1, column=5, value=f"Not in the quantity total (quoted as a word): {listed}")


def _write_summary(
    ws: Worksheet, project: ProjectLike, groups: list[tuple[str, list[BoqLine]]], label: str, stamp: str
) -> None:
    _title_block(ws, project, label, stamp)

    details = [
        ("Client", project.client),
        ("Consultant", project.consultant),
        ("Contractor", project.contractor),
        ("Location", project.location),
        ("Scope of Work", project.scope_of_work),
    ]
    row = HEADER_ROW
    for name, value in details:
        ws.cell(row=row, column=1, value=name).font = _BOLD
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    headers = ["System", "Lines", "Total Qty", "Quoted as a word", "Total Price"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = _BOLD
        cell.fill = _HEADER_FILL
    for system, lines in groups:
        row += 1
        totals = totals_for(system, lines)
        ws.cell(row=row, column=1, value=system)
        ws.cell(row=row, column=2, value=totals.lines)
        ws.cell(row=row, column=3, value=_number(totals.units))
        ws.cell(
            row=row,
            column=4,
            value=", ".join(f"{count} × {word}" for word, count in totals.by_word.items()) or None,
        )
        price = ws.cell(row=row, column=5, value=totals.total_price)
        price.number_format = _PRICE_FORMAT
    if not groups:
        ws.cell(row=row + 1, column=1, value="The BOQ has no lines.")

    for col, width in enumerate([18, 40, 12, 24, 14], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _sheet_title(system: str) -> str:
    # Excel refuses these characters in a sheet name and caps it at 31.
    return re.sub(r"[\[\]:*?/\\]", "-", system)[:31] or UNASSIGNED


def boq_workbook(
    project: ProjectLike,
    lines: Sequence[BoqLine],
    sheet_codes: Sequence[str | None],
    *,
    exported_by: str,
    exported_at: datetime,
    revision_label: str | None = None,
) -> bytes:
    label = "Bill of Quantities" + (f" — {revision_label}" if revision_label else "")
    stamp = f"Exported {exported_at:%d %b %Y %H:%M} UTC by {exported_by}"
    groups = group_by_system(lines, sheet_codes)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, project, groups, label, stamp)
    for system, members in groups:
        _write_system_sheet(
            workbook.create_sheet(_sheet_title(system)), project, system, members, label, stamp
        )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
