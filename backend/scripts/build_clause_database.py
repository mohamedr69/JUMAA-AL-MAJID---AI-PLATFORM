"""Build the clause database: every past clause with its reply and remark.

Reads the Responses sheet of the exported compliance knowledge base
(`data base/Compliance Knowledge Base.xlsx`) and writes a SQLite file that
Python can query directly (`data base/compliance_clauses.db`):

    clauses    one row per distinct clause (matched on its normalized text)
    responses  one row per distinct reply + remark given to a clause, with
               how many times it was used and the canonical answer
               (Comply, Noted, Not applicable, ...)
    sources    where each response was found: EP number, project,
               statement file, system, SL.NO
    clause_responses   a view: clause, reply, remark, answer, times_used
    clause_fts         full-text index over the clauses, for near matches

A CSV of the view is written beside it for opening in Excel, and
`Compliance_Response_Database.xlsx` -- the same records shaped as the
workbook the platform's knowledge base imports. Point
COMPLIANCE_KNOWLEDGE_SOURCE at the `data base` folder and an administrator's
"Update knowledge base" loads it.

    cd backend
    .\\venv\\Scripts\\python scripts\\build_clause_database.py
    .\\venv\\Scripts\\python scripts\\build_clause_database.py --lookup "The panel shall be UL listed"

Rebuilding replaces the database file; the workbook is only read.
No model is called.
"""

from __future__ import annotations

import argparse
import csv
import re
import sqlite3
import sys
import warnings
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
warnings.filterwarnings("ignore")

from app.compliance.statements import canonical  # noqa: E402

REPO = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = REPO / "data base" / "Compliance Knowledge Base.xlsx"
DEFAULT_OUTPUT = REPO / "data base" / "compliance_clauses.db"
# The file the platform's "Update knowledge base" reads, with
# COMPLIANCE_KNOWLEDGE_SOURCE pointing at this folder.
DEFAULT_IMPORT_WORKBOOK = REPO / "data base" / "Compliance_Response_Database.xlsx"

SCHEMA = """
CREATE TABLE clauses (
    id          INTEGER PRIMARY KEY,
    clause      TEXT NOT NULL,          -- the wording as first seen
    clause_key  TEXT NOT NULL UNIQUE    -- lower case, punctuation and spacing collapsed
);
CREATE TABLE responses (
    id          INTEGER PRIMARY KEY,
    clause_id   INTEGER NOT NULL REFERENCES clauses(id),
    reply       TEXT NOT NULL,          -- as written, '' when only a remark was given
    remark      TEXT NOT NULL,          -- '' when there is none
    answer      TEXT,                   -- canonical: Comply, Noted, Not applicable, ...
    times_used  INTEGER NOT NULL,
    UNIQUE (clause_id, reply, remark)
);
CREATE TABLE sources (
    id              INTEGER PRIMARY KEY,
    response_id     INTEGER NOT NULL REFERENCES responses(id),
    ep_number       TEXT,
    project         TEXT,
    statement_file  TEXT,
    system          TEXT,
    sl_no           TEXT
);
CREATE INDEX ix_responses_clause ON responses(clause_id);
CREATE INDEX ix_sources_response ON sources(response_id);
CREATE INDEX ix_sources_ep ON sources(ep_number);
CREATE VIEW clause_responses AS
    SELECT c.id AS clause_id, c.clause, r.id AS response_id, r.reply, r.remark, r.answer, r.times_used
    FROM clauses c JOIN responses r ON r.clause_id = c.id;
CREATE VIRTUAL TABLE clause_fts USING fts5(clause, content='clauses', content_rowid='id');
"""


def clause_key(text: str) -> str:
    return re.sub(r"\W+", " ", text.lower()).strip()


def clean(value) -> str:
    return re.sub(r"[ \t]+", " ", str(value)).strip() if value is not None else ""


def build(workbook: Path, output: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(workbook, read_only=True)
    if "Responses" not in wb.sheetnames:
        sys.exit(f"{workbook.name} has no Responses sheet")
    rows = wb["Responses"].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    col = {name: header.index(name) for name in
           ("EP number", "Project folder", "Statement file", "System", "SL.NO", "Specification clause", "Compliance response", "Remarks")}

    clauses: dict[str, int] = {}
    clause_text: list[tuple[int, str, str]] = []
    responses: dict[tuple[int, str, str], list] = {}   # key -> [id, count]
    sources: list[tuple] = []
    skipped = 0
    for row in rows:
        clause, reply, remark = clean(row[col["Specification clause"]]), clean(row[col["Compliance response"]]), clean(row[col["Remarks"]])
        key = clause_key(clause)
        if not key or not (reply or remark):
            skipped += 1
            continue
        if key not in clauses:
            clauses[key] = len(clauses) + 1
            clause_text.append((clauses[key], clause, key))
        rkey = (clauses[key], reply, remark)
        if rkey not in responses:
            responses[rkey] = [len(responses) + 1, 0]
        responses[rkey][1] += 1
        sources.append((responses[rkey][0], clean(row[col["EP number"]]), clean(row[col["Project folder"]]),
                        clean(row[col["Statement file"]]), clean(row[col["System"]]), clean(row[col["SL.NO"]])))

    output.parent.mkdir(parents=True, exist_ok=True)
    tmp = output.with_suffix(".tmp")
    tmp.unlink(missing_ok=True)
    db = sqlite3.connect(tmp)
    db.executescript(SCHEMA)
    db.executemany("INSERT INTO clauses (id, clause, clause_key) VALUES (?, ?, ?)", clause_text)
    db.executemany("INSERT INTO responses (id, clause_id, reply, remark, answer, times_used) VALUES (?, ?, ?, ?, ?, ?)",
                   [(rid, cid, reply, remark, canonical(reply) or (canonical(remark) if not reply else None), n)
                    for (cid, reply, remark), (rid, n) in responses.items()])
    db.executemany("INSERT INTO sources (response_id, ep_number, project, statement_file, system, sl_no) VALUES (?, ?, ?, ?, ?, ?)", sources)
    db.execute("INSERT INTO clause_fts(clause_fts) VALUES ('rebuild')")
    db.commit()

    csv_path = output.with_suffix(".csv")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["clause_id", "clause", "reply", "remark", "answer", "times_used"])
        writer.writerows(db.execute("SELECT clause_id, clause, reply, remark, answer, times_used FROM clause_responses "
                                    "ORDER BY clause_id, times_used DESC"))
    db.close()
    output.unlink(missing_ok=True)
    tmp.rename(output)

    print(f"{len(clauses):,} clauses, {len(responses):,} distinct responses, {len(sources):,} source rows "
          f"({skipped:,} rows without a clause or a reply left out)")
    print(f"Database: {output}")
    print(f"CSV:      {csv_path}")


# --- the workbook the platform's knowledge base imports ----------------------------------------

# The statement's system names to the knowledge base's system labels (policy.SYSTEM_TO_KNOWLEDGE).
SYSTEM_LABELS = {
    "Fire Alarm": "FA", "Voice Evacuation": "FA", "Emergency Light Monitoring": "EML",
    "Central Battery System": "CBS", "PA/VA & BGM": "OTHER: PA / PAVA (non fire)",
}
# The answer a reply reads as (statements.canonical) to the knowledge base's historical status.
ANSWER_STATUS = {
    "Comply": "Comply", "Noted": "Noted", "Complied with remark": "Comply with Qualification",
    "Not applicable": "Not Applicable", "By others": "By Others", "Deviation": "Not Comply / Deviation",
}


def _labels(system: str) -> list[str]:
    return sorted({SYSTEM_LABELS[s.strip()] for s in system.split(",") if s.strip() in SYSTEM_LABELS})


def _manufacturer(texts: list[str]) -> str | None:
    """The one manufacturer a statement names, when one clearly dominates;
    None when it names none or several."""
    from collections import Counter

    from app.knowledge import policy

    aliases = sorted(policy._ALIAS_LOOKUP.items(), key=lambda kv: -len(kv[0]))
    counts: Counter = Counter()
    for text in texts:
        upper = text.upper()
        for alias, maker in aliases:
            if re.search(rf"(?<![A-Z0-9]){re.escape(alias)}(?![A-Z0-9])", upper):
                counts[maker] += 1
    top = counts.most_common(2)
    if not top or (len(top) > 1 and top[0][1] < 3 * top[1][1]):
        return None
    return top[0][0]


def write_import_workbook(workbook: Path, target: Path) -> None:
    """Shape the past statements as the Compliance Response Database the
    importer reads (app.knowledge.importer.SHEETS): one source per statement
    file, one requirement per clause wording, one response per reply +
    remark + manufacturer.

    What the statements do not record is left honest: the manufacturer is
    taken from the statement's own text only when one maker clearly
    dominates it, and every pairing is recorded as medium confidence --
    a reply is written in as a draft the engineer checks, never trusted
    silently."""
    import openpyxl
    from collections import defaultdict
    from datetime import date

    from app.compliance.statements import canonical
    from app.knowledge.importer import SHEETS
    from app.knowledge.normalize import requirement_hash

    wb = openpyxl.load_workbook(workbook, read_only=True)
    rows = wb["Statements"].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    modified = {(clean(r[header.index("EP number")]), clean(r[header.index("Statement file")])): clean(r[header.index("File modified")])
                for r in rows}
    rows = wb["Responses"].iter_rows(values_only=True)
    header = [clean(h) for h in next(rows)]
    records = [dict(zip(header, (clean(v) for v in r))) for r in rows]
    wb.close()

    # Sources: one per statement file.
    by_source: dict[tuple, list[dict]] = defaultdict(list)
    for r in records:
        if r["Specification clause"] and (r["Compliance response"] or r["Remarks"]):
            by_source[(r["EP number"], r["Project folder"], r["Statement file"], r["System"])].append(r)
    sources, source_of, maker_of = [], {}, {}
    for n, (key, items) in enumerate(sorted(by_source.items()), 1):
        ep, project, filename, system = key
        sid = f"S{n:05d}"
        maker = _manufacturer([project, filename] + [f"{i['Compliance response']} {i['Remarks']}" for i in items]
                              + [i["Specification clause"] for i in items])
        source_of[key], maker_of[sid] = sid, maker
        sources.append({"source_id": sid, "filename": filename, "onedrive_link_or_id": f"{project}\\{filename}", "duplicate_copies": None,
                        "project": project, "job_number": ep, "client": None, "system": ", ".join(_labels(system)) or None,
                        "manufacturer": maker, "brand": maker, "specification_family": None, "section_numbers": None,
                        "document_revision": None, "document_date": modified.get((ep, filename)),
                        "document_review_status": None, "extraction_status": "compliance statement table", "sha1": None})

    requirements: dict[str, dict] = {}
    mappings: dict[tuple, dict] = {}
    responses: dict[tuple, dict] = {}
    links: dict[tuple, dict] = {}
    for key, items in by_source.items():
        sid = source_of[key]
        labels = _labels(key[3])
        for r in items:
            digest = requirement_hash(r["Specification clause"])
            req = requirements.get(digest)
            if req is None:
                req = requirements[digest] = {"requirement_id": f"R{len(requirements) + 1:06d}", "system": set(), "subsystem": None,
                                              "topic": None, "canonical_requirement_text": r["Specification clause"],
                                              "technical_variant": None, "merge_review_status": None, "numbers": None, "n_sources": set()}
            req["system"].update(labels)
            req["n_sources"].add(sid)
            mkey = (req["requirement_id"], sid)
            if mkey not in mappings:
                mappings[mkey] = {"mapping_id": f"M{len(mappings) + 1:07d}", "requirement_id": req["requirement_id"],
                                  "specification_family": None, "specification_title": None, "specification_edition": None,
                                  "section_number": None, "clause_number": None, "clause_label_as_printed": r["SL.NO"][:32] or None,
                                  "heading_context": None, "original_requirement_text": r["Specification clause"], "source_id": sid,
                                  "pdf_page": None, "printed_page": None, "extraction_method": "A_table", "pairing_confidence": "medium"}
            reply, remark = r["Compliance response"], r["Remarks"]
            answer = canonical(reply) or (None if reply else canonical(remark))
            rkey = (req["requirement_id"], reply, remark, maker_of[sid])
            resp = responses.get(rkey)
            if resp is None:
                resp = responses[rkey] = {"response_id": f"A{len(responses) + 1:07d}", "requirement_id": req["requirement_id"],
                                          "system": set(), "manufacturer": maker_of[sid], "brand": maker_of[sid],
                                          "model_configuration": None, "exact_historical_response": reply, "remarks_column": remark or None,
                                          "original_compliance_status": ANSWER_STATUS.get(answer, "Other / Free text"),
                                          "scope_conditions": None, "responsible_party": None, "cited_supporting_references": None,
                                          "review_flag": None, "specification_families": None, "n_sources": set()}
            resp["system"].update(labels)
            resp["n_sources"].add(sid)
            links.setdefault((resp["response_id"], sid), {"response_id": resp["response_id"], "mapping_id": mappings[mkey]["mapping_id"],
                                                          "source_id": sid, "pdf_page": None, "historical_review_status": None,
                                                          "consultant_comment": None, "superseded_status": None})

    for table in (requirements, responses):
        for row in table.values():
            row["system"] = ", ".join(sorted(row["system"])) or None
            row["n_sources"] = len(row["n_sources"])

    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.stem + ".tmp.xlsx")
    out = openpyxl.Workbook(write_only=True)
    data = {"SOURCES": sources, "REQUIREMENTS": list(requirements.values()), "SPECIFICATION_MAPPINGS": list(mappings.values()),
            "RESPONSES": list(responses.values()), "RESPONSE_SOURCES": list(links.values()), "MERGE_REVIEW": [], "REVIEW_ISSUES": []}
    for sheet, columns in SHEETS.items():
        ws = out.create_sheet(sheet)
        ws.append(list(columns))
        for row in data[sheet]:
            ws.append([row.get(c) for c in columns])
    readme = out.create_sheet("README")
    readme.append([f"Built {date.today():%d %b %Y} from past statements"])
    readme.append([f"Source: {workbook.name}, Responses sheet. Manufacturer inferred per statement; pairings recorded as medium confidence."])
    out.save(tmp)
    target.unlink(missing_ok=True)
    tmp.rename(target)

    known = sum(1 for s in sources if s["manufacturer"])
    print(f"Import workbook: {target}")
    print(f"  {len(sources):,} statements ({known:,} with a manufacturer), {len(requirements):,} requirements, "
          f"{len(responses):,} responses, {len(links):,} response-source links")


def lookup(output: Path, text: str, limit: int = 5) -> None:
    """The replies given to this clause before: exact wording first, else the nearest clauses."""
    db = sqlite3.connect(output)
    exact = db.execute("SELECT id, clause FROM clauses WHERE clause_key = ?", (clause_key(text),)).fetchall()
    if exact:
        matches = exact
    else:
        words = " OR ".join(f'"{w}"' for w in clause_key(text).split() if len(w) > 2)
        matches = db.execute("SELECT rowid, clause FROM clause_fts WHERE clause_fts MATCH ? ORDER BY rank LIMIT ?",
                             (words, limit)).fetchall() if words else []
    if not matches:
        print("No past clause found.")
    for cid, clause in matches:
        print(("EXACT: " if exact else "NEAR:  ") + clause[:300])
        for reply, remark, answer, n in db.execute(
                "SELECT reply, remark, answer, times_used FROM responses WHERE clause_id = ? ORDER BY times_used DESC", (cid,)):
            print(f"    [{n}x] reply: {reply or '-'} | answer: {answer or '-'}" + (f" | remark: {remark[:200]}" if remark else ""))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--import-workbook", type=Path, default=DEFAULT_IMPORT_WORKBOOK,
                        help="where to write the workbook the platform's knowledge base imports")
    parser.add_argument("--lookup", metavar="CLAUSE", help="query the built database instead of building it")
    args = parser.parse_args()
    if args.lookup:
        lookup(args.output, args.lookup)
    else:
        build(args.workbook, args.output)
        write_import_workbook(args.workbook, args.import_workbook)


if __name__ == "__main__":
    main()
