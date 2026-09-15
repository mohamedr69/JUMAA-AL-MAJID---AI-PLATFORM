"""Reading engineers' amplifier workbooks into a zone schedule.

The synthetic workbooks below each copy one layout found in the archive
(named per test); the live tests, gated like the extractor ones, read the
real files and pin them against the fixtures generated from them.
"""

import json
import os
from pathlib import Path

import openpyxl
import pytest

from app.schemas_design import VoiceEvacuationDesign
from app.services.ve_calculation import calculate
from app.services.ve_workbook_reader import (
    WorkbookReadError,
    parse_watts,
    read_amplifier_workbook,
    tap_from_heading,
)

FIXTURES = Path(__file__).parent / "fixtures"
LIVE_ROOT = os.environ.get("EP_PLATFORM_LIVE_ARCHIVE_ROOT")
requires_live_archive = pytest.mark.skipif(
    not LIVE_ROOT, reason="Set EP_PLATFORM_LIVE_ARCHIVE_ROOT to run against the real archive"
)


def _workbook(tmp_path: Path, cells: dict, merges: tuple = (), sheets: tuple = ()) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calc"
    for coordinate, value in cells.items():
        ws[coordinate] = value
    for merged in merges:
        ws.merge_cells(merged)
    for name in sheets:
        wb.create_sheet(name)
    path = tmp_path / "amp.xlsx"
    wb.save(path)
    return path


def _calculated(read):
    return calculate(
        VoiceEvacuationDesign(
            speaker_types=read.speaker_types,
            zones=read.zones,
            channels=read.channels,
            racks=read.racks,
            max_load_fraction=0.8,
        )
    )


@pytest.mark.parametrize(
    ("value", "watts"),
    [
        (50, 50),
        ("50 W", 50),
        ("50W", 50),
        ("250WATTS", 250),
        ("1 No 250 Watt Amplifier Upto 11th Floor", 250),
        ("1 No 50 watt Amplifier proposed", 50),
        ("2 x 50 W ", 100),
        ("SIGA-AA50 AMPLIFIER", None),
        (None, None),
    ],
)
def test_parse_watts(value, watts):
    assert parse_watts(value) == watts


@pytest.mark.parametrize(
    ("heading", "tap"),
    [
        ("Ceiling Speaker Qty  (1/2 Watt)", 0.5),
        ("Wall Speaker (W/P) Qty\n(1 Watt)", 1),
        ("CEILING MOUNTED SPEAKER\n (Corridor-0.75W)", 0.75),
        ("PARKING SPEAKER \n(Drive way-2W)", 2),
        ("WS-(.5W)", 0.5),
        ("WSF(W/P)-(2W)", 2),
        ("Ceiling Speaker", None),
        ("Wall Speaker flasher w/p", None),
    ],
)
def test_tap_from_heading(heading, tap):
    assert tap_from_heading(heading) == tap


def test_tapping_wattage_row_layout(tmp_path):
    """EP-24601 / EP-25176 / P09: headings, an optional row of models, a
    "Tapping Wattage" row, required watts / amplifier / rack merged down."""
    path = _workbook(
        tmp_path,
        {
            "A3": "TOWER AMPLIFIER LOAD CALCULATION",
            "C4": "Ceiling Speaker", "D4": "Wall Speaker (Stairs)", "E4": "Watts per Area",
            "F4": "Required Watts ", "G4": "Proposed Amplifier ", "H4": "RACK NAME",
            "C5": "EST-S186", "D5": "G4SRN",
            "A6": "Tapping Wattage", "C6": 0.75, "D6": 0.5,
            "A7": "Staircase -1", "B7": "CC2A", "D7": 25, "E7": 12.5, "F7": 20, "G7": 50, "H7": "APS-1",
            "A8": "Staircase -2", "B8": "CC2A", "D8": 15, "E8": 7.5,
            "A9": "Ground Floor", "B9": "CC2A", "C9": 12, "E9": 9, "F9": 39, "G9": "50 W",
            "B10": "CC2A", "C10": 40, "E10": 30,
            "A11": "Level 1", "B11": "CC2A", "C11": 4, "E11": 3, "F11": 3, "G11": 50, "H11": "APS-2",
            "A12": "Level 2", "E12": 0,
            "A14": "Result",
        },
        merges=("A9:A10", "F7:F8", "G7:G8", "F9:F10", "G9:G10", "H7:H10", "F11:F12", "G11:G12", "H11:H12"),
    )
    read = read_amplifier_workbook(path)
    assert [(s.name, s.model, s.tap_watts) for s in read.speaker_types] == [
        ("Ceiling Speaker", "EST-S186", 0.75),
        ("Wall Speaker (Stairs)", "G4SRN", 0.5),
    ]
    # The merged "Ground Floor" names both its rows; the "Result" note under
    # a blank row ends the schedule.
    assert [z.name for z in read.zones] == [
        "Staircase -1", "Staircase -2", "Ground Floor", "Ground Floor (2)", "Level 1", "Level 2",
    ]
    assert [(c.first_zone, c.amplifier_watts, c.sheet_required_watts) for c in read.channels] == [
        (0, 50, 20), (2, 50, 39), (4, 50, 3),
    ]
    assert [(r.name, r.first_channel) for r in read.racks] == [("APS-1", 0), ("APS-2", 2)]
    assert read.warnings == []

    result = _calculated(read)
    assert [c.required_watts for c in result.channels] == [20, 39, 3]
    assert [c.status for c in result.channels] == ["ok", "ok", "ok"]
    assert result.sheet_mismatches == 0


def test_uncalculated_formulas_are_refused_not_read_as_zero(tmp_path):
    """A workbook saved by something that writes formulas without computing
    them has no cached results: the counts would read as blank -- zero
    speakers -- and the amplifiers would look under-loaded."""
    cells = {
        "C4": "Ceiling Speaker", "D4": "Wall Speaker", "E4": "Watts per Area", "F4": "Required Watts", "G4": "Proposed Amplifier",
        "A6": "Tapping Wattage", "C6": 0.5, "D6": 1,
        "A7": "Level 1", "C7": 10, "D7": 2, "E7": 7, "F7": 7, "G7": 50,
        "A8": "Level 2", "C8": "=C7*2", "D8": 2, "E8": "=C8*C6+D8*D6",
    }
    with pytest.raises(WorkbookReadError, match="FORMULA_NOT_CALCULATED.*C8"):
        read_amplifier_workbook(_workbook(tmp_path, cells))

    # Only the sheet's own watts uncalculated: read, with the comparison skipped and said so.
    cells["C8"] = 20
    read = read_amplifier_workbook(_workbook(tmp_path, cells))
    assert [z.counts for z in read.zones][1] == {"ceiling_speaker": 20, "wall_speaker": 2}
    assert any("FORMULA_NOT_CALCULATED" in w for w in read.warnings)


def test_tap_in_heading_layout_with_totals_rows(tmp_path):
    """EP-15792 / EP-17428 / EP-13705: taps written into the headings, a
    label column on the left, "2 x 50 W" amplifiers and totals underneath."""
    path = _workbook(
        tmp_path,
        {
            "A3": "DEVICE", "B3": "Floor", "C3": "Ceiling Speaker Qty (1/2 Watt)",
            "D3": "Wall Speaker (W/P) Qty\n(1 Watt)", "E3": "Load in watt /Floor",
            "F3": "Total Wattage", "G3": "Total Wattage with 20 % Spare", "H3": "Proposed Amplifier",
            "A4": "AMPLIFIER-1", "B4": "Basement", "C4": 10, "D4": 8, "E4": 13, "F4": 30, "G4": 36, "H4": "2 x 50 W",
            "B5": "Ground Floor", "C5": 20, "D5": 7, "E5": 17,
            "A6": "AMPLIFIER-2", "B6": "Roof", "D6": 4, "E6": 4, "F6": 4, "G6": 4.8, "H6": "1 No 50 watt Amplifier",
            "B7": "Grand Total wattage", "C7": 30, "D7": 19, "E7": 34,
        },
        merges=("A4:A5", "F4:F5", "G4:G5", "H4:H5"),
    )
    read = read_amplifier_workbook(path)
    assert [s.tap_watts for s in read.speaker_types] == [0.5, 1]
    assert [z.name for z in read.zones] == ["Basement", "Ground Floor", "Roof"]
    assert [(c.label, c.amplifier_watts, c.sheet_required_watts) for c in read.channels] == [
        ("AMPLIFIER-1", 100, 30),
        ("AMPLIFIER-2", 50, 4),
    ]
    result = _calculated(read)
    assert [c.required_watts for c in result.channels] == [30, 4]


def test_headings_with_no_watts_per_area_column(tmp_path):
    """EP-23315: abbreviated speaker headings carrying their taps, and only
    the per-amplifier total ("AMP W") -- which here disagrees with the
    counts, and that is what the engineer needs to see."""
    path = _workbook(
        tmp_path,
        {
            "A1": "FLOOR", "B1": "WS-(.5W)", "C1": "WSF(W/P)-(2W)", "D1": "AMP W", "E1": "AMP",
            "A2": "ST-1", "B2": 5, "D2": 2.5, "E2": "AMP-1 LOC.GF",
            "A3": "B1", "B3": 16, "C3": 6, "D3": 20, "E3": "AMP-2 LOC.GF",
            "A4": "GF", "B4": 22,
            "A5": "TOTAL SPKR", "B5": 43, "C5": 6,
        },
        merges=("D3:D4", "E3:E4"),
    )
    read = read_amplifier_workbook(path)
    assert [z.name for z in read.zones] == ["ST-1", "B1", "GF"]
    assert [(c.first_zone, c.label) for c in read.channels] == [(0, "AMP-1 LOC.GF"), (1, "AMP-2 LOC.GF")]
    result = _calculated(read)
    assert [(c.required_watts, c.sheet_required_watts, c.sheet_mismatch) for c in result.channels] == [
        (2.5, 2.5, False),
        (31, 20, True),
    ]
    assert {c.status for c in result.channels} == {"no_rating"}


def test_headerless_layout(tmp_path):
    """EP-29495: speaker models over their taps, nothing labelled; channels
    and racks are merged text columns right of the watts."""
    path = _workbook(
        tmp_path,
        {
            "C1": "APS ",
            "C2": "EST-S186C", "D2": "757-3A-SS70",
            "C3": 0.75, "D3": 1,
            "B4": "UG", "E4": 0, "F4": "AMP-01", "G4": "APS-02", "H4": "GF FCC ROOM",
            "B5": "BF", "C5": 3, "D5": 16, "E5": 18.25,
            "B6": "GF", "C6": 5, "D6": 10, "E6": 13.75, "F6": "AMP-02",
            "B7": 1, "C7": 23, "D7": 5, "E7": 22.25, "F7": "AMP-01", "G7": "APS-03", "H7": "1ST",
        },
        merges=("F4:F5", "G4:G6", "H4:H6"),
    )
    read = read_amplifier_workbook(path)
    assert [(s.name, s.tap_watts) for s in read.speaker_types] == [("EST-S186C", 0.75), ("757-3A-SS70", 1)]
    assert [z.name for z in read.zones] == ["UG", "BF", "GF", "1"]
    assert [(c.first_zone, c.label) for c in read.channels] == [(0, "AMP-01"), (2, "AMP-02"), (3, "AMP-01")]
    assert [(r.name, r.location, r.first_channel) for r in read.racks] == [
        ("APS-02", "GF FCC ROOM", 0),
        ("APS-03", "1ST", 2),
    ]
    assert _calculated(read).sheet_mismatches == 0


def test_the_first_sheet_that_reads_is_used(tmp_path):
    path = _workbook(
        tmp_path,
        {"A1": "Notes only"},
        sheets=(),
    )
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("TW-A")
    for coordinate, value in {
        "A1": "Floor", "B1": "Ceiling Speaker (0.75W)", "C1": "Watts per Floor",
        "A2": "GF", "B2": 4, "C2": 3,
    }.items():
        ws[coordinate] = value
    wb.save(path)

    read = read_amplifier_workbook(path)
    assert (read.sheet, read.sheets) == ("TW-A", ["Calc", "TW-A"])
    with pytest.raises(WorkbookReadError, match="Calc"):
        read_amplifier_workbook(path, sheet="Calc")
    with pytest.raises(WorkbookReadError, match="no sheet named"):
        read_amplifier_workbook(path, sheet="Nope")


def test_counted_speakers_without_a_tap_are_refused(tmp_path):
    """Guessing the tap would put a made-up number into every total below."""
    path = _workbook(
        tmp_path,
        {
            "B1": "Ceiling Speaker", "C1": "Horn Speaker", "D1": "Watts per Area",
            "A2": "Tapping Wattage", "B2": 0.75,
            "A3": "GF", "B3": 4, "C3": 2, "D3": 5,
        },
    )
    with pytest.raises(WorkbookReadError, match="Horn Speaker"):
        read_amplifier_workbook(path)


def test_an_unrelated_workbook_is_refused(tmp_path):
    path = _workbook(tmp_path, {"A1": "Loop", "B1": "Devices", "A2": "L1", "B2": 120})
    with pytest.raises(WorkbookReadError, match="No sheet reads"):
        read_amplifier_workbook(path)


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path):
    path = tmp_path / "amp.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(WorkbookReadError, match="Could not open"):
        read_amplifier_workbook(path)


# --- the real archive --------------------------------------------------------


def _fixture(name: str) -> dict:
    return json.loads((FIXTURES / name).read_text())


@requires_live_archive
@pytest.mark.parametrize("fixture", ["ve_ep24601_tower_a.json", "ve_ep29495.json"])
def test_live_workbooks_read_as_their_fixtures(fixture):
    expected = _fixture(fixture)
    read = read_amplifier_workbook(Path(LIVE_ROOT) / expected["source"]["archive_path"], expected["source"]["sheet"])

    assert [s.tap_watts for s in read.speaker_types] == [s["tap_watts"] for s in expected["speaker_types"]]
    keys = [s.key for s in read.speaker_types]
    fixture_keys = [s["key"] for s in expected["speaker_types"]]
    assert len(read.zones) == len(expected["zones"])
    for zone, want in zip(read.zones, expected["zones"]):
        assert zone.sheet_watts == want["watts"]
        counts = {fixture_keys[keys.index(k)]: n for k, n in zone.counts.items()}
        assert {k: n for k, n in counts.items() if n} == {k: n for k, n in want["counts"].items() if n}

    assert [c.first_zone for c in read.channels] == [c["first_zone"] for c in expected["engineer_channels"]]
    for channel, want in zip(read.channels, expected["engineer_channels"]):
        assert channel.label == want.get("label")
        assert channel.amplifier_watts == want.get("amplifier_watts")
        assert channel.sheet_required_watts == want.get("required_watts")
    assert [(r.name, r.first_channel) for r in read.racks] == [
        (r["name"], r["first_channel"]) for r in expected["engineer_aps"]
    ]
    assert read.warnings == []


@requires_live_archive
def test_live_ep23315_amp2_formula_misses_two_floors():
    read = read_amplifier_workbook(
        Path(LIVE_ROOT)
        / "Aceptech Elect. & Mech. Contr. LLC/EP-23315 Hotel 2B+G+6 Typ. on Plot 6757715"
        / "Shop Dwg/FIRE ALARM SYSTEM/REV-01/AMP & BPS-.xlsx"
    )
    channels = _calculated(read).channels
    assert [(c.label, c.required_watts, c.sheet_required_watts) for c in channels] == [
        ("AMP-1 LOC.GF", 26, 26),
        ("AMP-2 LOC.GF", 53, 38.5),
        ("AMP-3 LOC.L3", 28, 28),
    ]


@requires_live_archive
def test_live_ep15792_basement_row_disagrees_with_its_counts():
    read = read_amplifier_workbook(
        Path(LIVE_ROOT) / "Al Arabia EMW/EP-15792 - G+4 and 40 Villas @ JVC/Amplifier Calculation -LAVIE.xlsx"
    )
    mismatches = [(z.name, z.watts, z.sheet_watts) for z in _calculated(read).zones if z.sheet_mismatch]
    assert mismatches == [("BASEMENT FLOOR", 83.5, 81)]


@requires_live_archive
def test_live_workbook_without_taps_is_refused():
    """EP-16884 lists a speaker total per circuit with no tap per type."""
    with pytest.raises(WorkbookReadError):
        read_amplifier_workbook(
            Path(LIVE_ROOT)
            / "AG Engineering/EP-16884 IMPZ District Cooling Plant/Shop Drawings/PAVA/REV-01/Amplifier Calcualtion.xlsx"
        )
