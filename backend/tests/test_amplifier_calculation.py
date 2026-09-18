"""Voice evacuation amplifier loading, worked out from the floor-wise BOQ.

The arithmetic is simple; the decisions around it are what these hold to:
a floor is never split between amplifiers, the limit is a fraction of the
rating rather than the rating, and only speakers load an amplifier.
"""

import openpyxl

from app.core.config import get_settings
from app.services import amplifier_calculation
from app.services.amplifier_calculation import AMPLIFIER_PART, calculate

from .conftest import login

settings = get_settings()

# The speaker database, as the design rules hold it.
TAPS = {
    "EST-S186C": {"part_no": "EST-S186C", "description": "Ceiling speaker",
                  "taps": [0.25, 0.5, 1, 2], "default_tap": 1},
    "EST-S1814": {"part_no": "EST-S1814", "description": "Wall speaker",
                  "taps": [0.5, 1, 2, 4], "default_tap": 2},
}


def _schedule(floors: list[str], lines: list[tuple[str, str, str | None, dict[str, int]]]) -> dict:
    """A floor-wise BOQ as the schedule reader produces one: (description,
    device, settled part number, {floor: count})."""
    return {
        "floors": floors,
        "items": [
            {"description": description, "device": device, "system": "FAS",
             "material": {"part_no": part, "description": "", "manufacturer": None} if part else None,
             "per_floor": dict(counts), "total": sum(counts.values())}
            for description, device, part, counts in lines
        ],
    }


# --- what loads an amplifier ------------------------------------------------------------------


def test_a_floors_load_is_its_speakers_times_their_tapping():
    schedule = _schedule(
        ["GF", "Level 1"],
        [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 10, "Level 1": 14}),
         ("Wall Speaker", "Speaker", "EST-S1814", {"GF": 2})],
    )
    result = calculate(schedule, taps=TAPS).as_dict()

    # Sorted by part number, so "EST-S1814" comes before "EST-S186C".
    assert [column["key"] for column in result["columns"]] == ["EST-S1814", "EST-S186C"]
    ground, first = result["floors"]
    assert ground["counts"] == {"EST-S186C": 10, "EST-S1814": 2}
    assert ground["watts"] == 10 * 1 + 2 * 2          # 14 W
    assert first["watts"] == 14 * 1
    assert result["total_watts"] == 28 and result["total_speakers"] == 26


def test_only_speakers_load_an_amplifier():
    """A sounder, a horn or a flasher is on a notification circuit."""
    schedule = _schedule(
        ["GF"],
        [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 10}),
         ("Wall Sounder", "Sounder", None, {"GF": 40}),
         ("Wall Sounder Flasher WP", "Sounder", None, {"GF": 40}),
         ("Smoke Detector", "Smoke detector", None, {"GF": 90})],
    )
    result = calculate(schedule, taps=TAPS).as_dict()
    assert [column["key"] for column in result["columns"]] == ["EST-S186C"]
    assert result["floors"][0]["watts"] == 10


def test_a_speaker_with_no_tapping_set_adds_nothing_and_says_so():
    schedule = _schedule(["GF"], [("Ceiling Speaker", "Speaker", "SPK-UNKNOWN", {"GF": 10})])
    result = calculate(schedule, taps=TAPS).as_dict()
    assert result["columns"][0]["tap"] is None
    assert result["floors"][0]["watts"] == 0
    assert any("no tapping set" in warning for warning in result["warnings"])


def test_a_line_not_settled_as_a_part_is_named_by_the_schedule_and_flagged():
    schedule = _schedule(["GF"], [("Ceiling Speaker", "Speaker", None, {"GF": 10})])
    result = calculate(schedule, taps=TAPS).as_dict()
    assert result["columns"][0]["key"] == "Ceiling Speaker"
    assert result["columns"][0]["unsettled"] is True
    assert any("not settled as a part" in warning for warning in result["warnings"])


# --- filling the amplifiers --------------------------------------------------------------------


def test_floors_are_combined_until_the_next_would_pass_the_limit():
    """At the 0.8 rule a SIGA-AA50 takes 40 W. Three 15 W floors are two
    amplifiers, not one: 15 + 15 fits, the third would make 45."""
    floors = ["GF", "Level 1", "Level 2"]
    schedule = _schedule(floors, [("Ceiling Speaker", "Speaker", "EST-S186C",
                                   {floor: 15 for floor in floors})])
    result = calculate(schedule, taps=TAPS, fraction=0.8).as_dict()

    assert result["limit_watts"] == 40
    assert [(a["name"], a["floors"], a["watts"]) for a in result["amplifiers"]] == [
        (f"AA50-1", ["GF", "Level 1"], 30.0),
        (f"AA50-2", ["Level 2"], 15.0),
    ]
    assert [floor["amplifier"] for floor in result["floors"]] == ["AA50-1", "AA50-1", "AA50-2"]


def test_a_floor_that_exactly_reaches_the_limit_still_fits():
    floors = ["GF", "Level 1"]
    schedule = _schedule(floors, [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 20, "Level 1": 20})])
    result = calculate(schedule, taps=TAPS).as_dict()
    assert len(result["amplifiers"]) == 1
    assert result["amplifiers"][0]["watts"] == 40


def test_a_floor_with_no_speakers_takes_no_amplifier():
    schedule = _schedule(["UG", "GF"], [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 10})])
    result = calculate(schedule, taps=TAPS).as_dict()
    assert result["floors"][0]["watts"] == 0 and result["floors"][0]["amplifier"] is None
    assert len(result["amplifiers"]) == 1


def test_a_floor_bigger_than_one_amplifier_is_reported_not_split():
    """Splitting a floor across two amplifiers is a decision an engineer
    makes, not something to do quietly."""
    schedule = _schedule(["GF", "Level 1"],
                         [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 55, "Level 1": 10})])
    result = calculate(schedule, taps=TAPS).as_dict()

    big = result["amplifiers"][0]
    assert big["floors"] == ["GF"] and big["watts"] == 55 and big["over_limit"] is True
    # ...and nothing else is put on it.
    assert result["amplifiers"][1]["floors"] == ["Level 1"]
    assert any("more speakers than one" in warning for warning in result["warnings"])


# --- cabinets ------------------------------------------------------------------------------------


def test_two_amplifiers_make_a_cabinet_and_an_odd_one_still_needs_its_own():
    floors = [f"Level {n}" for n in range(1, 6)]
    schedule = _schedule(floors, [("Ceiling Speaker", "Speaker", "EST-S186C",
                                   {floor: 40 for floor in floors})])
    result = calculate(schedule, taps=TAPS).as_dict()

    assert len(result["amplifiers"]) == 5
    assert [(c["name"], c["amplifiers"]) for c in result["cabinets"]] == [
        ("APS-1", ["AA50-1", "AA50-2"]),
        ("APS-2", ["AA50-3", "AA50-4"]),
        ("APS-3", ["AA50-5"]),
    ]
    assert result["amplifiers"][0]["cabinet"] == "APS-1"
    assert amplifier_calculation.cabinets_for(5) == 3
    assert amplifier_calculation.cabinets_for(4) == 2


def test_the_spare_capacity_is_shown_beside_the_load_not_folded_into_it():
    """The amplifiers are assigned on the plain sum against a limit that
    already carries its own headroom; multiplying the two would count the
    same margin twice."""
    schedule = _schedule(["GF"], [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 20})])
    result = calculate(schedule, taps=TAPS).as_dict()
    assert result["total_watts"] == 20
    assert result["total_watts_with_spare"] == 25        # 20 x 1.25
    assert result["amplifiers"][0]["watts"] == 20        # assigned on the plain sum


# --- through the API ------------------------------------------------------------------------------


def _login(client):
    return login(client, settings.default_admin_email, settings.default_admin_password)


def _upload_schedule(client, project_id: int, tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "FLOOR WISE"
    for index, heading in enumerate(["DRG.NAME", "GF", "1 to 3"], start=1):
        sheet.cell(row=1, column=index, value=heading)
    sheet.cell(row=2, column=1, value="Ceiling Speaker")
    sheet.cell(row=2, column=2, value=10)
    sheet.cell(row=2, column=3, value=14)
    path = tmp_path / "floors.xlsx"
    book.save(path)
    return client.post(f"/projects/{project_id}/floor-schedule",
                       files={"file": (path.name, path.read_bytes(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_the_tab_says_what_is_missing_before_a_schedule_is_read(client):
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30910", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    body = client.get(f"/projects/{project_id}/design/amplifier").json()
    assert body["schedule_file"] is None
    assert any("No floor-wise BOQ has been read" in warning for warning in body["result"]["warnings"])


def test_the_speaker_database_is_offered_and_a_tapping_it_does_not_have_is_refused(client, tmp_path):
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30911", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    _upload_schedule(client, project_id, tmp_path)

    database = {row["part_no"]: row for row in client.get("/design/speakers").json()}
    # The 70 V line figures off the EST-S186 datasheet.
    assert "EST-S186C" in database and database["EST-S186C"]["taps"] == [0.37, 0.75, 1.5, 3]
    assert database["EST-S186C"]["default_tap"] is None

    refused = client.put(f"/projects/{project_id}/design/amplifier",
                         json={"taps": {"EST-S186C": 7}, "counts": {}})
    assert refused.status_code == 400 and "not tapped at 7" in refused.json()["detail"]

    saved = client.put(f"/projects/{project_id}/design/amplifier",
                       json={"taps": {"EST-S186C": 0.75}, "counts": {}})
    assert saved.status_code == 200, saved.text
    # The schedule's speaker line is not settled as a part, so the column is
    # named after the schedule's own wording and carries no tapping yet.
    columns = client.get(f"/projects/{project_id}/design/amplifier").json()["result"]["columns"]
    assert [column["key"] for column in columns] == ["Ceiling Speaker"]
    assert columns[0]["unsettled"] is True and columns[0]["tap"] is None


def test_the_schedule_is_worked_out_from_the_boq_every_time(client, db_session, tmp_path):
    """Nothing about the loading is stored: a speaker added to a floor
    shows without the amplifier tab being re-imported."""
    from app.models import ProjectFloorSchedule

    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30912", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    _upload_schedule(client, project_id, tmp_path)
    # Settle the speaker line as a part the database knows.
    row = db_session.query(ProjectFloorSchedule).filter(
        ProjectFloorSchedule.project_id == project_id).one()
    result = dict(row.result)
    items = [dict(item) for item in result["items"]]
    items[0]["material"] = {"part_no": "EST-S186C", "description": "Ceiling speaker", "manufacturer": "EDWARDS"}
    result["items"] = items
    row.result = result
    db_session.commit()

    body = client.get(f"/projects/{project_id}/design/amplifier").json()["result"]
    assert [column["key"] for column in body["columns"]] == ["EST-S186C"]
    assert body["total_speakers"] == 10 + 14 * 3
    # Nothing is chosen for the engineer, so until a tapping is set the
    # speakers carry no load and the page says why.
    assert body["total_watts"] == 0
    assert any("no tapping set" in warning for warning in body["warnings"])

    client.put(f"/projects/{project_id}/design/amplifier",
               json={"taps": {"EST-S186C": 1.5}, "counts": {}})
    body = client.get(f"/projects/{project_id}/design/amplifier").json()["result"]
    # GF 10 + three typical floors of 14, at the 1.5 W tap.
    assert body["total_watts"] == (10 + 14 * 3) * 1.5
    assert body["amplifier_part"] == AMPLIFIER_PART
    assert len(body["amplifiers"]) == 3 and len(body["cabinets"]) == 2


# --- the speaker database is what the datasheets say ---------------------------------------------


def test_the_tappings_are_the_datasheets_own_figures(client):
    """The database was first seeded from a design mock-up, and its
    figures were wrong: an EST-S186 taps at 3 / 1.5 / 0.75 / 0.37 W on a
    70 V line, not at the quarter-watt steps a Genesis speaker uses. The
    amplifier calculation is only as good as this table, so it is held to
    the datasheets in the library."""
    _login(client)
    database = {row["part_no"]: row for row in client.get("/design/speakers").json()}

    # EST-S186 datasheet, page 2: "Tappings 70 V line, W -- 3 / 1,5 / 0,75 / 0,37".
    assert database["EST-S186C"]["taps"] == [0.37, 0.75, 1.5, 3]
    # Genesis G4 and Integrity 757: 1/4, 1/2, 1 and 2 watt operation.
    assert database["G4SRN"]["taps"] == [0.25, 0.5, 1, 2]
    assert database["757-3A-SS70"]["taps"] == [0.25, 0.5, 1, 2]

    # A datasheet says what a speaker can be set to, not what this company
    # sets it to, so nothing is chosen for the engineer.
    assert all(row["default_tap"] is None for row in database.values())

    # The speakers that were seeded from the mock-up with no datasheet
    # behind them are not offered at all.
    assert not ({"EST-S1830", "EST-S1814", "EST-SPK-15", "EST-SPK-30", "EST-SPK-50"} & set(database))


def test_a_tapping_the_datasheet_does_not_offer_is_refused(client, tmp_path):
    """1 W is a Genesis tap, not an EST-S186 one; on a 70 V line the
    EST-S186 nearest steps are 0.75 and 1.5."""
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30913", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    refused = client.put(f"/projects/{project_id}/design/amplifier",
                         json={"taps": {"EST-S186C": 1}, "counts": {}})
    assert refused.status_code == 400
    assert "0.37, 0.75, 1.5, 3" in refused.json()["detail"].replace(" W", "")

    assert client.put(f"/projects/{project_id}/design/amplifier",
                      json={"taps": {"EST-S186C": 1.5}, "counts": {}}).status_code == 200


# --- one number, two tabs --------------------------------------------------------------------


def _settle(db_session, project_id: int, part: str = "EST-S186C"):
    """Settle the schedule's speaker line as a part the database knows."""
    from app.models import ProjectFloorSchedule

    row = db_session.query(ProjectFloorSchedule).filter(
        ProjectFloorSchedule.project_id == project_id).one()
    result = dict(row.result)
    items = [dict(item) for item in result["items"]]
    items[0]["material"] = {"part_no": part, "description": "Ceiling speaker", "manufacturer": "EDWARDS"}
    result["items"] = items
    row.result = result
    db_session.commit()


def test_a_count_changed_on_the_amplifier_tab_changes_the_boq(client, db_session, tmp_path):
    """The two tabs are two views of one number. A second copy kept on the
    amplifier page would be a way for them to disagree."""
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30914", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    _upload_schedule(client, project_id, tmp_path)
    _settle(db_session, project_id)
    client.put(f"/projects/{project_id}/design/amplifier", json={"taps": {"EST-S186C": 1.5}})

    before = client.get(f"/projects/{project_id}/design/amplifier").json()["result"]
    assert before["floors"][0]["counts"]["EST-S186C"] == 10

    changed = client.patch(f"/projects/{project_id}/design/amplifier/counts",
                           json={"floor": "GF", "part_no": "EST-S186C", "count": 18})
    assert changed.status_code == 200, changed.text
    assert changed.json()["result"]["floors"][0]["counts"]["EST-S186C"] == 18

    # ...and the BOQ Floor Wise tab shows it, because that is where it was set.
    boq = client.get(f"/projects/{project_id}/floor-schedule").json()["result"]
    assert boq["items"][0]["per_floor"]["GF"] == 18
    assert boq["items"][0]["edited"] is True


def test_a_count_changed_on_the_boq_changes_the_amplifier(client, db_session, tmp_path):
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30915", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    _upload_schedule(client, project_id, tmp_path)
    _settle(db_session, project_id)
    client.put(f"/projects/{project_id}/design/amplifier", json={"taps": {"EST-S186C": 1.5}})

    row = client.get(f"/projects/{project_id}/floor-schedule").json()["result"]["items"][0]["row"]
    client.patch(f"/projects/{project_id}/floor-schedule/items/{row}", json={"floor": "GF", "quantity": 2})

    result = client.get(f"/projects/{project_id}/design/amplifier").json()["result"]
    assert result["floors"][0]["counts"]["EST-S186C"] == 2
    assert result["floors"][0]["watts"] == 3          # 2 x 1.5 W


def test_a_part_ordered_for_two_boq_lines_is_not_changed_blindly(client, db_session, tmp_path):
    """Which of the two lines gained a speaker is a question only the
    engineer can answer, so it is asked rather than guessed."""
    from app.models import ProjectFloorSchedule

    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30916", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "FLOOR WISE"
    for index, heading in enumerate(["DRG.NAME", "GF"], start=1):
        sheet.cell(row=1, column=index, value=heading)
    sheet.cell(row=2, column=1, value="Ceiling Speaker")
    sheet.cell(row=2, column=2, value=10)
    sheet.cell(row=3, column=1, value="Wall Speaker")
    sheet.cell(row=3, column=2, value=4)
    path = tmp_path / "two.xlsx"
    book.save(path)
    client.post(f"/projects/{project_id}/floor-schedule",
                files={"file": (path.name, path.read_bytes(),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    # Both lines settled as the same part.
    row = db_session.query(ProjectFloorSchedule).filter(
        ProjectFloorSchedule.project_id == project_id).one()
    result = dict(row.result)
    items = [dict(item) for item in result["items"]]
    for item in items:
        item["material"] = {"part_no": "EST-S186C", "description": "Ceiling speaker", "manufacturer": "EDWARDS"}
    result["items"] = items
    row.result = result
    db_session.commit()

    refused = client.patch(f"/projects/{project_id}/design/amplifier/counts",
                           json={"floor": "GF", "part_no": "EST-S186C", "count": 5})
    assert refused.status_code == 409
    assert "BOQ Floor Wise tab" in refused.json()["detail"]


def test_the_calculation_exports_as_a_pdf(client, db_session, tmp_path):
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30917", "project_name": "Titania",
                                                "design_sheets": []}).json()["id"]
    _upload_schedule(client, project_id, tmp_path)
    _settle(db_session, project_id)
    client.put(f"/projects/{project_id}/design/amplifier", json={"taps": {"EST-S186C": 1.5}})

    export = client.get(f"/projects/{project_id}/design/amplifier/export.pdf")
    assert export.status_code == 200
    assert export.headers["content-type"] == "application/pdf"
    assert export.content.startswith(b"%PDF")
    assert "Amplifier Calculation.pdf" in export.headers["content-disposition"]


# --- the audio riser module ----------------------------------------------------------------------


def test_every_floor_with_speakers_is_fed_through_one_module():
    """One SIGA-CC2A a floor. A floor with no speakers has nothing to
    feed and takes none."""
    floors = ["UG", "GF", "Level 1"]
    schedule = _schedule(floors, [("Ceiling Speaker", "Speaker", "EST-S186C",
                                   {"GF": 10, "Level 1": 14})])
    result = calculate(schedule, taps=TAPS).as_dict()

    assert result["module_part"] == "SIGA-CC2A"
    assert [floor["modules"] for floor in result["floors"]] == [0, 1, 1]
    assert result["total_modules"] == 2


def test_a_floor_over_the_modules_limit_needs_another_and_says_so():
    """A module carries 35 W. The riser is drawn off this, so a floor that
    needs two is worth knowing about before it is drawn."""
    schedule = _schedule(["GF", "Level 1"],
                         [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 40, "Level 1": 10})])
    result = calculate(schedule, taps=TAPS, module={"part_no": "SIGA-CC2A", "max_watts": 35}).as_dict()

    ground, first = result["floors"]
    assert ground["watts"] == 40 and ground["modules"] == 2      # 40 W over a 35 W module
    assert first["watts"] == 10 and first["modules"] == 1
    assert result["total_modules"] == 3
    assert any("more than one SIGA-CC2A can take at 35 W" in warning for warning in result["warnings"])


def test_a_floor_exactly_on_the_modules_limit_still_takes_one():
    schedule = _schedule(["GF"], [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 35})])
    result = calculate(schedule, taps=TAPS, module={"max_watts": 35}).as_dict()
    assert result["floors"][0]["watts"] == 35 and result["floors"][0]["modules"] == 1
    assert not any("more than one" in warning for warning in result["warnings"])


def test_the_module_and_its_limit_come_from_the_design_rule(client):
    """The part and the limit are a design decision, kept where they can be
    corrected without a release."""
    _login(client)
    schedule = _schedule(["GF"], [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 30})])
    # A rule naming another module, with a smaller limit.
    result = calculate(schedule, taps=TAPS,
                       module={"part_no": "SIGA-CC1", "max_watts": 20}).as_dict()
    assert result["module_part"] == "SIGA-CC1"
    assert result["module_max_watts"] == 20
    assert result["floors"][0]["modules"] == 2                   # 30 W over a 20 W module


# --- the staircases, on circuits of their own -------------------------------------------------

STAIR_TAPS = {
    **TAPS,
    "G4SRN": {"part_no": "G4SRN", "description": "Wall speaker", "taps": [0.25, 0.5, 1, 2], "default_tap": 0.5},
}
STAIRCASE = {"parts": ["G4SRN"]}


def test_staircase_speakers_are_taken_off_the_floors():
    """A staircase speaker never shares a floor's circuit or amplifier: it
    leaves the Speakers tab altogether and is worked out on its own."""
    floors = ["GF", "Level 1"]
    schedule = _schedule(floors, [
        ("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 10, "Level 1": 10}),
        ("Wall Speaker", "Speaker", "G4SRN", {"GF": 2, "Level 1": 2}),
    ])
    result = calculate(schedule, taps=STAIR_TAPS, staircase=STAIRCASE).as_dict()

    assert [column["key"] for column in result["columns"]] == ["EST-S186C"]
    assert result["total_speakers"] == 20 and result["total_watts"] == 20
    stair = result["staircase"]
    assert [column["key"] for column in stair["columns"]] == ["G4SRN"]
    assert (stair["total_speakers"], stair["total_watts"]) == (4, 2.0)
    # What the job orders counts both.
    assert result["job"]["speakers"] == 24 and result["job"]["watts"] == 22.0


def test_a_floors_count_of_staircase_speakers_is_its_number_of_stairs():
    """One speaker a floor in each stair: the building has as many stairs as
    its busiest floor, and a stair reaches only the floors that count it."""
    schedule = _schedule(["B1", "GF", "Level 1"], [
        ("Wall Speaker", "Speaker", "G4SRN", {"B1": 2, "GF": 4, "Level 1": 2}),
    ])
    stair = calculate(schedule, taps=STAIR_TAPS, staircase=STAIRCASE).as_dict()["staircase"]

    assert stair["stairs"] == 4
    assert [(c["name"], c["stair"], c["floors"]) for c in stair["circuits"]] == [
        ("ST1-1", 1, ["B1", "GF", "Level 1"]),
        ("ST2-1", 2, ["B1", "GF", "Level 1"]),
        ("ST3-1", 3, ["GF"]),
        ("ST4-1", 4, ["GF"]),
    ]
    by_floor = {row["floor"]: row["circuits"] for row in stair["floors"]}
    assert by_floor["GF"] == ["ST1-1", "ST2-1", "ST3-1", "ST4-1"] and by_floor["B1"] == ["ST1-1", "ST2-1"]


def test_a_stair_starts_a_new_circuit_on_a_new_module_past_its_limit():
    """A stair's circuit takes floors until the next would pass the limit;
    then a new circuit starts, each on a module of its own."""
    floors = [f"Level {n}" for n in range(1, 11)]
    # One stair, a 4 W speaker a floor: eight floors are 32 W, a ninth
    # would be 36 W, over a SIGA-CC2A's 35 W.
    schedule = _schedule(floors, [("Stair Speaker", "Speaker", "G4SRN", {floor: 1 for floor in floors})])
    taps = {**STAIR_TAPS, "G4SRN": {**STAIR_TAPS["G4SRN"], "taps": [4], "default_tap": 4}}
    result = calculate(schedule, taps=taps, staircase=STAIRCASE,
                       module={"part_no": "SIGA-CC2A", "max_watts": 35}).as_dict()
    stair = result["staircase"]

    assert [(c["name"], len(c["floors"]), c["watts"]) for c in stair["circuits"]] == [
        ("ST1-1", 8, 32.0), ("ST1-2", 2, 8.0)]
    assert stair["total_circuits"] == 2 and result["job"]["modules"] == 2


def test_a_staircase_circuit_is_held_to_the_module_rating_where_lower():
    """A SIGA-AA50 may be loaded to 40 W, but a SIGA-CC2A on a 70 V line is
    rated 35 W: the circuit is held to the lower, or the module could not
    switch it. Where the module allows more, the amplifier's 40 W holds."""
    schedule = _schedule(["GF"], [("Wall Speaker", "Speaker", "G4SRN", {"GF": 1})])
    at_70v = calculate(schedule, taps=STAIR_TAPS, staircase=STAIRCASE, module={"max_watts": 35}).as_dict()
    at_25v = calculate(schedule, taps=STAIR_TAPS, staircase=STAIRCASE, module={"max_watts": 50}).as_dict()
    assert at_70v["staircase"]["circuit_limit_watts"] == 35
    assert at_25v["staircase"]["circuit_limit_watts"] == 40


def test_staircase_circuits_are_fed_whole_and_share_the_floors_cabinets():
    """Circuits go onto amplifiers whole, as floors do; the cabinets are one
    set for the job, so an odd amplifier on each tab shares a cabinet."""
    schedule = _schedule(["GF", "Level 1"], [
        ("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 20}),
        ("Wall Speaker", "Speaker", "G4SRN", {"GF": 2, "Level 1": 2}),
    ])
    result = calculate(schedule, taps=STAIR_TAPS, staircase=STAIRCASE).as_dict()
    stair = result["staircase"]

    # Two stairs of 1 W each fit one amplifier together.
    assert [(a["name"], a["circuits"], a["watts"]) for a in stair["amplifiers"]] == [
        (f"{AMPLIFIER_PART.replace('SIGA-', '')}-ST1", ["ST1-1", "ST2-1"], 2.0)]
    # One floor amplifier and one staircase amplifier: one cabinet, not two.
    assert [cabinet["amplifiers"] for cabinet in result["cabinets"]] == [["AA50-1", "AA50-ST1"]]
    assert result["job"] == {"amplifiers": 2, "modules": 1 + 2, "cabinets": 1, "speakers": 24, "watts": 22.0}


def test_a_line_worded_as_a_staircase_is_one_whatever_it_is_ordered_as():
    schedule = _schedule(["GF"], [
        ("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 4}),
        ("Staircase Speaker", "Speaker", "EST-S1814", {"GF": 2}),
    ])
    result = calculate(schedule, taps=TAPS, staircase={"parts": []}).as_dict()
    assert [column["key"] for column in result["columns"]] == ["EST-S186C"]
    assert [column["key"] for column in result["staircase"]["columns"]] == ["EST-S1814"]


def test_without_the_staircase_rule_every_speaker_is_a_floors():
    schedule = _schedule(["GF"], [("Wall Speaker", "Speaker", "G4SRN", {"GF": 2})])
    result = calculate(schedule, taps=STAIR_TAPS).as_dict()
    assert [column["key"] for column in result["columns"]] == ["G4SRN"] and result["staircase"] is None
