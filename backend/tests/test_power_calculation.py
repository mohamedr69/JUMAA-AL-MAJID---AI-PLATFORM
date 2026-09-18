"""The 24 V power calculation, worked out from the floor-wise BOQ.

The same shape as the amplifier's, for everything that runs on direct
current. Two things are particular to it and are what these hold: a
speaker counts only for its flasher, and a sounder base is powered but
needs no module.
"""

import openpyxl

from app.core.config import get_settings
from app.services.power_calculation import calculate

from .conftest import login

settings = get_settings()

# The device currents, as the design rules hold them.
CURRENTS = {
    "G1ARN": {
        "part_no": "G1ARN", "description": "Genesis G1 compact wall horn",
        "currents": [{"ma": 13, "label": "C-LOW / T-LOW, 16-33 VDC"},
                     {"ma": 23, "label": "C-HIGH / T-HIGH, 16-33 VDC"}],
    },
    "202-7A-T": {
        "part_no": "202-7A-T", "description": "Xenon flasher",
        "currents": [{"ma": 90, "label": "15/75 cd, 24 Vdc"}, {"ma": 180, "label": "110 cd, 24 Vdc"}],
    },
    "SIGA-LPS": {
        "part_no": "SIGA-LPS", "description": "Audible (sounder) base",
        "currents": [{"ma": 15, "label": "high dBA, 24 Vdc"}],
    },
}


def _schedule(floors: list[str], lines: list[tuple[str, str, str | None, dict[str, int]]]) -> dict:
    return {
        "floors": floors,
        "items": [
            {"description": description, "device": device, "system": "FAS",
             "material": {"part_no": part, "description": "", "manufacturer": None} if part else None,
             "per_floor": dict(counts), "total": sum(counts.values())}
            for description, device, part, counts in lines
        ],
    }


# --- what draws 24 V ---------------------------------------------------------------------------


def test_a_floors_load_is_its_appliances_times_their_current():
    schedule = _schedule(
        ["GF", "Level 1"],
        [("Wall Sounder", "Sounder", "G1ARN", {"GF": 4, "Level 1": 2}),
         ("Wall Sounder Flasher WP", "Strobe", "202-7A-T", {"GF": 2})],
    )
    result = calculate(schedule, currents=CURRENTS,
                       chosen={"G1ARN": 23, "202-7A-T": 90}).as_dict()

    assert [column["key"] for column in result["columns"]] == ["202-7A-T", "G1ARN"]
    ground, first = result["floors"]
    assert ground["current_ma"] == 4 * 23 + 2 * 90          # 272 mA
    assert first["current_ma"] == 2 * 23
    assert result["total_ma"] == 318 and result["total_amps"] == 0.318


def test_a_plain_speaker_draws_nothing_but_a_speaker_flasher_does():
    """The speaker is on the amplifier's 70 V line; the flasher is not."""
    schedule = _schedule(
        ["GF"],
        [("Ceiling Speaker", "Speaker", "EST-S186C", {"GF": 20}),
         ("Wall Sounder Flasher WP", "Speaker/strobe", "202-7A-T", {"GF": 3})],
    )
    result = calculate(schedule, currents=CURRENTS, chosen={"202-7A-T": 90}).as_dict()
    assert [column["key"] for column in result["columns"]] == ["202-7A-T"]
    assert result["floors"][0]["current_ma"] == 270


def test_a_detector_without_a_sounder_base_is_not_a_24v_load():
    schedule = _schedule(["GF"], [("Smoke Detector", "Smoke detector", "SIGA-OSD-FCN", {"GF": 40})])
    result = calculate(schedule, currents=CURRENTS).as_dict()
    assert result["columns"] == []
    assert any("nothing to draw" in warning for warning in result["warnings"])


def test_a_device_with_no_current_chosen_draws_nothing_and_says_so():
    schedule = _schedule(["GF"], [("Wall Sounder", "Sounder", "G1ARN", {"GF": 4})])
    result = calculate(schedule, currents=CURRENTS).as_dict()
    # Two figures on the datasheet, so the engineer chooses.
    assert result["columns"][0]["current_ma"] is None
    assert result["floors"][0]["current_ma"] == 0
    assert any("no current set" in warning for warning in result["warnings"])


def test_a_device_with_one_figure_on_its_datasheet_takes_it():
    schedule = _schedule(["GF"], [("Smoke with Sounder Base", "Smoke detector with sounder base",
                                   "SIGA-LPS", {"GF": 6})])
    result = calculate(schedule, currents=CURRENTS).as_dict()
    assert result["columns"][0]["current_ma"] == 15
    assert result["floors"][0]["current_ma"] == 90


# --- the module ----------------------------------------------------------------------------------


def test_a_floor_with_a_notification_circuit_takes_one_module():
    schedule = _schedule(["UG", "GF"], [("Wall Sounder", "Sounder", "G1ARN", {"GF": 4})])
    result = calculate(schedule, currents=CURRENTS, chosen={"G1ARN": 23}).as_dict()
    assert result["module_part"] == "SIGA-CC1"
    assert [floor["modules"] for floor in result["floors"]] == [0, 1]
    assert result["total_modules"] == 1


def test_a_sounder_base_is_powered_but_needs_no_module():
    """It sits under a Signature detector on the loop and is addressed
    there, so a floor carrying only bases draws power and takes none."""
    schedule = _schedule(
        ["GF", "Level 1"],
        [("Smoke with Sounder Base", "Smoke detector with sounder base", "SIGA-LPS",
          {"GF": 6, "Level 1": 4}),
         ("Wall Sounder", "Sounder", "G1ARN", {"Level 1": 2})],
    )
    result = calculate(schedule, currents=CURRENTS, chosen={"G1ARN": 23}).as_dict()

    base = next(column for column in result["columns"] if column["key"] == "SIGA-LPS")
    assert base["needs_module"] is False
    ground, first = result["floors"]
    # The ground floor has bases only: powered, no module.
    assert ground["current_ma"] == 6 * 15 and ground["modules"] == 0
    # The first floor also has a sounder, so it has a circuit to drive.
    assert first["current_ma"] == 4 * 15 + 2 * 23 and first["modules"] == 1
    assert result["total_modules"] == 1


def test_a_pair_takes_the_current_of_the_half_that_draws_it():
    """"SIGA-OSD-FCN + SIGA-LPS" is a detector and its audible base; the
    base is what draws 24 V."""
    schedule = _schedule(["GF"], [("Smoke with Sounder Base", "Smoke detector with sounder base",
                                   "SIGA-OSD-FCN + SIGA-LPS", {"GF": 6})])
    result = calculate(schedule, currents=CURRENTS).as_dict()
    assert result["columns"][0]["current_ma"] == 15
    assert result["floors"][0]["current_ma"] == 90


# --- the supplies ----------------------------------------------------------------------------------


def test_floors_are_combined_until_the_next_would_pass_the_supplys_limit():
    floors = ["GF", "Level 1", "Level 2"]
    schedule = _schedule(floors, [("Wall Sounder", "Sounder", "G1ARN", {floor: 100 for floor in floors})])
    # A 1 A supply worked to 80%: 800 mA. Each floor draws 100 x 5 = 500 mA.
    result = calculate(schedule, currents=CURRENTS, chosen={"G1ARN": 5},
                       supply={"part_no": "BPS6A", "amps": 1, "fraction": 0.8}).as_dict()

    assert result["limit_ma"] == 800
    assert [(s["name"], s["floors"], s["current_ma"]) for s in result["supplies"]] == [
        ("BPS-1", ["GF"], 500.0),
        ("BPS-2", ["Level 1"], 500.0),
        ("BPS-3", ["Level 2"], 500.0),
    ]


def test_a_floor_drawing_more_than_one_supply_can_give_is_reported():
    schedule = _schedule(["GF"], [("Wall Sounder", "Sounder", "G1ARN", {"GF": 100})])
    result = calculate(schedule, currents=CURRENTS, chosen={"G1ARN": 23},
                       supply={"part_no": "BPS6A", "amps": 1, "fraction": 0.8}).as_dict()
    assert result["supplies"][0]["over_limit"] is True
    assert any("more than one BPS6A can give" in warning for warning in result["warnings"])


# --- through the API ---------------------------------------------------------------------------------


def _login(client):
    return login(client, settings.default_admin_email, settings.default_admin_password)


def _upload(client, project_id: int, tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "FLOOR WISE"
    for index, heading in enumerate(["DRG.NAME", "GF", "1 to 3"], start=1):
        sheet.cell(row=1, column=index, value=heading)
    sheet.cell(row=2, column=1, value="Wall Sounder")
    sheet.cell(row=2, column=2, value=4)
    sheet.cell(row=2, column=3, value=2)
    path = tmp_path / "floors.xlsx"
    book.save(path)
    return client.post(f"/projects/{project_id}/floor-schedule",
                       files={"file": (path.name, path.read_bytes(),
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_the_device_currents_are_the_datasheets_own_figures(client):
    _login(client)
    database = {row["part_no"]: row for row in client.get("/design/device-currents").json()}

    # Genesis G1 datasheet, page 3: horns draw 13 mA low and 23 mA high on VDC.
    horn = {entry["ma"] for entry in database["G1ARN"]["currents"]}
    assert {13, 23}.issubset(horn)
    # 202 series datasheet, page 3: 90 mA at 15/75 cd, 180 mA at 110 cd.
    flasher = {entry["ma"] for entry in database["202-7A-T"]["currents"]}
    assert {90, 180}.issubset(flasher)
    # Every figure says what condition it is for, so a choice can be made.
    assert all(entry.get("label") for row in database.values() for entry in row["currents"])


def test_a_current_the_datasheet_does_not_give_is_refused(client, tmp_path):
    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30920", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    refused = client.put(f"/projects/{project_id}/design/power", json={"currents": {"G1ARN": 50}})
    assert refused.status_code == 400 and "does not draw 50 mA" in refused.json()["detail"]

    assert client.put(f"/projects/{project_id}/design/power",
                      json={"currents": {"G1ARN": 23}}).status_code == 200


def test_the_power_schedule_is_worked_out_from_the_boq(client, db_session, tmp_path):
    from app.models import ProjectFloorSchedule

    _login(client)
    project_id = client.post("/projects", json={"ep_number": "30921", "project_name": "T",
                                                "design_sheets": []}).json()["id"]
    _upload(client, project_id, tmp_path)
    row = db_session.query(ProjectFloorSchedule).filter(
        ProjectFloorSchedule.project_id == project_id).one()
    result = dict(row.result)
    items = [dict(item) for item in result["items"]]
    items[0]["material"] = {"part_no": "G1ARN", "description": "horn", "manufacturer": "EDWARDS"}
    result["items"] = items
    row.result = result
    db_session.commit()

    client.put(f"/projects/{project_id}/design/power", json={"currents": {"G1ARN": 23}})
    body = client.get(f"/projects/{project_id}/design/power").json()["result"]
    assert [column["key"] for column in body["columns"]] == ["G1ARN"]
    assert body["total_devices"] == 4 + 2 * 3
    assert body["total_ma"] == (4 + 6) * 23
    assert body["module_part"] == "SIGA-CC1" and body["total_modules"] == 4

    export = client.get(f"/projects/{project_id}/design/power/export.pdf")
    assert export.status_code == 200 and export.content.startswith(b"%PDF")
