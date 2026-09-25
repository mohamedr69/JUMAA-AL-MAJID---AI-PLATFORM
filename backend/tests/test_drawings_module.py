"""The Drawings page's records: IFC defines the floors, the shop drawings
define the submission history, the two systems share nothing but the
floors, a file found is not a revision submitted, and what an engineer
confirmed stands over the folder and the AI."""
from __future__ import annotations

import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import ezdxf
import pytest

from app.ai import provider as ai_provider
from app.ai.provider import AiResponse, TextPart, Usage
from app.core.config import get_settings
from app.core.timeutils import utc_now
from app.models import (DrawingIssue, Project, ProjectBuildingFloor, ProjectDocument, ProjectShopDrawing,
                        ShopDrawingCandidate, ShopDrawingRevision)
from app.services import document_sync, shop_drawings
from app.services.document_control import ControlledDocument

from .conftest import login

settings = get_settings()
FA, EL = "EP-SDW-FA", "EP-SDW-EL"


def _admin(client):
    assert login(client, settings.default_admin_email, settings.default_admin_password).status_code == 200


def _project(client, tmp_path, ep: str, systems=None) -> tuple[int, Path]:
    folder = tmp_path / f"EP-{ep}"
    folder.mkdir(parents=True, exist_ok=True)
    body = {"ep_number": ep, "project_name": "Drawings", "design_sheets": [], "source_folder_path": str(folder)}
    if systems is not None:
        body["systems"] = systems
    return client.post("/projects", json=body).json()["id"], folder


BOTH = [{"name": "Fire Alarm", "brand": "EDWARDS", "method_statement": True, "drawing": True},
        {"name": "Emergency Light Monitoring", "brand": "MENVIER", "method_statement": True, "drawing": True}]


def _doc(floor, revision, status, *, ref, system="FAS", day=1, reply=None, name=None, path=None, category="drawings",
         source="document"):
    folder = "FA" if system == "FAS" else "ELS"
    return ControlledDocument(system, name or f"{floor or ''} PLAN", path or f"03- Drawings/SD/{folder}/{revision}/{ref}.pdf",
                              datetime(2026, 9, day), ref, revision, status, floor=floor, reply_text=reply,
                              category=category, source=source)


def _seed(db, pid: int, folder: Path, records: list[ControlledDocument], *, sha: dict[str, str] | None = None) -> None:
    """The document index as a sync would leave it: one row per file, its
    records on it. Existing rows for a path are replaced."""
    sha = sha or {}
    for record in records:
        existing = db.query(ProjectDocument).filter(ProjectDocument.project_id == pid,
                                                    ProjectDocument.relative_path == record.path).first()
        data = {**asdict(record), "modified": record.modified.isoformat()}
        data.pop("superseded", None)
        if existing is None:
            db.add(ProjectDocument(project_id=pid, role="document", path=str(folder / record.path), relative_path=record.path,
                                   filename=Path(record.path).name, state="fresh", findings=[], acknowledged=[],
                                   sha256=sha.get(record.path, f"sha-{Path(record.path).stem}-{record.revision}"),
                                   extracted={"records": [data], "notes": []}))
        else:
            existing.extracted = {"records": [data], "notes": []}
            existing.state = "fresh"
            if record.path in sha:
                existing.sha256 = sha[record.path]
    db.query(Project).filter(Project.id == pid).update({"documents_synced_at": utc_now()})
    db.commit()


def _remove(db, pid: int, path: str) -> None:
    row = db.query(ProjectDocument).filter(ProjectDocument.project_id == pid, ProjectDocument.relative_path == path).one()
    row.state = "removed"
    db.query(Project).filter(Project.id == pid).update({"documents_synced_at": utc_now()})
    db.commit()


def _reconcile(client, pid: int) -> dict:
    response = client.post(f"/projects/{pid}/drawings/reconcile")
    assert response.status_code == 200, response.text
    return response.json()


def _log(client, pid: int, system: str) -> dict:
    response = client.get(f"/projects/{pid}/drawings/log?system={system}")
    assert response.status_code == 200, response.text
    return response.json()


def _row(log: dict, reference: str) -> dict:
    return next(r for r in log["rows"] if r["reference"] == reference)


def _cells(row: dict) -> dict[str, str]:
    return {rev: cell["status"] for rev, cell in row["cells"].items()}


def _dxf(path: Path, floors: list[str]) -> Path:
    """An IFC drawing with a plan sheet per floor, its title block naming
    the floor (the way test_ifc_boq's building set is drawn)."""
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 4
    block = doc.blocks.new("SD")
    block.add_circle((0, 0), 200)
    block.add_text("S", height=150).set_placement((-50, -75))
    msp = doc.modelspace()
    title_block = doc.blocks.new("TITLE")
    title_block.add_attdef("TITLE", (0, 0), dxfattribs={"height": 5})
    for i, floor in enumerate(floors):
        x = i * 100_000
        for k in range(3):
            msp.add_blockref("SD", (x + 5000 + k * 3000, 5000), dxfattribs={"layer": "E-FIRE"})
        lay = doc.layouts.new(f"FA-{i + 1:02d}")
        lay.add_viewport(center=(200, 150), size=(400, 300), view_center_point=(x + 20_000, 5000), view_height=30_000)
        lay.add_blockref("TITLE", (10, 10)).add_auto_attribs({"TITLE": f"{floor} PLAN"})
    doc.saveas(path)
    return path


def _import_ifc(client, pid: int, path: Path):
    response = client.post(f"/projects/{pid}/ifc-drawings", files={"file": (path.name, path.read_bytes(), "application/dxf")})
    assert response.status_code == 201, response.text
    return response.json()


# --- FAS / ELS: shared floors, separate state -------------------------------------------------------


def test_the_two_systems_share_floors_and_nothing_else(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93001", BOTH)
    _seed(db_session, pid, folder, [
        _doc("Level 1", "R0", "rejected", ref=f"{FA}-L01"), _doc("Level 1", "R1", "approved", ref=f"{FA}-L01", day=2),
        _doc("Level 1", "R0", "UR", ref=f"{EL}-L01", system="ELS"),
    ])
    _reconcile(client, pid)
    fas, els = _log(client, pid, "FAS"), _log(client, pid, "ELS")
    assert fas["systems"] == ["FAS", "ELS"] and fas["system"] == "FAS" and els["system"] == "ELS"
    fa = _row(fas, f"{FA}-L01")
    assert (fa["latest_revision"], fa["latest_status"]) == ("R1", "approved") and _cells(fa)["R0"] == "not_approved"
    el = _row(els, f"{EL}-L01")
    assert (el["latest_revision"], el["latest_status"]) == ("R0", "under_review")
    assert all(r["reference"] != f"{FA}-L01" for r in els["rows"]) and all(r["reference"] != f"{EL}-L01" for r in fas["rows"])
    # The same floor, one row each side.
    assert fa["floor_keys"] == el["floor_keys"] == ["L1"]
    # A system the project has not got is refused, never quietly answered with another's log.
    assert client.get(f"/projects/{pid}/drawings/log?system=PAVA").status_code == 404
    summary = {s["code"]: s for s in client.get(f"/projects/{pid}/drawings/summary").json()["systems"]}
    assert summary["FAS"]["approved"] == 1 and summary["ELS"]["under_review"] == 1 and summary["ELS"]["approved"] == 0


# --- IFC provides floors only -----------------------------------------------------------------------


def test_ifc_gives_the_log_its_floors_and_changes_no_shop_drawing(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93002")
    _seed(db_session, pid, folder, [_doc("Ground Floor", "R0", "approved", ref=f"{FA}-GF")])
    _reconcile(client, pid)
    before = _row(_log(client, pid, "FAS"), f"{FA}-GF")
    _import_ifc(client, pid, _dxf(tmp_path / "FA-IFC.dxf", ["GROUND FLOOR", "1ST FLOOR", "2ND FLOOR"]))
    log = _log(client, pid, "FAS")
    keys = [f.floor_key for f in db_session.query(ProjectBuildingFloor).filter(ProjectBuildingFloor.project_id == pid)
            .order_by(ProjectBuildingFloor.sort_order)]
    assert keys == ["GF", "L1", "L2"]
    after = _row(log, f"{FA}-GF")
    assert (after["latest_revision"], after["latest_status"], _cells(after)) == (
        before["latest_revision"], before["latest_status"], _cells(before))
    # The new floors: a row each, not submitted, and no revision record made up for them.
    assert [(r["floor"], r["source"], r["latest_status"]) for r in log["rows"] if r["source"] == "ifc_floor"] == [
        ("1st Floor", "ifc_floor", "not_submitted"), ("2nd Floor", "ifc_floor", "not_submitted")]
    assert db_session.query(ShopDrawingRevision).count() == 1
    # Nothing of the IFC drawing's is a shop drawing's.
    assert not any("FA 10" in str(v) or "FA-IFC" in str(v) for r in log["rows"] for v in r.values())


def test_a_typical_ifc_sheet_is_a_floor_each_and_a_removed_floor_keeps_its_history(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93003")
    _import_ifc(client, pid, _dxf(tmp_path / "typ.dxf", ["GROUND FLOOR", "TYPICAL 3RD TO 5TH FLOOR", "ROOF"]))
    keys = [f.floor_key for f in db_session.query(ProjectBuildingFloor).filter(ProjectBuildingFloor.project_id == pid)
            .order_by(ProjectBuildingFloor.sort_order)]
    assert keys == ["GF", "L3", "L4", "L5", "RF"]
    log = _log(client, pid, "FAS")
    assert [r["floor"] for r in log["rows"]] == ["Ground Floor", "Level 3", "Level 4", "Level 5", "Roof"]

    # L4 has history; the roof has none. A new IFC drops both.
    _seed(db_session, pid, folder, [_doc("Level 4", "R0", "ANN", ref=f"{FA}-L04")])
    _reconcile(client, pid)
    first = client.get(f"/projects/{pid}/ifc-drawings").json()[0]
    revised = client.post(f"/projects/{pid}/ifc-drawings", data={"supersedes_id": str(first["id"])},
                          files={"file": ("typ.dxf", _dxf(tmp_path / "typ2.dxf", ["GROUND FLOOR", "3RD FLOOR"]).read_bytes(),
                                          "application/dxf")})
    assert revised.status_code == 201, revised.text
    db_session.expire_all()
    floors = {f.floor_key: f for f in db_session.query(ProjectBuildingFloor).filter(ProjectBuildingFloor.project_id == pid)}
    assert floors["L4"].active and floors["RF"].active is False and floors["L5"].active is False
    log = _log(client, pid, "FAS")
    l4 = _row(log, f"{FA}-L04")
    assert (l4["latest_revision"], l4["latest_status"]) == ("R0", "approved_as_noted")
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()
    assert [(i["kind"], i["floor_key"]) for i in issues["system_checks"]] == [("floor_not_in_ifc", "L4")]
    assert "Roof" not in [r["floor"] for r in log["rows"]]


# --- a file found is not a revision submitted -------------------------------------------------------


def test_r1_found_after_r0_approved_is_a_candidate(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93010")
    _seed(db_session, pid, folder, [_doc("Level 5", "R0", "approved", ref=f"{FA}-L05"),
                                    _doc("Level 5", "R1", "UR", ref=f"{FA}-L05", day=2)])
    _reconcile(client, pid)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L05")
    assert _cells(row)["R0"] == "approved" and row["cells"]["R1"]["label"] == "—"
    assert row["cells"]["R1"]["status"] not in ("under_review",) and "candidate" in row["cells"]["R1"]
    assert (row["latest_revision"], row["latest_status"], row["latest_note"]) == ("R0", "approved", "R1 available")
    assert [h["label"] for h in row["hints"]] == ["R1 available"]
    assert db_session.query(ShopDrawingRevision).count() == 1
    (candidate,) = db_session.query(ShopDrawingCandidate).all()
    assert candidate.revision == "R1" and candidate.candidate_status == "available"
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()
    assert [(i["kind"], i["severity"], i["source"]) for i in issues["system_checks"]] == [("revision_candidate", "info", "system")]
    summary = client.get(f"/projects/{pid}/drawings/summary").json()["systems"][0]
    assert summary["approved"] == 1 and summary["under_review"] == 0 and summary["candidates"] == 1 and summary["review_items"] == 1

    # Confirmed by the engineer: R1 is the official revision, under review.
    confirmed = client.post(f"/projects/{pid}/drawings/candidates/{candidate.id}/confirm", json={"submission_reference": "SD-249"}).json()
    assert (confirmed["latest_revision"], confirmed["latest_status"]) == ("R1", "under_review")
    assert confirmed["cells"]["R0"]["status"] == "approved" and confirmed["cells"]["R1"]["source"] == "engineer"
    assert client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["total"] == 0
    # And the next sync does not take it back.
    _reconcile(client, pid)
    again = _row(_log(client, pid, "FAS"), f"{FA}-L05")
    assert (again["latest_revision"], again["latest_status"]) == ("R1", "under_review")
    kinds = [e["kind"] for e in client.get(f"/projects/{pid}/drawings/activity?system=FAS").json()["events"]]
    assert "revision.confirmed" in kinds and "revision.detected" in kinds


def test_r1_answered_by_the_consultant_is_the_latest_revision(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93011")
    _seed(db_session, pid, folder, [_doc("Level 6", "R0", "approved", ref=f"{FA}-L06"),
                                    _doc("Level 6", "R1", "ANN", ref=f"{FA}-L06", day=2, reply="Incorporate comments")])
    _reconcile(client, pid)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L06")
    assert (row["latest_revision"], row["latest_status"]) == ("R1", "approved_as_noted")
    assert _cells(row) == {"R0": "approved", "R1": "approved_as_noted", "R2": "not_submitted"}
    assert row["hints"] == [] and row["candidates"] == []


def test_r1_after_a_rejection_is_the_resubmission(client, db_session, tmp_path):
    """After "Not Approved" the folder's R1 is the resubmission, filed
    where submissions are filed: under review, and R0 keeps its answer."""
    _admin(client)
    pid, folder = _project(client, tmp_path, "93012")
    _seed(db_session, pid, folder, [_doc("Level 7", "R0", "rejected", ref=f"{FA}-L07", reply="Revise"),
                                    _doc("Level 7", "R1", "UR", ref=f"{FA}-L07", day=2)])
    _reconcile(client, pid)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L07")
    assert _cells(row) == {"R0": "not_approved", "R1": "under_review", "R2": "not_submitted"}
    assert (row["latest_revision"], row["latest_status"]) == ("R1", "under_review")
    # Synced again and again: R0 never becomes "not submitted".
    for _ in range(2):
        _reconcile(client, pid)
    assert _cells(_row(_log(client, pid, "FAS"), f"{FA}-L07"))["R0"] == "not_approved"


def test_a_skipped_revision_is_flagged_not_invented(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93013")
    _seed(db_session, pid, folder, [_doc("Level 8", "R0", "rejected", ref=f"{FA}-L08"),
                                    _doc("Level 8", "R2", "UR", ref=f"{FA}-L08", day=3)])
    _reconcile(client, pid)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L08")
    assert row["cells"]["R1"]["status"] == "reply_not_found" and row["cells"]["R1"]["path"] is None
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]
    assert [(i["kind"], i["detail"]["revision"]) for i in issues] == [("revision_gap", "R1")]


# --- candidates: ignored, changed, in conflict ------------------------------------------------------


def test_an_ignored_candidate_is_not_asked_again_until_the_file_changes(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93020")
    path = f"03- Drawings/SD/FA/R1/{FA}-L09.pdf"
    _seed(db_session, pid, folder, [_doc("Level 9", "R0", "approved", ref=f"{FA}-L09"),
                                    _doc("Level 9", "R1", "UR", ref=f"{FA}-L09", day=2)], sha={path: "aaa"})
    _reconcile(client, pid)
    (candidate,) = db_session.query(ShopDrawingCandidate).all()
    client.post(f"/projects/{pid}/drawings/candidates/{candidate.id}/ignore", json={"reason": "a draft"})
    assert client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["total"] == 0
    _reconcile(client, pid)
    assert client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["total"] == 0
    row = _row(_log(client, pid, "FAS"), f"{FA}-L09")
    assert row["latest_revision"] == "R0" and row["hints"] == []
    # The file changed: a new candidate, asked about afresh.
    _seed(db_session, pid, folder, [_doc("Level 9", "R1", "UR", ref=f"{FA}-L09", day=3)], sha={path: "bbb"})
    _reconcile(client, pid)
    db_session.expire_all()
    statuses = sorted((c.file_sha256, c.candidate_status) for c in db_session.query(ShopDrawingCandidate).all())
    assert statuses == [("aaa", "ignored"), ("bbb", "available")]
    assert client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["total"] == 1


def test_two_files_for_one_revision(client, db_session, tmp_path):
    """The same content twice is a copy; a package beside its bare sheet
    is one submission; the stamped copy under Received is the answer. Two
    different files side by side under different names are a conflict."""
    _admin(client)
    pid, folder = _project(client, tmp_path, "93021")
    a = f"03- Drawings/SD/FA/R0/{FA}-L10.pdf"
    package = f"03- Drawings/SD/FA/R0/{FA}-L10 Shop Drawing for Level 10 Floor Plan.pdf"
    received = f"03- Drawings/SD/FA/R0/Received/{FA}-L10.pdf"
    other = f"03- Drawings/SD/FA/R0/OLD-SHEET-L10.pdf"
    _seed(db_session, pid, folder, [_doc("Level 10", "R0", "UR", ref=f"{FA}-L10", path=a),
                                    _doc("Level 10", "R0", "UR", ref=f"{FA}-L10", path=package),
                                    _doc("Level 10", "R0", "ANN", ref=f"{FA}-L10", path=received)],
          sha={a: "sheet", package: "package", received: "stamped"})
    _reconcile(client, pid)
    assert [i["kind"] for i in client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]] == []
    _seed(db_session, pid, folder, [_doc("Level 10", "R0", "UR", ref=f"{FA}-L10", path=other, day=2)], sha={other: "different"})
    _reconcile(client, pid)
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]
    assert [(i["kind"], i["severity"]) for i in issues] == [("revision_conflict", "error")]
    assert "OLD-SHEET-L10.pdf" in issues[0]["text"]


def test_a_reference_used_twice_for_one_floor_is_a_conflict(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93022")
    _seed(db_session, pid, folder, [_doc("Level 11", "R0", "approved", ref=f"{FA}-0105"),
                                    _doc("Level 11", "R1", "UR", ref=f"{FA}-0150", day=2)])
    _reconcile(client, pid)
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]
    conflict = next(i for i in issues if i["kind"] == "reference_conflict")
    assert conflict["severity"] == "error" and {r["reference"] for r in conflict["detail"]["references"]} == {f"{FA}-0105", f"{FA}-0150"}
    # The same sheet re-issued with its placeholder filled in is one drawing, not a conflict.
    pid2, folder2 = _project(client, tmp_path, "93025")
    _seed(db_session, pid2, folder2, [_doc("L22", "R0", "rejected", ref="BBY006-GME-SDW-FP-FA-ZZZ-ZZZ-010009"),
                                      _doc("L22", "R1", "UR", ref="BBY006-GME-SDW-FP-FA-ZZZ-L22-010009", day=2)])
    _reconcile(client, pid2)
    assert [i["kind"] for i in client.get(f"/projects/{pid2}/drawings/issues?system=FAS").json()["system_checks"]] == []


def test_a_drawing_whose_reference_names_another_system_is_not_logged_under_this_one(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93023", BOTH)
    _seed(db_session, pid, folder, [_doc("Level 12", "R0", "UR", ref="EP-SDW-EL-LI-L12", system="ELS"),
                                    _doc("Level 12", "R0", "UR", ref="EP-SDW-FP-FA-L12", system="ELS", path="03- Drawings/SD/ELS/R0/stray.pdf")])
    _reconcile(client, pid)
    els = _log(client, pid, "ELS")
    assert [r["reference"] for r in els["rows"] if r["reference"]] == ["EP-SDW-EL-LI-L12"]
    issues = client.get(f"/projects/{pid}/drawings/issues?system=ELS").json()["system_checks"]
    assert [(i["kind"], i["detail"]["named_system"]) for i in issues] == [("system_mismatch", "FAS")]


def test_one_submission_for_several_floors_is_one_drawing_and_a_row_per_floor(client, db_session, tmp_path):
    """"Basement 4, 3, 2" is three floors drawn apart and sent together:
    one drawing record under the one reference, and a row for each
    basement in the log -- each with the same revisions."""
    _admin(client)
    pid, folder = _project(client, tmp_path, "93024")
    _seed(db_session, pid, folder, [
        _doc("Basement 4, 3, 2", "R0", "rejected", ref=f"{FA}-BSM", name=f"{FA}-BSM Shop Drawing for Basement 4, 3, 2 Floor Plan"),
        _doc("Basement 4, 3, 2", "R1", "UR", ref=f"{FA}-BSM", day=2, name=f"{FA}-BSM Shop Drawing for Basement 4, 3, 2 Floor Plan"),
    ])
    _reconcile(client, pid)
    _reconcile(client, pid)                                        # and again: the same drawing, not a second R0
    drawings = db_session.query(ProjectShopDrawing).all()
    assert len(drawings) == 1 and sorted(drawings[0].floor_keys) == ["B2", "B3", "B4"] and not drawings[0].typical
    assert [r.revision for r in drawings[0].revisions] == ["R0", "R1"]
    log = _log(client, pid, "FAS")
    assert [(r["floor"], r["reference"], r["latest_revision"]) for r in log["rows"]] == [
        ("Basement 4", f"{FA}-BSM", "R1"), ("Basement 3", f"{FA}-BSM", "R1"), ("Basement 2", f"{FA}-BSM", "R1")]
    assert log["counts"] == {"under_review": 3}


# --- what stands: engineer over folder, history over sync ------------------------------------------


def test_an_engineers_status_stands_over_the_next_sync(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93030")
    _seed(db_session, pid, folder, [_doc("Level 13", "R0", "UR", ref=f"{FA}-L13")])
    _reconcile(client, pid)
    drawing = db_session.query(ProjectShopDrawing).one()
    set_ = client.put(f"/projects/{pid}/drawings/sd/{drawing.id}/revisions/R0",
                      json={"status": "approved", "note": "Approval received by email"}).json()
    assert set_["latest_status"] == "approved" and set_["cells"]["R0"]["confirmed"]
    # The folder still says under review, and now says not approved: the engineer's word stands, and it is an issue.
    _seed(db_session, pid, folder, [_doc("Level 13", "R0", "rejected", ref=f"{FA}-L13", day=2)])
    _reconcile(client, pid)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L13")
    assert row["latest_status"] == "approved"
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]
    assert [i["kind"] for i in issues] == ["status_conflict"]
    assert client.put(f"/projects/{pid}/drawings/sd/{drawing.id}/revisions/R0", json={"status": "maybe"}).status_code == 422
    assert client.put(f"/projects/{pid}/drawings/sd/{drawing.id}/revisions/X", json={"status": "approved"}).status_code == 422


def test_a_file_that_disappears_changes_no_status(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93031")
    path = f"03- Drawings/SD/FA/R0/{FA}-L14.pdf"
    _seed(db_session, pid, folder, [_doc("Level 14", "R0", "approved", ref=f"{FA}-L14", path=path)])
    _reconcile(client, pid)
    _remove(db_session, pid, path)
    _reconcile(client, pid)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L14")
    assert (row["latest_revision"], row["latest_status"]) == ("R0", "approved") and row["cells"]["R0"]["source_missing"]
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]
    assert [i["kind"] for i in issues] == ["source_missing"]


def test_the_page_reads_the_records_not_the_index(client, db_session, tmp_path, monkeypatch):
    """Once the records are up to date, a GET reads them: the index is not
    re-read, the folder is not walked, no IFC drawing is resolved."""
    _admin(client)
    pid, folder = _project(client, tmp_path, "93032")
    _seed(db_session, pid, folder, [_doc("Level 15", "R0", "ANN", ref=f"{FA}-L15")])
    _reconcile(client, pid)

    def boom(*_args, **_kwargs):
        raise AssertionError("the page re-read the index")

    monkeypatch.setattr(document_sync, "log_records", boom)
    from app.ifc import resolve as ifc_resolve

    monkeypatch.setattr(ifc_resolve, "resolved_drawing", boom)
    log = _log(client, pid, "FAS")
    assert _row(log, f"{FA}-L15")["latest_status"] == "approved_as_noted"
    assert client.get(f"/projects/{pid}/drawings/summary").status_code == 200
    assert client.get(f"/projects/{pid}/drawings/issues").status_code == 200


def test_a_project_synced_before_the_records_existed_catches_up_once(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93033")
    _seed(db_session, pid, folder, [_doc("Level 16", "R0", "UR", ref=f"{FA}-L16")])
    assert db_session.get(Project, pid).drawings_reconciled_at is None
    log = _log(client, pid, "FAS")             # no reconcile asked for: the page catches up by itself
    assert _row(log, f"{FA}-L16")["latest_status"] == "under_review"
    db_session.expire_all()
    assert db_session.get(Project, pid).drawings_reconciled_at is not None


# --- one drawing: detail, correction, activity ------------------------------------------------------


def test_the_details_panel_and_a_correction(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93040")
    _seed(db_session, pid, folder, [_doc("Level 17", "R0", "rejected", ref=f"{FA}-L17", reply="Fix it"),
                                    _doc("Level 17", "R1", "approved", ref=f"{FA}-L17", day=2)])
    _reconcile(client, pid)
    drawing = db_session.query(ProjectShopDrawing).one()
    detail = client.get(f"/projects/{pid}/drawings/sd/{drawing.id}").json()
    assert detail["system"] == "FAS" and detail["floor"] == "Level 17" and detail["latest_revision"] == "R1"
    assert [(r["revision"], r["status"]) for r in detail["revision_history"]] == [("R0", "not_approved"), ("R1", "approved")]
    assert detail["floor_source"] == {"L17": {"source": "shop_drawing", "ifc_sheet": None, "active": True}}
    assert [e["kind"] for e in detail["events"]][-1] == "drawing.found"
    corrected = client.patch(f"/projects/{pid}/drawings/sd/{drawing.id}", json={"drawing_reference": f"{FA}-L17-A", "remarks": "Renumbered"}).json()
    assert corrected["reference"] == f"{FA}-L17-A" and corrected["confirmed"] and corrected["remarks"] == "Renumbered"
    assert client.patch(f"/projects/{pid}/drawings/sd/{drawing.id}", json={"floor_keys": ["L99"]}).status_code == 422
    assert client.patch(f"/projects/{pid}/drawings/sd/{drawing.id}", json={"drawing_reference": "  "}).status_code == 422
    # A sync finds the old reference again: the corrected drawing keeps its floors, the old one is another record.
    _reconcile(client, pid)
    db_session.expire_all()
    references = sorted(d.drawing_reference for d in db_session.query(ProjectShopDrawing).all())
    assert references == [f"{FA}-L17", f"{FA}-L17-A"]
    events = client.get(f"/projects/{pid}/drawings/activity?system=FAS").json()["events"]
    assert "reference.corrected" in [e["kind"] for e in events]
    assert client.get(f"/projects/{pid}/drawings/activity?all_systems=true").status_code == 200


def test_an_issue_an_engineer_resolves_stays_resolved(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "93041")
    _seed(db_session, pid, folder, [_doc("Level 18", "R0", "rejected", ref=f"{FA}-L18"),
                                    _doc("Level 18", "R2", "UR", ref=f"{FA}-L18", day=3)])
    _reconcile(client, pid)
    (issue,) = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["system_checks"]
    assert client.post(f"/projects/{pid}/drawings/issues/{issue['id']}/resolve", json={"resolution": "R1 was sent by hand"}).status_code == 200
    _reconcile(client, pid)
    assert client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["total"] == 0
    assert db_session.get(DrawingIssue, issue["id"]).resolution == "R1 was sent by hand"


# --- export, open folder --------------------------------------------------------------------------


def test_the_export_is_the_systems_own_and_shows_a_candidate_as_a_hint(client, db_session, tmp_path):
    from openpyxl import load_workbook

    _admin(client)
    pid, folder = _project(client, tmp_path, "93050", BOTH)
    _seed(db_session, pid, folder, [_doc("Level 19", "R0", "approved", ref=f"{EL}-L19", system="ELS"),
                                    _doc("Level 19", "R1", "UR", ref=f"{EL}-L19", system="ELS", day=2)])
    _reconcile(client, pid)
    export = client.get(f"/projects/{pid}/drawings/log/export.xlsx?system=ELS")
    assert export.status_code == 200 and "Drawings Log ELS.xlsx" in export.headers["content-disposition"]
    ws = load_workbook(filename=__import__("io").BytesIO(export.content)).active
    header = [c.value for c in ws[4]]
    row = next(r for r in ws.iter_rows(min_row=5, values_only=True) if r[2] == f"{EL}-L19")
    assert header[4:6] == ["R0", "R1"] and row[4] == "Approved" and row[5] == "—"
    assert row[header.index("Latest Revision")] == "R0" and row[header.index("Issues / Hints")] == "R1 available"
    assert "Drawings Log FAS.xlsx" in client.get(f"/projects/{pid}/drawings/log/export.xlsx?system=FAS").headers["content-disposition"]
    assert client.get(f"/projects/{pid}/drawings/log/export.xlsx?system=PAVA").status_code == 404


def test_open_folder_is_the_systems_own_and_only_for_a_browser_on_this_pc(client, db_session, tmp_path, monkeypatch):
    from app.routers import drawings as drawings_router

    _admin(client)
    pid, folder = _project(client, tmp_path, "93051", BOTH)
    (folder / "03- Drawings" / "SD" / "ELS").mkdir(parents=True, exist_ok=True)
    # The test client is not a browser on this PC.
    assert client.post(f"/projects/{pid}/drawings/open-folder", json={"system": "ELS"}).status_code == 403

    def request(host="127.0.0.1", headers=None):
        return SimpleNamespace(client=SimpleNamespace(host=host), headers={"host": "localhost:8000", **(headers or {})})

    assert drawings_router._local_desktop(request()) is None
    assert drawings_router._local_desktop(request(host="10.0.0.5")) is not None
    assert "proxy" in drawings_router._local_desktop(request(headers={"x-forwarded-for": "10.0.0.5"}))
    assert "another address" in drawings_router._local_desktop(request(headers={"origin": "http://intranet:5173"}))
    assert drawings_router._local_desktop(request(headers={"origin": "http://localhost:5173"})) is None
    monkeypatch.setattr(settings, "desktop_actions_enabled", False)
    assert "switched off" in drawings_router._local_desktop(request())
    monkeypatch.setattr(settings, "desktop_actions_enabled", True)

    opened = []
    monkeypatch.setattr(drawings_router.subprocess, "Popen", lambda args: opened.append(args))
    monkeypatch.setattr(drawings_router, "_local_desktop", lambda _request: None)
    monkeypatch.setattr(drawings_router.os, "name", "nt")
    assert client.post(f"/projects/{pid}/drawings/open-folder", json={"system": "ELS"}).status_code == 204
    assert opened[-1][-1].endswith(str(Path("03- Drawings") / "SD" / "ELS"))
    assert client.post(f"/projects/{pid}/drawings/open-folder", json={"system": "PAVA"}).status_code == 404
    assert client.post(f"/projects/{pid}/drawings/open-folder", json={"path": "../../etc"}).status_code == 403


# --- the AI: exceptions only, never over the engineer -----------------------------------------------


class FakeAi:
    name, ready, status = "fake", True, "a fake provider (tests)"

    def __init__(self, answer):
        self.answer = answer
        self.requests = []

    def complete(self, request):
        self.requests.append(request)
        payload = json.loads(next(p.text for p in request.parts if isinstance(p, TextPart)))
        data = self.answer(payload) if callable(self.answer) else self.answer
        return AiResponse(data=data, usage=Usage(input_tokens=200, output_tokens=30), model="fake-model", latency_ms=1)


@pytest.fixture()
def ai(monkeypatch):
    monkeypatch.setattr(settings, "ai_enabled", True)
    holder = {"provider": None}

    def use(answer):
        holder["provider"] = FakeAi(answer)
        ai_provider.set_provider(holder["provider"])
        return holder["provider"]

    yield use
    ai_provider.set_provider(None)


def _reply(reference, revision, status, *, path, words="approved as noted"):
    return ControlledDocument(None, "Reply to consultant comments", path, datetime(2026, 9, 5), reference, revision, status,
                              None, words, 1, source="reply", category="reply")


def test_a_stray_reply_the_ai_cannot_place_is_an_issue_not_a_status(client, db_session, tmp_path, ai):
    fake = ai({"drawing_reference": None, "revision": None, "status": "unknown", "confidence": 0.2,
               "reason_code": "INSUFFICIENT_EVIDENCE", "requires_engineer": True})
    _admin(client)
    pid, folder = _project(client, tmp_path, "93060")
    _seed(db_session, pid, folder, [_doc("Level 20", "R0", "UR", ref=f"{FA}-L20-010020"),
                                    _reply(f"{FA}-L2O-010020", "R0", "ANN", path="03- Drawings/SD/FA/R0/Received/reply.pdf")])
    project = db_session.get(Project, pid)
    shop_drawings.reconcile(db_session, project, ai=True)
    assert fake.requests and "93060" not in fake.requests[0].parts[0].text
    row = _row(_log(client, pid, "FAS"), f"{FA}-L20-010020")
    assert row["latest_status"] == "under_review"
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()
    assert [(i["kind"], i["source"]) for i in issues["ai_review"]] == [("reply_unmatched", "ai")]
    assert issues["system_checks"] == []
    # Unchanged files: the cache answers, the model is not asked again.
    shop_drawings.reconcile(db_session, project, ai=True)
    assert len(fake.requests) == 1


def test_a_reply_the_ai_places_with_confidence_sets_the_status_but_never_over_an_engineer(client, db_session, tmp_path, ai):
    def answer(payload):
        return {"drawing_reference": payload["candidates"][0]["reference"], "revision": "R0", "status": "approved_as_noted",
                "confidence": 0.99, "reason_code": "REFERENCE_PARTIAL_MATCH", "requires_engineer": False}

    fake = ai(answer)
    _admin(client)
    pid, folder = _project(client, tmp_path, "93061")
    _seed(db_session, pid, folder, [_doc("Level 21", "R0", "UR", ref=f"{FA}-L21-010021"),
                                    _reply(f"{FA}-L2I-010021", "R0", "ANN", path="03- Drawings/SD/FA/R0/Received/reply.pdf")])
    project = db_session.get(Project, pid)
    shop_drawings.reconcile(db_session, project, ai=True)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L21-010021")
    assert row["latest_status"] == "approved_as_noted" and row["cells"]["R0"]["source"] == "ai"
    assert client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()["total"] == 0

    # The engineer says approved. The AI, asked again about a changed reply, says not approved: the engineer stands.
    drawing = db_session.query(ProjectShopDrawing).one()
    client.put(f"/projects/{pid}/drawings/sd/{drawing.id}/revisions/R0", json={"status": "approved"})
    fake.answer = lambda payload: {"drawing_reference": payload["candidates"][0]["reference"], "revision": "R0",
                                   "status": "not_approved", "confidence": 0.99, "reason_code": "REFERENCE_PARTIAL_MATCH",
                                   "requires_engineer": False}
    _seed(db_session, pid, folder, [_reply(f"{FA}-L2I-010021", "R0", "rejected", path="03- Drawings/SD/FA/R0/Received/reply.pdf",
                                           words="not approved")],
          sha={"03- Drawings/SD/FA/R0/Received/reply.pdf": "changed"})
    db_session.expire_all()
    shop_drawings.reconcile(db_session, db_session.get(Project, pid), ai=True)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L21-010021")
    assert row["latest_status"] == "approved" and row["cells"]["R0"]["confirmed"]
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()
    assert [i["kind"] for i in issues["ai_review"]] == ["ai_suggestion"]


def test_an_ai_outage_never_breaks_the_drawings(client, db_session, tmp_path, ai, monkeypatch):
    fake = ai({})
    fake.complete = lambda request: AiResponse(data=None, error="transport", error_detail="gone", model="fake")
    _admin(client)
    pid, folder = _project(client, tmp_path, "93062")
    _seed(db_session, pid, folder, [_doc("Level 22", "R0", "approved", ref=f"{FA}-L22-010022"),
                                    _reply(f"{FA}-L22-010O22", "R0", "ANN", path="03- Drawings/SD/FA/R0/Received/reply.pdf")])
    shop_drawings.reconcile(db_session, db_session.get(Project, pid), ai=True)
    row = _row(_log(client, pid, "FAS"), f"{FA}-L22-010022")
    assert row["latest_status"] == "approved"
    issues = client.get(f"/projects/{pid}/drawings/issues?system=FAS").json()
    assert [i["kind"] for i in issues["ai_review"]] == ["ai_review_required"]


def test_normal_rows_cost_no_ai_call(client, db_session, tmp_path, ai):
    fake = ai({})
    _admin(client)
    pid, folder = _project(client, tmp_path, "93063")
    _seed(db_session, pid, folder, [_doc("Level 23", "R0", "rejected", ref=f"{FA}-L23"), _doc("Level 23", "R1", "approved", ref=f"{FA}-L23", day=2),
                                    _reply(f"{FA}-L23", "R1", "approved", path="03- Drawings/SD/FA/R1/Received/reply.pdf")])
    shop_drawings.reconcile(db_session, db_session.get(Project, pid), ai=True)
    assert fake.requests == []
