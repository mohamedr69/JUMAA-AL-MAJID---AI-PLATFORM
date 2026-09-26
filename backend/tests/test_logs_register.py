"""Project > Logs > Drawings: the consolidated register, one row per logical
shop drawing, read from the same records the Drawings page edits -- so
the two cannot disagree about a floor, a reference, a revision's status
or what is the latest. A drawing schedule entry is evidence on the record,
never a row of its own; systems come from the project; a cable has no
drawings register.
"""
import io

from openpyxl import load_workbook

from .test_drawings_module import BOTH, EL, FA, _admin, _cells, _doc, _dxf, _import_ifc, _log, _project, _reconcile, _row, _seed


def _register(client, pid: int, **params) -> dict:
    query = "&".join(f"{k}={v}" for k, v in params.items())
    response = client.get(f"/projects/{pid}/logs/drawings" + (f"?{query}" if query else ""))
    assert response.status_code == 200, response.text
    return response.json()


def test_the_register_and_the_drawings_page_cannot_disagree(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "95001", BOTH)
    _import_ifc(client, pid, _dxf(tmp_path / "ifc.dxf", ["GROUND FLOOR", "1ST FLOOR", "2ND FLOOR"]))
    _seed(db_session, pid, folder, [
        _doc("Level 1", "R0", "rejected", ref=f"{FA}-L01", reply="Revise"), _doc("Level 1", "R1", "approved", ref=f"{FA}-L01", day=2),
        _doc("Ground Floor", "R0", "ANN", ref=f"{FA}-GF"),
        _doc("Level 1", "R0", "UR", ref=f"{EL}-L01", system="ELS"),
    ])
    _reconcile(client, pid)
    everything = _register(client, pid)
    assert everything["selected_system"] == "ALL"
    assert [(s["code"], s["drawings"], s["material_submittals"]) for s in everything["systems"]] == [("FAS", True, True), ("ELS", True, True)]
    pages = {code: _log(client, pid, code) for code in ("FAS", "ELS")}
    assert everything["pagination"]["total"] == len(pages["FAS"]["rows"]) + len(pages["ELS"]["rows"])
    assert everything["summary"]["total"] == everything["pagination"]["total"]
    # Every row with a drawing says exactly what the Drawings page says of it.
    compared = 0
    for row in everything["rows"]:
        if row["drawing_id"] is None:
            continue
        page_rows = [r for r in pages[row["system"]]["rows"] if r["id"] == row["drawing_id"]]
        page_row = page_rows[0]
        assert row["floor_keys"] == [k for r in page_rows for k in r["floor_keys"]]
        assert (row["reference"], row["latest_revision"], row["latest_status"]) == (
            page_row["reference"], page_row["latest_revision"], page_row["latest_status"])
        assert {rev: c["status"] for rev, c in row["cells"].items() if rev in page_row["cells"]} == _cells(page_row)
        assert [h["label"] for h in row["hints"]] == [h["label"] for h in page_row["hints"]]
        compared += 1
    assert compared == 3
    # One system: its rows alone, the register's own, and the system column implied.
    fas = _register(client, pid, system="FAS")
    assert fas["selected_system"] == "FAS" and {r["system"] for r in fas["rows"]} == {"FAS"}
    assert fas["pagination"]["total"] == len(pages["FAS"]["rows"]) and fas["folder"] == "03- Drawings/SD/FA"
    l01 = next(r for r in fas["rows"] if r["reference"] == f"{FA}-L01")
    assert {rev: c["status"] for rev, c in l01["cells"].items()} == {"R0": "not_approved", "R1": "approved", "R2": "not_submitted"}
    assert (l01["latest_revision"], l01["latest_status"]) == ("R1", "approved")
    # The record itself is the Drawings page's own details.
    record = client.get(f"/projects/{pid}/logs/drawings/{l01['drawing_id']}").json()
    assert record == client.get(f"/projects/{pid}/drawings/sd/{l01['drawing_id']}").json()
    assert [e["kind"] for e in record["evidence"]] == ["drawing", "reply", "drawing", "reply"]


def test_one_drawing_for_several_floors_is_one_row(client, db_session, tmp_path):
    """"Basement 4, 3, 2" is one issued drawing: the Drawings page lists it
    under each floor, the register once, its floors together."""
    _admin(client)
    pid, folder = _project(client, tmp_path, "95006")
    _seed(db_session, pid, folder, [_doc("Basement 4, 3, 2", "R0", "ANN", ref=f"{FA}-BSM-ZZZ-010001", name="BASEMENTS FLOOR PLAN")])
    _reconcile(client, pid)
    page = [r for r in _log(client, pid, "FAS")["rows"] if r["id"]]
    assert [r["floor_keys"] for r in page] == [["B4"], ["B3"], ["B2"]]
    (row,) = _register(client, pid)["rows"]
    assert (row["floor_keys"], row["floors"], row["floor"]) == (["B4", "B3", "B2"], 3, "Basement 4, 3, 2")
    assert (row["latest_revision"], row["latest_status"]) == ("R0", "approved_as_noted")
    assert _register(client, pid)["summary"]["total"] == 1
    assert [r["reference"] for r in _register(client, pid, floor="B3")["rows"]] == [f"{FA}-BSM-ZZZ-010001"]


def test_a_schedule_entry_is_evidence_on_the_record_not_a_row(client, db_session, tmp_path):
    """The project's shop drawings log plans FA 102 for basement 4; the
    drawing filed is BBY006-...-B04-010025 at R0 (not approved) and R1.
    One row: the drawing, its title as the title block reads, the
    schedule entry as evidence -- never "FA 102, R0, View schedule" beside it."""
    _admin(client)
    pid, folder = _project(client, tmp_path, "95002")
    reference = "BBY006-GME-SDW-FP-FA-BSM-B04-010025"
    _seed(db_session, pid, folder, [
        _doc("BASEMENT- 4", "R0", "UR", ref="FA 102", name="BASEMENT- 4 FLOOR PLAN", path="04- Drawings/Shop drawings log.pdf",
             source="drawing schedule"),
        _doc("BASEMENT-4", "R0", "rejected", ref=reference, name="BASEMENT-4 FLOOR PLAN FIRE ALARM LAYOUT", reply="Revise as marked"),
        _doc("BASEMENT-4", "R1", "UR", ref=reference, name="BASEMENT-4 FLOOR PLAN FIRE ALARM LAYOUT", day=2),
    ])
    _reconcile(client, pid)
    out = _register(client, pid)
    b4 = [r for r in out["rows"] if "B4" in r["floor_keys"]]
    assert len(b4) == 1
    (row,) = b4
    assert row["reference"] == reference and row["title"] == "Basement-4 Floor Plan Fire Alarm Layout"
    assert row["schedule"]["reference"] == "FA 102" and row["schedule"]["path"] == "04- Drawings/Shop drawings log.pdf"
    assert {rev: c["status"] for rev, c in row["cells"].items()} == {"R0": "not_approved", "R1": "under_review", "R2": "not_submitted"}
    assert (row["latest_revision"], row["latest_status"]) == ("R1", "under_review")
    assert not any(r["reference"] == "FA 102" for r in out["rows"])
    record = client.get(f"/projects/{pid}/logs/drawings/{row['drawing_id']}").json()
    assert [e["kind"] for e in record["evidence"]] == ["schedule", "drawing", "reply", "drawing"]
    assert record["evidence"][0]["label"] == "Drawing schedule: FA 102"
    # And the schedule entry is found by search, as the record's.
    assert [r["reference"] for r in _register(client, pid, search="FA 102")["rows"]] == [reference]


def test_search_filters_and_pages_read_the_records(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "95003")
    _seed(db_session, pid, folder, [
        _doc("Level 1", "R0", "approved", ref=f"{FA}-L01", name="LEVEL 1 FIRE ALARM LAYOUT"),
        _doc("Level 2", "R0", "UR", ref=f"{FA}-L02", name="LEVEL 2 FIRE ALARM LAYOUT"),
        _doc("Basement 4", "R0", "rejected", ref=f"{FA}-B04", name="BASEMENT 4 FIRE ALARM LAYOUT"),
        _doc("Basement 4", "R2", "UR", ref=f"{FA}-B04", name="BASEMENT 4 FIRE ALARM LAYOUT", day=3),   # R1 missing: a gap
    ])
    _reconcile(client, pid)
    assert [r["reference"] for r in _register(client, pid, search="FA-L01")["rows"]] == [f"{FA}-L01"]
    assert [r["reference"] for r in _register(client, pid, search="Basement 4")["rows"]] == [f"{FA}-B04"]
    assert [r["reference"] for r in _register(client, pid, search="B4")["rows"]] == [f"{FA}-B04"]
    assert [r["reference"] for r in _register(client, pid, status="approved")["rows"]] == [f"{FA}-L01"]
    assert [r["reference"] for r in _register(client, pid, revision="R2")["rows"]] == [f"{FA}-B04"]
    gapped = _register(client, pid, issues_only="true")["rows"]
    assert [r["reference"] for r in gapped] == [f"{FA}-B04"] and "revision_gap" in [h["kind"] for h in gapped[0]["hints"]]
    assert _register(client, pid, floor="L2")["rows"][0]["reference"] == f"{FA}-L02"
    paged = _register(client, pid, page_size="7", page="9")
    assert paged["pagination"] == {"page": 1, "page_size": 25, "pages": 1, "total": 3}
    assert paged["available_revisions"] == ["R0", "R1", "R2"]
    assert paged["summary"]["total"] == 3 and paged["summary"]["review_items"] == 1


def test_only_a_drawn_system_has_a_drawings_register(client, db_session, tmp_path):
    """A full-package project has fire-rated cables: a material submittal
    and a sample, never a shop drawing. ALL covers the drawn systems; the
    cable's register is refused, not answered with another's rows."""
    _admin(client)
    folder = tmp_path / "EP-95004"
    folder.mkdir()
    pid = client.post("/projects", json={"ep_number": "95004", "project_name": "Cables", "design_sheets": [],
                                         "source_folder_path": str(folder), "scope_of_work": "Full Package",
                                         "systems": BOTH}).json()["id"]
    systems = client.get(f"/projects/{pid}/logs/systems").json()
    assert [(s["code"], s["drawings"], s["material_submittals"], s["samples"]) for s in systems["systems"]] == [
        ("FAS", True, True, True), ("ELS", True, True, True), ("FRC", False, True, True)]
    assert client.get(f"/projects/{pid}/logs/drawings?system=FRC").status_code == 404
    assert [s["code"] for s in _register(client, pid)["systems"] if s["drawings"]] == ["FAS", "ELS"]


def test_the_export_has_the_register_and_the_revision_details(client, db_session, tmp_path):
    _admin(client)
    pid, folder = _project(client, tmp_path, "95005", BOTH)
    _seed(db_session, pid, folder, [
        _doc("Level 1", "R0", "rejected", ref=f"{FA}-L01"), _doc("Level 1", "R1", "ANN", ref=f"{FA}-L01", day=2),
        _doc("Level 1", "R0", "approved", ref=f"{EL}-L01", system="ELS"),
    ])
    _reconcile(client, pid)
    everything = client.get(f"/projects/{pid}/logs/drawings/export.xlsx")
    assert everything.status_code == 200 and "Drawings Register All.xlsx" in everything.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(everything.content))
    assert wb.sheetnames == ["Drawings Register All", "Revision Details"]
    rows = list(wb["Drawings Register All"].iter_rows(min_row=4, values_only=True))
    assert rows[0][:5] == ("#", "System", "Floor", "Drawing Title", "Drawing Reference")
    assert sorted((r[1], r[4]) for r in rows[1:]) == [("ELS", f"{EL}-L01"), ("FAS", f"{FA}-L01")]
    details = list(wb["Revision Details"].iter_rows(min_row=1, values_only=True))
    assert details[0][:5] == ("System", "Floor", "Drawing Reference", "Revision", "Status")
    assert sorted((d[0], d[3], d[4]) for d in details[1:]) == [("ELS", "R0", "Approved"), ("FAS", "R0", "Not Approved"),
                                                               ("FAS", "R1", "Approved as Noted")]
    one = client.get(f"/projects/{pid}/logs/drawings/export.xlsx?system=ELS")
    assert "Drawings Register ELS.xlsx" in one.headers["content-disposition"]
    assert [r[1] for r in load_workbook(io.BytesIO(one.content))["Drawings Register ELS"].iter_rows(min_row=5, values_only=True)] == ["ELS"]
