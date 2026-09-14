from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import RoleEnum
from app.services.boq_export import parse_quantity

from .conftest import login, make_user
from .test_projects import _login_admin, _valid_project_payload


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("12", Decimal(12)),
        (" 7 ", Decimal(7)),
        ("1,200", Decimal(1200)),
        ("2.5", Decimal("2.5")),
        ("Lot", None),
        ("", None),
        (None, None),
        ("1.2.3", None),
        ("-3", None),
    ],
)
def test_parse_quantity(text, expected):
    assert parse_quantity(text) == expected


def _lines() -> list[dict]:
    return [
        {"system_code": "FAS", "group_heading": "Main Panel", "manufacturer": "EDWARDS",
         "catalog_no": "4-CPU", "description": "Central Processor Module", "quantity": "2",
         "unit": "Nos", "unit_price": "250.00", "total_price": "500.00"},
        {"system_code": "FAS", "catalog_no": "SIGA-PS", "description": "Photo Smoke Detector",
         "quantity": "1,200", "remarks": "Includes spares"},
        {"system_code": "FAS", "description": "Commissioning", "quantity": "Lot"},
        # Emergency lighting (ELS, CBS, EML) is one system: stored as ELS.
        {"system_code": "ELS", "catalog_no": "VG-B", "description": "VisionGuard Basisversion",
         "quantity": "1"},
        {"system_code": None, "description": "Loose item", "quantity": "3"},
    ]


def _export(client, pid):
    resp = client.get(f"/projects/{pid}/boq/export.xlsx")
    assert resp.status_code == 200
    return resp, load_workbook(BytesIO(resp.content))


def test_export_has_a_summary_then_one_sheet_per_system_in_tab_order(client):
    _login_admin(client)
    pid = client.post("/projects", json=_valid_project_payload("36000")).json()["id"]
    client.put(f"/projects/{pid}/boq", json=_lines())

    resp, wb = _export(client, pid)
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="EP-36000 BOQ.xlsx"' in resp.headers["content-disposition"]
    # the project's Design Sheet (FAS) first, then other systems, then none
    assert wb.sheetnames == ["Summary", "FAS", "ELS", "Unassigned"]


def test_export_system_sheet_lines_and_totals(client):
    _login_admin(client)
    pid = client.post("/projects", json=_valid_project_payload("36100")).json()["id"]
    client.put(f"/projects/{pid}/boq", json=_lines())

    _, wb = _export(client, pid)
    fas = wb["FAS"]
    assert fas["A1"].value == "EP-36100 — IVY Garden 2"
    assert [c.value for c in fas[5]] == [
        "#", "Group", "Manufacturer", "Model / Part No.", "Description",
        "Qty", "Unit", "Unit Price", "Total Price", "Remarks",
    ]
    assert [c.value for c in fas[6]] == [
        1, "Main Panel", "EDWARDS", "4-CPU", "Central Processor Module",
        2, "Nos", 250, 500, None,
    ]
    # numeric quantities are numbers Excel can sum -- separators and all --
    # while "Lot" stays as the sheet wrote it
    assert fas["F7"].value == 1200
    assert fas["F8"].value == "Lot"

    # live formulas, so the total follows edits made in Excel
    assert fas["E9"].value == "Total"
    assert fas["F9"].value == "=SUM(F6:F8)"
    assert fas["I9"].value == "=SUM(I6:I8)"
    assert fas["E10"].value == "Not in the quantity total (quoted as a word): 1 × Lot"


def test_export_summary_totals(client):
    _login_admin(client)
    pid = client.post("/projects", json=_valid_project_payload("36200")).json()["id"]
    client.put(f"/projects/{pid}/boq", json=_lines())

    _, wb = _export(client, pid)
    summary = wb["Summary"]
    rows = {
        row[0]: row[1:]
        for row in summary.iter_rows(min_row=1, values_only=True)
        if row[0] in ("FAS", "ELS", "Unassigned", "Client")
    }
    assert rows["Client"][0] == "Samana"
    # system: lines, total qty, quoted as a word, total price
    assert rows["FAS"][:4] == (3, 1202, "1 × Lot", 500)
    assert rows["ELS"][:4] == (1, 1, None, None)
    assert rows["Unassigned"][:4] == (1, 3, None, None)


def test_export_of_an_empty_boq_still_opens(client):
    _login_admin(client)
    pid = client.post("/projects", json=_valid_project_payload("36300")).json()["id"]

    _, wb = _export(client, pid)
    assert wb.sheetnames == ["Summary"]
    assert any(
        cell == "The BOQ has no lines." for row in wb["Summary"].iter_rows(values_only=True) for cell in row
    )


def test_export_open_to_viewers_and_404_for_missing(client, db_session):
    _login_admin(client)
    pid = client.post("/projects", json=_valid_project_payload("36400")).json()["id"]
    client.put(f"/projects/{pid}/boq", json=_lines())
    assert client.get("/projects/999999/boq/export.xlsx").status_code == 404
    client.post("/auth/logout")

    make_user(db_session, "viewer6@ep-platform.com", RoleEnum.viewer)
    login(client, "viewer6@ep-platform.com")
    assert client.get(f"/projects/{pid}/boq/export.xlsx").status_code == 200
