"""A real floor-wise schedule, read as the engineer wrote it.

This is the Binghatti Titania schedule (EP-30880): a fire alarm block and
an emergency lighting block in one sheet, the description column headed
DRG.NAME, a `Structural slab` among the floors, ranges written "3rd to
16th floor", and the sheet's own "No. of Floors" row stating what each
column stands for.

It is here because every one of those broke the reader when it first met
a real sheet: the serial column was taken for the description, so every
item read "A", "2", "3"; the structural slab was not taken for a floor,
so 28 heat detectors were dropped; and the floors came out sorted into an
order the sheet never had.
"""

from pathlib import Path

import openpyxl

from app.services import floor_schedule

HEADINGS = [
    "S.NO", "DRG.NAME",
    "3rd Basement", "2nd Basement", "1st Basement", "Ground floor",
    "1st Podium", "2nd Podium", "3rd Podium", "4th Podium",
    "1st Floor", "2nd Floor", "3rd to 16th floor", "Mech floor", "Structural slab",
    "17th to 22nd floor", "23rd to 27th floor", "28th to 32nd Floor", "Roof Floor",
    "TOTAL",
]

# (S.No, name, the floor cells left to right, the sheet's own total)
ROWS: list[tuple] = [
    ("A", "No. of Floors", [1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 14, 1, 1, 6, 5, 5, 1], 43),
    (2, "Smoke Detector", [4, 4, 4, 23, 8, 4, 4, 4, 38, 30, 476, 4, 2, 204, 170, 170, 14], 1163),
    (3, "Smoke with Sounder Base", [None] * 8 + [19, 22, 364, None, None, 156, 130, 130, None], 821),
    (4, "Heat Detector", [None, None, None, 7, 2, None, None, 2, 4, 4, 84, 19, 28, 36, 30, 30, 5], 251),
    (5, "Multisensor", [None, None, None, 1, None, None, None, None, 14, 17, 266, None, None, 114, 95, 95, None], 602),
    (6, "Manual Pull Station", [None] * 8 + [2, 2, 28, 2, None, 12, 10, 10, None], 66),
    (7, "Manual Pull Station WP", [2, 2, 2, 3, 2, 2, 2, 2] + [None] * 8 + [1], 18),
    (8, "Telephone Jack", [2, 2, 2, 4, 2, 2, 2, 2, 2, 2, 28, None, None, 12, 10, 10, 2], 86),
    (9, "Telephone Jack-Lift", [None, None, None, 9] + [None] * 13, 9),
    (10, "Wall Speaker", [2, 2, 2, 4, 2, 2, 2, 2, 2, 2, 28, None, None, 12, 10, 10, 1], 85),
    (11, "Ceiling Speaker", [3, 3, 3, 12, 6, 3, 3, 3, 19, 11, 168, 3, None, 72, 60, 60, None], 429),
    (12, "Wall Sounder", [None] * 8 + [2, 2, 28, 2, None, 12, 10, 10, None], 66),
    (13, "Wall Sounder Flasher WP", [14, 14, 14, 4, 10, 10, 10, 10, 3] + [None] * 7 + [1], 90),
    (14, "MM for FM200", [None, None, None, 1] + [None] * 13, 1),
    (15, "Control Module", [None] * 16 + [9], 9),
    (16, "Emergency Luminaire - Surface Mounted (Driveway)", [63, 63, 63, 42, 36, 35, 36, 36] + [None] * 9, 374),
    (17, "Side wall Mounted emergency", [None] * 8 + [3] + [None] * 8, 3),
    (18, "Surface Emergency(MEP Rm)", [None] * 13 + [19, 28, None, None], 47),
    (19, "Emergency Luminaire - Surface Mounted (Service Room)",
     [5, 5, 5, 21, 8, 5, 5, 7, 10, 10, 140, 6, 6, 60, 50, 50, 11], 404),
    (20, "Emergency Luminaire - Recess Mounted", [3, 3, 3, 13, 6, 3, 3, 3, 19, 11, 168, 2, None, 72, 60, 60, None], 429),
    (21, "Exit Light - Above Door", [2, 2, 2, 11, 2, 2, 2, 2, 5, 4, 56, 3, None, 24, 25, 25, None], 167),
    (22, "Exit Light - Above Door WP", [2, 2, 2, 5, 4, 2, 2, 3, 3, None, None, 3, None, None, None, None, 5], 33),
    (23, "Directional Exit Light (Corridor)", [None, None, None, 3, None, None, None, None, 4, 4, 70,
                                               None, None, 30, 25, 25, None], 161),
    (24, "Directional Exit Light (Driveway)", [5, 5, 5, None, 4, 7, 4, 4] + [None, None, None, 2] + [None] * 5, 36),
    (25, "Self Contained EM Light", [None, None, None, 6] + [None] * 13, 6),
]

# The row whose quantities are not numbers at all: the sheet writes the
# arithmetic across a merged cell ("7*2=14(Stair), 7*9=63(Lift)") and
# totals it 77. It is here because such a row must still reach the page.
NOTE_ROW = (1, "Smoke for every 23 meters", "7*2=14(Stair), 7*9=63(Lift)", 77)

# The sheet's own grand total is 5433. Two rows -- Telephone Jack and Wall
# Speaker -- could not be read off the screenshot cell by cell (each is
# two short of the total printed beside it), so the workbook built here
# totals each row from its own cells, as the real sheet does. What is
# asserted against the screenshot is every row that was transcribed
# exactly, which is all but those two.
SCREENSHOT_TOTALS: dict[str, int] = {
    "Smoke Detector": 1163,
    "Smoke with Sounder Base": 821,
    "Heat Detector": 251,
    "Multisensor": 602,
    "Manual Pull Station": 66,
    "Manual Pull Station WP": 18,
    "Telephone Jack-Lift": 9,
    "Ceiling Speaker": 429,
    "Wall Sounder": 66,
    "Wall Sounder Flasher WP": 90,
    "MM for FM200": 1,
    "Control Module": 9,
    "Emergency Luminaire - Surface Mounted (Driveway)": 374,
    "Side wall Mounted emergency": 3,
    "Surface Emergency(MEP Rm)": 47,
    "Emergency Luminaire - Surface Mounted (Service Room)": 404,
    "Emergency Luminaire - Recess Mounted": 429,
    "Exit Light - Above Door": 167,
    "Exit Light - Above Door WP": 33,
    "Directional Exit Light (Corridor)": 161,
    "Directional Exit Light (Driveway)": 36,
    "Self Contained EM Light": 6,
}


def _row_total(cells: list) -> int:
    return sum(value for value in cells if isinstance(value, (int, float)))


def _titania(tmp_path: Path) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "FLOOR WISE"
    sheet.cell(row=1, column=2, value="Binghatti Titania")
    for index, heading in enumerate(HEADINGS, start=1):
        sheet.cell(row=3, column=index, value=heading)
    for offset, (serial, name, cells, total) in enumerate(ROWS, start=4):
        sheet.cell(row=offset, column=1, value=serial)
        sheet.cell(row=offset, column=2, value=name)
        for index, value in enumerate(cells, start=3):
            if value is not None:
                sheet.cell(row=offset, column=index, value=value)
        # The sheet totals each row from its own cells.
        sheet.cell(row=offset, column=len(HEADINGS),
                   value=total if name == "No. of Floors" else _row_total(cells))
    # ...and the one row whose quantities are written as a sum across a
    # merged cell rather than as numbers under the floors.
    note_row = len(ROWS) + 4
    serial, name, text, total = NOTE_ROW
    sheet.cell(row=note_row, column=1, value=serial)
    sheet.cell(row=note_row, column=2, value=name)
    sheet.merge_cells(start_row=note_row, start_column=9, end_row=note_row, end_column=13)
    sheet.cell(row=note_row, column=9, value=text)
    sheet.cell(row=note_row, column=len(HEADINGS), value=total)
    path = tmp_path / "EP-30880 FLOOR WISE BOQ.xlsx"
    book.save(path)
    return path


def test_the_description_column_is_the_one_carrying_the_names(tmp_path):
    """Headed DRG.NAME, with a serial column beside it. Taking the serial
    column for the description made every item read "A", "2", "3"."""
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    names = [item["description"] for item in read["items"]]
    assert "Smoke Detector" in names and "Self Contained EM Light" in names
    assert not any(name.strip().isdigit() or name.strip() == "A" for name in names)


def test_the_sheets_own_floor_order_is_kept(tmp_path):
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    assert read["floors"][:10] == [
        "3rd Basement", "2nd Basement", "1st Basement", "Ground floor",
        "1st Podium", "2nd Podium", "3rd Podium", "4th Podium", "1st Floor", "2nd Floor",
    ]
    # The ranges are written out where they stand, and the oddities between
    # them are floors too.
    assert "Level 16" in read["floors"] and "Mech floor" in read["floors"]
    assert "Structural slab" in read["floors"]
    assert read["floors"][-1] == "Roof Floor"
    assert read["floors"].index("Mech floor") > read["floors"].index("Level 16")
    # 8 named + 2 + 14 + Mech + slab + 6 + 5 + 5 + roof
    assert len(read["floors"]) == 43


def test_a_structural_slab_is_a_floor_devices_are_counted_on(tmp_path):
    """It reads nothing like a storey and it carries 28 heat detectors.
    Left out, they were dropped and the row came to 223 instead of 251."""
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    heat = next(item for item in read["items"] if item["description"] == "Heat Detector")
    assert heat["per_floor"]["Structural slab"] == 28
    assert heat["total"] == 251


def test_every_row_adds_up_to_what_the_sheet_totals_it(tmp_path):
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    by_name = {item["description"]: item for item in read["items"]}
    for _serial, name, cells, _total in ROWS:
        if name == "No. of Floors":
            continue
        assert by_name[name]["total"] == _row_total(cells), name
    # ...and those transcribed exactly off the schedule come to what it
    # prints beside them.
    for name, total in SCREENSHOT_TOTALS.items():
        assert by_name[name]["total"] == total, f"{name}: {by_name[name]['total']} != {total}"


def test_a_row_written_as_a_sum_across_a_merged_cell_still_reaches_the_page(tmp_path):
    """"Smoke for every 23 meters" carries "7*2=14(Stair), 7*9=63(Lift)"
    across a merged cell and is totalled 77. There is no quantity per
    floor to read, and dropping the line would take 77 devices out of the
    BOQ without saying so."""
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    note = next(item for item in read["items"] if item["description"] == "Smoke for every 23 meters")
    assert note["per_floor"] == {} and note["stated_total"] == 77
    assert note["device"] == "Smoke detector" and note["system"] == "FAS"
    assert any("could not be read as numbers" in warning for warning in read["warnings"])
    # The two totals differ by exactly what could not be counted per floor.
    assert read["stated_grand_total"] - read["grand_total"] == 77


def test_the_floor_count_row_is_a_check_not_an_item(tmp_path):
    """The sheet states how many floors each column stands for. Counted as
    an item it adds a phantom device to every floor; read as a check it
    confirms "3rd to 16th floor" is the fourteen the platform made of it."""
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    assert "No. of Floors" not in [item["description"] for item in read["items"]]
    # The sheet and the reader agree, so nothing is said about it.
    assert not any("floor" in warning and "was read as" in warning for warning in read["warnings"])


def test_a_column_the_sheet_counts_differently_is_reported(tmp_path):
    path = _titania(tmp_path)
    book = openpyxl.load_workbook(path)
    # The sheet now claims the 3rd-to-16th column is nine floors.
    book["FLOOR WISE"].cell(row=4, column=13, value=9)
    book.save(path)

    read = floor_schedule.read(path).as_dict()
    assert any("is 9 floors; it was read as 14" in warning for warning in read["warnings"])


def test_every_line_names_its_device_and_its_system(tmp_path):
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    named = {item["description"]: (item["device"], item["system"]) for item in read["items"]}

    assert named["Smoke Detector"] == ("Smoke detector", "FAS")
    assert named["Smoke with Sounder Base"] == ("Smoke detector with sounder base", "FAS")
    assert named["Multisensor"] == ("Multisensor detector", "FAS")
    assert named["Manual Pull Station WP"] == ("Manual call point", "FAS")
    assert named["Telephone Jack-Lift"] == ("Fire telephone", "FAS")
    assert named["MM for FM200"] == ("Monitor module", "FAS")
    # ...and the emergency lighting block reads as its own system.
    assert named["Emergency Luminaire - Recess Mounted"] == ("Emergency light", "ELS")
    assert named["Exit Light - Above Door"] == ("Exit light", "ELS")
    assert named["Directional Exit Light (Driveway)"] == ("Exit light", "ELS")
    assert named["Self Contained EM Light"] == ("Emergency light", "ELS")
    assert named["Surface Emergency(MEP Rm)"] == ("Emergency light", "ELS")

    # Every line on this sheet is named; none falls between the systems.
    assert all(device for device, _ in named.values())
    assert all(system in ("FAS", "ELS") for _, system in named.values())


def test_the_two_systems_account_for_every_device_on_the_sheet(tmp_path):
    """The tab shows a system at a time, so what the systems come to must
    be what the whole sheet comes to -- otherwise a line is on screen
    somewhere and missing from the arithmetic, or on neither tab at all."""
    read = floor_schedule.read(_titania(tmp_path)).as_dict()
    assert set(read["systems"]) == {"FAS", "ELS"}
    assert sum(read["systems"].values()) == read["grand_total"]
    assert read["grand_total"] == sum(item["total"] for item in read["items"])
