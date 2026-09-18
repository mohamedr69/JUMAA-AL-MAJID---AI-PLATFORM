"""The floor-wise BOQ read off the schedule an engineer keeps in Excel.

The tab's whole point is the typical column: a schedule writes "1 to 13"
once and the platform owes the engineer the thirteen floors. These tests
hold that, and the thing that makes it dangerous -- a typical column's
quantity is per floor, which is thirteen times the alternative, so it is
checked against the sheet's own total rather than assumed.
"""

from pathlib import Path

import openpyxl

from app.core.config import get_settings
from app.models import Project, ProjectBoqItem, ProjectFloorSchedule
from app.services import floor_schedule

from .conftest import login

settings = get_settings()


def _login(client):
    return login(client, settings.default_admin_email, settings.default_admin_password)


def _project(client, ep_number: str = "30880") -> int:
    return client.post("/projects", json={"ep_number": ep_number, "project_name": "Titania", "design_sheets": []}).json()["id"]


def _workbook(tmp_path: Path, rows: list[list], headings: list[str], name: str = "schedule.xlsx",
              title: str | None = None) -> Path:
    """A floor-wise schedule: a title row, a header row, then the items."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "FLOOR WISE"
    start = 1
    if title:
        sheet.cell(row=1, column=1, value=title)
        start = 3
    for index, heading in enumerate(headings, start=1):
        sheet.cell(row=start, column=index, value=heading)
    for offset, row in enumerate(rows, start=start + 1):
        for index, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=offset, column=index, value=value)
    path = tmp_path / name
    book.save(path)
    return path


def _upload(client, project_id: int, path: Path):
    return client.post(
        f"/projects/{project_id}/floor-schedule",
        files={"file": (path.name, path.read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


# --- a typical column is not one floor ----------------------------------------------------------


def test_a_column_for_a_range_of_floors_becomes_a_row_per_floor(tmp_path):
    path = _workbook(
        tmp_path,
        headings=["S.No", "Catalogue No", "Description", "Unit", "B1", "GF", "1 to 13", "ROOF"],
        rows=[
            [1, "SIGA-PS", "Photoelectric smoke detector", "Nos", 4, 6, 5, 2],
            [2, "SIGA-278", "Manual call point", "Nos", 2, 2, 1, None],
        ],
        title="BINGHATTI TITANIA - FIRE ALARM FLOOR WISE BOQ",
    )
    read = floor_schedule.read(path).as_dict()

    # The thirteen typical floors are shown one by one, in building order.
    assert read["floors"] == ["B1", "GF", "Level 1", "Level 2", "Level 3", "Level 4", "Level 5", "Level 6",
                              "Level 7", "Level 8", "Level 9", "Level 10", "Level 11", "Level 12", "Level 13",
                              "ROOF"]
    smoke = read["items"][0]
    assert smoke["catalog_no"] == "SIGA-PS" and smoke["unit"] == "Nos"
    assert smoke["per_floor"]["B1"] == 4 and smoke["per_floor"]["ROOF"] == 2
    # Five on each of the thirteen typical floors.
    assert smoke["per_floor"]["Level 1"] == 5 and smoke["per_floor"]["Level 13"] == 5
    assert smoke["total"] == 4 + 6 + (5 * 13) + 2


def test_a_blank_floor_is_not_a_line_of_zero(tmp_path):
    path = _workbook(
        tmp_path,
        headings=["Description", "B1", "GF", "ROOF"],
        rows=[["Manual call point", 2, None, 0]],
    )
    item = floor_schedule.read(path).as_dict()["items"][0]
    # The item is simply not on those floors.
    assert item["per_floor"] == {"B1": 2}
    assert "GF" not in item["per_floor"] and "ROOF" not in item["per_floor"]


def test_the_floors_follow_the_sheets_own_order(tmp_path):
    """A schedule is written deepest basement first and roof last, and that
    order is the engineer's own.

    Reordering it by what the platform makes of each heading only ever
    loses to the sheet: "3rd Basement" and "1st Podium" are plain to a
    reader and awkward to a regular expression, and sorting them put the
    basements upside down and the podiums among the levels. A range column
    takes its place where it stands.
    """
    path = _workbook(
        tmp_path,
        headings=["Description", "3rd Basement", "2nd Basement", "1st Basement", "Ground floor",
                  "1st Podium", "1st Floor", "3rd to 6th floor", "Roof Floor"],
        rows=[["Smoke detector", 1, 1, 1, 1, 1, 1, 4, 1]],
    )
    assert floor_schedule.read(path).as_dict()["floors"] == [
        "3rd Basement", "2nd Basement", "1st Basement", "Ground floor", "1st Podium", "1st Floor",
        "Level 3", "Level 4", "Level 5", "Level 6", "Roof Floor",
    ]


def test_a_merged_banner_over_the_floor_columns_is_read(tmp_path):
    """A schedule usually heads its typical columns with a merged banner
    and the floors under it."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "BOQ"
    sheet.cell(row=1, column=1, value="Description")
    sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
    sheet.cell(row=1, column=2, value="TYPICAL FLOOR")
    sheet.cell(row=2, column=2, value="1 to 5")
    sheet.cell(row=2, column=3, value="ROOF")
    sheet.cell(row=3, column=1, value="Smoke detector")
    sheet.cell(row=3, column=2, value=3)
    sheet.cell(row=3, column=3, value=1)
    path = tmp_path / "merged.xlsx"
    book.save(path)

    read = floor_schedule.read(path).as_dict()
    assert read["floors"] == ["Level 1", "Level 2", "Level 3", "Level 4", "Level 5", "ROOF"]
    assert read["items"][0]["total"] == 3 * 5 + 1


# --- what a typical quantity means ---------------------------------------------------------------


def test_the_sheets_own_total_settles_that_a_typical_quantity_is_per_floor(tmp_path):
    path = _workbook(
        tmp_path,
        headings=["Description", "GF", "1 to 10", "Total"],
        rows=[["Smoke detector", 4, 5, 54]],          # 4 + 5x10
    )
    read = floor_schedule.read(path).as_dict()
    assert read["typical_reading"] == "per_floor"
    assert read["items"][0]["total"] == 54
    assert "10 of 1 rows" not in read["typical_reason"]


def test_a_sheet_that_totals_across_the_range_is_read_that_way(tmp_path):
    """Some schedules write the whole job's quantity against the typical
    column rather than one floor's. The sheet's own total says which, and
    it is believed over the convention."""
    path = _workbook(
        tmp_path,
        headings=["Description", "GF", "1 to 10", "Total"],
        rows=[["Smoke detector", 4, 50, 54]],         # 4 + 50 shared across ten
    )
    read = floor_schedule.read(path).as_dict()
    assert read["typical_reading"] == "across"
    assert read["items"][0]["per_floor"]["Level 1"] == 5
    assert read["items"][0]["total"] == 54


def test_without_a_total_the_per_floor_reading_stands_and_says_so(tmp_path):
    path = _workbook(
        tmp_path,
        headings=["Description", "GF", "1 to 13"],
        rows=[["Smoke detector", 4, 5]],
    )
    read = floor_schedule.read(path).as_dict()
    assert read["typical_reading"] == "per_floor"
    assert "nothing confirms it" in read["typical_reason"]
    # The engineer is told the reading rests on a convention.
    assert any("counted on every floor" in warning for warning in read["warnings"])


def test_a_row_that_does_not_add_up_to_its_stated_total_is_reported(tmp_path):
    path = _workbook(
        tmp_path,
        headings=["Description", "GF", "1 to 10", "Total"],
        rows=[["Smoke detector", 4, 5, 54], ["Heat detector", 1, 1, 99]],
    )
    read = floor_schedule.read(path).as_dict()
    assert any("do not add up to the total" in warning for warning in read["warnings"])


# --- through the API ------------------------------------------------------------------------------


def test_the_workbook_is_read_once_and_the_tab_reads_it_from_the_database(client, db_session, tmp_path):
    _login(client)
    project_id = _project(client)
    assert client.get(f"/projects/{project_id}/floor-schedule").json()["result"] is None

    path = _workbook(
        tmp_path,
        headings=["Catalogue No", "Description", "Unit", "B1", "GF", "1 to 13", "ROOF"],
        rows=[["SIGA-PS", "Photoelectric smoke detector", "Nos", 4, 6, 5, 2]],
        name="EP-30880 FLOOR WISE.xlsx",
    )
    response = _upload(client, project_id, path)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["file_name"] == "EP-30880 FLOOR WISE.xlsx" and body["sheet_name"] == "FLOOR WISE"
    assert len(body["result"]["floors"]) == 16
    assert body["result"]["grand_total"] == 4 + 6 + (5 * 13) + 2

    stored = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert stored["result"] == body["result"]
    assert db_session.query(ProjectFloorSchedule).filter(
        ProjectFloorSchedule.project_id == project_id).count() == 1

    assert client.delete(f"/projects/{project_id}/floor-schedule").status_code == 204
    assert client.get(f"/projects/{project_id}/floor-schedule").json()["result"] is None


def test_a_file_that_is_not_a_schedule_says_so_rather_than_failing(client, tmp_path):
    _login(client)
    project_id = _project(client, "30881")
    book = openpyxl.Workbook()
    book.active["A1"] = "Nothing of interest"
    path = tmp_path / "empty.xlsx"
    book.save(path)

    response = _upload(client, project_id, path)
    assert response.status_code == 422
    assert "floor" in response.json()["detail"].lower()


def test_only_an_excel_workbook_is_accepted(client, tmp_path):
    _login(client)
    project_id = _project(client, "30882")
    path = tmp_path / "drawing.dxf"
    path.write_text("not a workbook")
    assert _upload(client, project_id, path).status_code == 400


def test_the_schedule_is_checked_against_the_design_sheet_boq(client, db_session, tmp_path):
    """The two are read from different documents and should agree; where
    they do not, the difference is the question."""
    _login(client)
    project_id = _project(client, "30883")
    db_session.add(ProjectBoqItem(project_id=project_id, position=0, catalog_no="SIGA-PS",
                                  description="Photoelectric smoke detector", quantity="72",
                                  manufacturer="EDWARDS", system_code="FAS", unit="Nos"))
    db_session.add(ProjectBoqItem(project_id=project_id, position=1, catalog_no="SIGA-HFS",
                                  description="Heat detector", quantity="10",
                                  manufacturer="EDWARDS", system_code="FAS", unit="Nos"))
    db_session.commit()

    path = _workbook(
        tmp_path,
        headings=["Catalogue No", "Description", "B1", "GF", "1 to 13"],
        rows=[["SIGA-PS", "Photoelectric smoke detector", 4, 6, 5],   # 75
              ["SIGA-UM", "Universal module", 1, 1, None]],           # not in the BOQ
    )
    _upload(client, project_id, path)

    check = client.get(f"/projects/{project_id}/floor-schedule/check")
    assert check.status_code == 200, check.text
    body = check.json()
    assert body["matched"] == 1
    smoke = body["rows"][0]
    assert smoke["schedule_total"] == 75 and smoke["boq_quantity"] == 72
    assert smoke["difference"] == 3
    assert [entry["catalog_no"] for entry in body["only_in_schedule"]] == ["SIGA-UM"]
    assert [entry["catalog_no"] for entry in body["only_in_boq"]] == ["SIGA-HFS"]


# --- what each line is ----------------------------------------------------------------------------


def test_every_line_says_what_device_it_is(tmp_path):
    """A schedule is a list of devices; the platform reads which, the same
    way it reads a symbol's legend on a drawing."""
    path = _workbook(
        tmp_path,
        headings=["Catalogue No", "Description", "GF"],
        rows=[["SIGA-PS", "Photoelectric smoke detector", 4],
              ["SIGA-HFS", "Heat detector, rate of rise", 2],
              ["SIGA-278", "Intelligent manual call point", 1],
              ["G4S-C", "Ceiling speaker 6W", 3]],
    )
    read = floor_schedule.read(path).as_dict()
    assert [item["device"] for item in read["items"]] == [
        "Smoke detector", "Heat detector", "Manual call point", "Speaker",
    ]
    # ...and which family it is ordered under, as the drawings tab groups.
    assert [item["family"] for item in read["items"]] == [
        "Detectors", "Detectors", "Pull station", "Speaker",
    ]
    assert read["devices"] == ["Heat detector", "Manual call point", "Smoke detector", "Speaker"]


def test_a_line_naming_no_known_device_is_read_all_the_same(tmp_path):
    """The quantities are what the tab is for; an unrecognised wording
    leaves the device blank rather than dropping the line or guessing."""
    path = _workbook(
        tmp_path,
        headings=["Description", "GF"],
        rows=[["Cable tray accessories", 12]],
    )
    read = floor_schedule.read(path).as_dict()
    assert read["items"][0]["device"] is None
    assert read["items"][0]["per_floor"]["GF"] == 12
    assert any("1 line names no device the platform knows" in warning for warning in read["warnings"])


def test_the_device_is_read_from_the_catalogue_number_when_the_wording_does_not_say(tmp_path):
    path = _workbook(tmp_path, headings=["Catalogue No", "Description", "GF"],
                     rows=[["SIGA-PS", "Ceiling mounted, addressable, 24V", 4]])
    assert floor_schedule.read(path).as_dict()["items"][0]["device"] == "Smoke detector"


# --- the workbook is kept in the project's folder --------------------------------------------------


def test_the_workbook_is_filed_in_the_projects_design_folder(client, db_session, tmp_path):
    """The schedule an engineer works from belongs with the project's own
    documents, so it is kept under 03- Design beside the submittals."""
    _login(client)
    project_id = _project(client, "30884")
    folder = tmp_path / "EP-30884 TITANIA"
    folder.mkdir()
    project = db_session.get(Project, project_id)
    project.source_folder_path = str(folder)
    db_session.commit()

    path = _workbook(tmp_path, headings=["Description", "GF", "1 to 3"],
                     rows=[["Photoelectric smoke detector", 4, 5]],
                     name="EP-30884 FLOOR WISE BOQ.xlsx")
    body = _upload(client, project_id, path).json()

    assert body["archive_path"] == "03- Design/EP-30884 FLOOR WISE BOQ.xlsx"
    assert body["filed_note"] is None
    filed = folder / "03- Design" / "EP-30884 FLOOR WISE BOQ.xlsx"
    assert filed.is_file() and filed.read_bytes() == path.read_bytes()


def test_a_project_whose_folder_is_not_on_this_pc_is_still_read(client, tmp_path):
    """The archive is wherever OneDrive syncs it; on a PC without this
    project's folder the schedule is read and shown, and the page says why
    it was not filed."""
    _login(client)
    project_id = _project(client, "30885")
    path = _workbook(tmp_path, headings=["Description", "GF"], rows=[["Smoke detector", 4]])
    body = _upload(client, project_id, path).json()

    assert body["result"]["items"][0]["per_floor"]["GF"] == 4
    assert body["archive_path"] is None
    assert "not reachable on this PC" in body["filed_note"]


def test_nothing_is_filed_when_the_workbook_is_not_a_schedule(client, db_session, tmp_path):
    """A file that turned out not to be a schedule is not left in the
    archive."""
    _login(client)
    project_id = _project(client, "30886")
    folder = tmp_path / "EP-30886"
    folder.mkdir()
    project = db_session.get(Project, project_id)
    project.source_folder_path = str(folder)
    db_session.commit()

    book = openpyxl.Workbook()
    book.active["A1"] = "Nothing of interest"
    path = tmp_path / "notes.xlsx"
    book.save(path)

    assert _upload(client, project_id, path).status_code == 422
    assert not (folder / "03- Design").exists()


# --- quantities set by hand -------------------------------------------------------------------


def test_a_quantity_can_be_raised_and_lowered_and_everything_follows(client, db_session, tmp_path):
    _login(client)
    project_id = _project(client, "30890")
    path = _workbook(tmp_path, headings=["Description", "GF", "1 to 3"],
                     rows=[["Photoelectric smoke detector", 4, 5], ["Heat detector", 1, 1]])
    body = _upload(client, project_id, path).json()
    row = body["result"]["items"][0]["row"]
    assert body["result"]["grand_total"] == (4 + 15) + (1 + 3)

    raised = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}",
                          json={"floor": "GF", "quantity": 9})
    assert raised.status_code == 200, raised.text
    result = raised.json()["result"]
    smoke = next(item for item in result["items"] if item["row"] == row)
    assert smoke["per_floor"]["GF"] == 9
    # The line's total, the floor's, the system's and the whole schedule's.
    assert smoke["total"] == 9 + 15
    assert result["totals"]["GF"] == 9 + 1
    assert result["grand_total"] == (9 + 15) + (1 + 3)
    assert result["systems"]["FAS"] == result["grand_total"]

    lowered = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}",
                           json={"floor": "GF", "quantity": 2}).json()["result"]
    assert next(i for i in lowered["items"] if i["row"] == row)["per_floor"]["GF"] == 2


def test_taking_a_quantity_to_nothing_leaves_the_floor_blank_not_zero(client, tmp_path):
    """A blank means the item is not on that floor; a zero would read as a
    BOQ line of none."""
    _login(client)
    project_id = _project(client, "30891")
    path = _workbook(tmp_path, headings=["Description", "GF", "ROOF"],
                     rows=[["Photoelectric smoke detector", 4, 2]])
    row = _upload(client, project_id, path).json()["result"]["items"][0]["row"]

    result = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}",
                          json={"floor": "GF", "quantity": 0}).json()["result"]
    item = result["items"][0]
    assert "GF" not in item["per_floor"] and item["per_floor"]["ROOF"] == 2
    assert item["total"] == 2


def test_an_edit_is_kept_and_read_back(client, tmp_path):
    _login(client)
    project_id = _project(client, "30892")
    path = _workbook(tmp_path, headings=["Description", "GF"], rows=[["Smoke detector", 4]])
    row = _upload(client, project_id, path).json()["result"]["items"][0]["row"]
    client.patch(f"/projects/{project_id}/floor-schedule/items/{row}", json={"floor": "GF", "quantity": 11})

    stored = client.get(f"/projects/{project_id}/floor-schedule").json()["result"]
    assert stored["items"][0]["per_floor"]["GF"] == 11
    assert stored["items"][0]["edited"] is True


def test_a_floor_or_a_line_that_is_not_there_is_refused(client, tmp_path):
    _login(client)
    project_id = _project(client, "30893")
    path = _workbook(tmp_path, headings=["Description", "GF"], rows=[["Smoke detector", 4]])
    row = _upload(client, project_id, path).json()["result"]["items"][0]["row"]

    assert client.patch(f"/projects/{project_id}/floor-schedule/items/{row}",
                        json={"floor": "Level 9", "quantity": 1}).status_code == 404
    assert client.patch(f"/projects/{project_id}/floor-schedule/items/999",
                        json={"floor": "GF", "quantity": 1}).status_code == 404


# --- which part a line is ordered as ------------------------------------------------------------


def test_a_line_is_settled_as_one_of_the_projects_own_materials(client, db_session, tmp_path):
    from app.models import ProjectProposedMaterial

    _login(client)
    project_id = _project(client, "30894")
    db_session.add(ProjectProposedMaterial(project_id=project_id, system_code="FAS", catalog_no="SIGA-PS",
                                           description="Photoelectric smoke detector", manufacturer="EDWARDS"))
    db_session.add(ProjectProposedMaterial(project_id=project_id, system_code="FAS", catalog_no="SB-01",
                                           description="Back box for detector", manufacturer="EDWARDS"))
    db_session.commit()

    path = _workbook(tmp_path, headings=["Description", "GF"], rows=[["Photoelectric smoke detector", 4]])
    row = _upload(client, project_id, path).json()["result"]["items"][0]["row"]

    offered = client.get(f"/projects/{project_id}/floor-schedule/materials?system=FAS").json()
    assert [material["part_no"] for material in offered] == ["SIGA-PS"]   # the back box is not offered

    chosen = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}/material",
                          json={"part_no": "SIGA-PS"})
    assert chosen.status_code == 200, chosen.text
    assert chosen.json()["result"]["items"][0]["material"]["part_no"] == "SIGA-PS"

    # A part that is not proposed for this project cannot be chosen.
    assert client.patch(f"/projects/{project_id}/floor-schedule/items/{row}/material",
                        json={"part_no": "SB-01"}).status_code == 400
    # ...and the choice can be taken back.
    cleared = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}/material", json={"part_no": None})
    assert "material" not in cleared.json()["result"]["items"][0]


def test_the_schedule_exports_as_a_pdf(client, tmp_path):
    _login(client)
    project_id = _project(client, "30895")
    path = _workbook(tmp_path, headings=["Description", "B1", "GF", "1 to 13"],
                     rows=[["Photoelectric smoke detector", 4, 6, 5],
                           ["Exit Light - Above Door", 2, 2, 1]])
    _upload(client, project_id, path)

    export = client.get(f"/projects/{project_id}/floor-schedule/export.pdf")
    assert export.status_code == 200
    assert export.headers["content-type"] == "application/pdf"
    assert export.content.startswith(b"%PDF")
    assert "BOQ Floor Wise.pdf" in export.headers["content-disposition"]


def test_a_range_that_does_not_divide_evenly_still_totals_exactly(tmp_path):
    """Five devices across thirteen typical floors is not a whole number
    per floor. Rounding each floor and then adding them up made a row of
    17 read 17.005; the totals are summed from what was read."""
    path = _workbook(
        tmp_path,
        headings=["Description", "B1", "GF", "1 to 13", "ROOF", "Total"],
        rows=[["Smoke detector", 4, 6, 5, 2, 17]],
    )
    read = floor_schedule.read(path).as_dict()
    assert read["typical_reading"] == "across"
    item = read["items"][0]
    assert item["total"] == 17 and read["grand_total"] == 17
    assert read["totals"]["B1"] == 4 and read["totals"]["ROOF"] == 2
    # ...and the engineer is told why a floor shows a fraction.
    assert any("does not divide evenly" in warning for warning in read["warnings"])


def test_stepping_a_quantity_on_such_a_sheet_keeps_the_total_whole(client, tmp_path):
    _login(client)
    project_id = _project(client, "30896")
    path = _workbook(
        tmp_path,
        headings=["Description", "B1", "GF", "1 to 13", "ROOF", "Total"],
        rows=[["Smoke detector", 4, 6, 5, 2, 17]],
    )
    row = _upload(client, project_id, path).json()["result"]["items"][0]["row"]

    result = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}",
                          json={"floor": "GF", "quantity": 7}).json()["result"]
    assert result["items"][0]["total"] == 18
    assert result["grand_total"] == 18
    assert result["systems"]["FAS"] == 18


# --- kept in step with the workbook in the project folder ---------------------------------------


def _project_folder(client, db_session, ep_number: str, tmp_path: Path):
    """A project whose OneDrive folder is reachable, and that folder."""
    from app.models import Project

    _login(client)
    project_id = _project(client, ep_number)
    folder = tmp_path / f"EP-{ep_number}"
    folder.mkdir()
    project = db_session.get(Project, project_id)
    project.source_folder_path = str(folder)
    db_session.commit()
    return project_id, folder


def _design_workbook(folder: Path, name: str, gf: int) -> Path:
    design = folder / "03- Design"
    design.mkdir(exist_ok=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "FLOOR WISE"
    for index, heading in enumerate(["DRG.NAME", "GF", "1 to 3"], start=1):
        sheet.cell(row=1, column=index, value=heading)
    sheet.cell(row=2, column=1, value="Photoelectric smoke detector")
    sheet.cell(row=2, column=2, value=gf)
    sheet.cell(row=2, column=3, value=5)
    path = design / name
    book.save(path)
    return path


def test_a_workbook_already_filed_is_read_without_being_uploaded(client, db_session, tmp_path):
    """The schedule an engineer keeps lives in 03- Design. Asking for it to
    be uploaded again would be asking them to do the platform's filing."""
    project_id, folder = _project_folder(client, db_session, "30900", tmp_path)
    _design_workbook(folder, "EP-30900 FLOOR WISE BOQ.xlsx", gf=4)

    body = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert body["result"] is not None
    assert body["file_name"] == "EP-30900 FLOOR WISE BOQ.xlsx"
    assert body["source_path"] == "03- Design/EP-30900 FLOOR WISE BOQ.xlsx"
    assert body["result"]["items"][0]["per_floor"]["GF"] == 4
    assert "03- Design" in body["filed_note"]


def test_the_workbook_is_read_again_only_when_it_changes(client, db_session, tmp_path):
    project_id, folder = _project_folder(client, db_session, "30901", tmp_path)
    _design_workbook(folder, "schedule.xlsx", gf=4)

    first = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert first["filed_note"] is not None                      # read the first time
    again = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert again["filed_note"] is None                          # nothing to do
    assert again["updated_at"] == first["updated_at"]

    _design_workbook(folder, "schedule.xlsx", gf=9)             # the engineer edits it in Excel
    after = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert "has changed" in after["filed_note"]
    assert after["result"]["items"][0]["per_floor"]["GF"] == 9


def test_a_quantity_set_by_hand_survives_the_workbook_being_read_again(client, db_session, tmp_path):
    """A correction is the engineer's own knowledge; a re-read of the sheet
    must not throw it away."""
    project_id, folder = _project_folder(client, db_session, "30902", tmp_path)
    _design_workbook(folder, "schedule.xlsx", gf=4)
    body = client.get(f"/projects/{project_id}/floor-schedule").json()
    row = body["result"]["items"][0]["row"]

    edited = client.patch(f"/projects/{project_id}/floor-schedule/items/{row}",
                          json={"floor": "GF", "quantity": 11}).json()["result"]
    assert edited["items"][0]["per_floor"]["GF"] == 11

    # The sheet now says something else on that floor; the correction stands.
    _design_workbook(folder, "schedule.xlsx", gf=6)
    after = client.get(f"/projects/{project_id}/floor-schedule").json()["result"]
    assert after["items"][0]["per_floor"]["GF"] == 11
    assert after["items"][0]["edited"] is True
    # ...and the floors the engineer did not touch follow the sheet.
    assert after["items"][0]["per_floor"]["Level 1"] == 5


def test_a_correction_whose_line_has_gone_is_reported(client, db_session, tmp_path):
    project_id, folder = _project_folder(client, db_session, "30903", tmp_path)
    _design_workbook(folder, "schedule.xlsx", gf=4)
    row = client.get(f"/projects/{project_id}/floor-schedule").json()["result"]["items"][0]["row"]
    client.patch(f"/projects/{project_id}/floor-schedule/items/{row}", json={"floor": "GF", "quantity": 11})

    # The line is renamed in Excel, so the correction has nothing to land on.
    book = openpyxl.load_workbook(folder / "03- Design" / "schedule.xlsx")
    book["FLOOR WISE"].cell(row=2, column=1, value="Heat detector")
    book.save(folder / "03- Design" / "schedule.xlsx")

    after = client.get(f"/projects/{project_id}/floor-schedule").json()["result"]
    assert any("no longer on the sheet" in warning for warning in after["warnings"])


def test_a_project_with_no_workbook_filed_is_still_asked_for_one(client, db_session, tmp_path):
    project_id, _folder = _project_folder(client, db_session, "30904", tmp_path)
    body = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert body["result"] is None and body["source_path"] is None


def test_excels_own_lock_file_is_not_mistaken_for_the_schedule(client, db_session, tmp_path):
    """Excel leaves "~$name.xlsx" beside an open workbook."""
    project_id, folder = _project_folder(client, db_session, "30905", tmp_path)
    _design_workbook(folder, "schedule.xlsx", gf=4)
    (folder / "03- Design" / "~$schedule.xlsx").write_bytes(b"not a workbook")

    body = client.get(f"/projects/{project_id}/floor-schedule").json()
    assert body["file_name"] == "schedule.xlsx"


def test_a_spacer_row_holding_a_stray_number_is_not_an_item(tmp_path):
    """One real sheet carries a row whose name cell is just "0". It would
    otherwise reach the page as a line with no device and no quantity, and
    make a system tab of its own."""
    path = _workbook(
        tmp_path,
        headings=["DRG.NAME", "GF", "Total"],
        rows=[["Smoke detector", 4, 4], ["0", None, 0]],
    )
    read = floor_schedule.read(path).as_dict()
    assert [item["description"] for item in read["items"]] == ["Smoke detector"]
    assert set(read["systems"]) == {"FAS"}
