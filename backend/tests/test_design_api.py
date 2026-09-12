"""The Voice Evacuation design API: import from the project's folder, edit,
recalculate, and the load limit that comes from the DesignRule."""

from pathlib import Path

import openpyxl

from app.core.config import get_settings
from app.models import DesignRule, ProjectDesign, RoleEnum

from .conftest import login, make_user

settings = get_settings()


def _login_admin(client):
    return login(client, settings.default_admin_email, settings.default_admin_password)


def _amplifier_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    for coordinate, value in {
        "C1": "Ceiling Speaker", "D1": "Watts per Area", "E1": "Required Watts", "F1": "Proposed Amplifier",
        "A2": "Tapping Wattage", "C2": 0.75,
        "A3": "GF", "C3": 32, "D3": 24, "E3": 42, "F3": 50,
        "A4": "L1", "C4": 24, "D4": 18,
        "A5": "L2", "C5": 40, "D5": 30, "E5": 30, "F5": 50,
    }.items():
        ws[coordinate] = value
    ws.merge_cells("E3:E4")
    ws.merge_cells("F3:F4")
    wb.save(path)


def _project(client, folder: Path) -> int:
    resp = client.post(
        "/projects",
        json={"ep_number": "24601", "project_name": "Tower", "source_folder_path": str(folder), "design_sheets": []},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _setup(client, tmp_path) -> int:
    folder = tmp_path / "EP-24601 Tower"
    _amplifier_workbook(folder / "Design" / "AMP TA CALC.xlsx")
    (folder / "Design" / "BOQ.xlsx").write_bytes(b"")
    (folder / "Design" / "~$AMP TA CALC.xlsx").write_bytes(b"")
    _login_admin(client)
    return _project(client, folder)


def test_before_import_there_is_no_design_but_the_rule_is_shown(client, tmp_path):
    project_id = _setup(client, tmp_path)
    body = client.get(f"/projects/{project_id}/design/ve").json()
    assert body["design"] is None and body["result"] is None
    assert body["rule"]["data"] == {"fraction": 0.8}
    assert "2026-09-11" in body["rule"]["source"]


def test_workbook_candidates_put_likely_ones_first_and_skip_lock_files(client, tmp_path):
    project_id = _setup(client, tmp_path)
    body = client.get(f"/projects/{project_id}/design/ve/workbooks").json()
    assert [(c["path"], c["likely"]) for c in body] == [
        ("Design/AMP TA CALC.xlsx", True),
        ("Design/BOQ.xlsx", False),
    ]


def test_import_calculates_against_the_rule(client, tmp_path):
    project_id = _setup(client, tmp_path)
    resp = client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"})
    assert resp.status_code == 200, resp.text
    body = resp.json()

    design, result = body["design"], body["result"]
    assert design["source"]["path"] == "Design/AMP TA CALC.xlsx"
    assert design["max_load_fraction"] == 0.8
    assert design["max_load_rule_id"] == body["rule"]["id"]
    assert [(c["required_watts"], c["limit_watts"], c["status"]) for c in result["channels"]] == [
        (42, 40, "over_limit"),
        (30, 40, "ok"),
    ]
    assert result["channels_failing"] == 1
    assert body["updated_by"] == "Platform Administrator"

    # Stored: a fresh read gives the same.
    assert client.get(f"/projects/{project_id}/design/ve").json()["result"] == result


def test_a_later_rule_version_does_not_change_an_imported_design(client, db_session, tmp_path):
    project_id = _setup(client, tmp_path)
    client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"})

    old = db_session.query(DesignRule).filter(DesignRule.category == "ve.limit").one()
    from app.core.timeutils import utc_now

    old.superseded_at = utc_now()
    db_session.add(DesignRule(category=old.category, key=old.key, version=2, data={"fraction": 0.9}, source="test"))
    db_session.commit()

    body = client.get(f"/projects/{project_id}/design/ve").json()
    assert body["design"]["max_load_fraction"] == 0.8
    assert body["rule"]["version"] == 1
    # Re-importing is what picks up the new rule.
    body = client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"}).json()
    assert body["design"]["max_load_fraction"] == 0.9
    assert body["result"]["channels_failing"] == 0


def test_import_only_reads_inside_the_project_folder(client, tmp_path):
    project_id = _setup(client, tmp_path)
    _amplifier_workbook(tmp_path / "elsewhere.xlsx")
    for path in ("../elsewhere.xlsx", str(tmp_path / "elsewhere.xlsx"), "Design/missing.xlsx"):
        resp = client.post(f"/projects/{project_id}/design/ve/import", json={"path": path})
        assert resp.status_code == 400, path


def test_an_unreadable_workbook_is_reported_not_stored(client, tmp_path):
    project_id = _setup(client, tmp_path)
    resp = client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/BOQ.xlsx"})
    assert resp.status_code == 422
    assert "Could not open" in resp.json()["detail"]
    assert client.get(f"/projects/{project_id}/design/ve").json()["design"] is None


def test_edits_recalculate_but_cannot_move_the_limit(client, tmp_path):
    project_id = _setup(client, tmp_path)
    design = client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"}).json()[
        "design"
    ]
    # The engineer moves L1 onto the second channel and asks for a 100 W amp there.
    design["channels"][1]["first_zone"] = 1
    design["channels"][1]["amplifier_watts"] = 100
    design["max_load_fraction"] = 1.0
    design["source"] = None

    body = client.put(f"/projects/{project_id}/design/ve", json=design).json()
    assert [(c["required_watts"], c["status"]) for c in body["result"]["channels"]] == [(24, "ok"), (48, "ok")]
    assert body["design"]["max_load_fraction"] == 0.8
    assert body["design"]["source"]["path"] == "Design/AMP TA CALC.xlsx"


def test_an_inconsistent_edit_is_rejected(client, tmp_path):
    project_id = _setup(client, tmp_path)
    design = client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"}).json()[
        "design"
    ]
    design["channels"][1]["first_zone"] = 9
    assert client.put(f"/projects/{project_id}/design/ve", json=design).status_code == 422


def test_editing_before_an_import_is_refused(client, tmp_path):
    project_id = _setup(client, tmp_path)
    design = {"speaker_types": [], "zones": [], "max_load_fraction": 0.8}
    assert client.put(f"/projects/{project_id}/design/ve", json=design).status_code == 409


def test_viewers_can_read_but_not_import(client, db_session, tmp_path):
    project_id = _setup(client, tmp_path)
    make_user(db_session, "viewer@ep-platform.com", RoleEnum.viewer)
    login(client, "viewer@ep-platform.com")
    assert client.get(f"/projects/{project_id}/design/ve").status_code == 200
    assert client.get(f"/projects/{project_id}/design/ve/workbooks").status_code == 403
    resp = client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"})
    assert resp.status_code == 403


def test_deleting_a_project_removes_its_design(client, db_session, tmp_path):
    project_id = _setup(client, tmp_path)
    client.post(f"/projects/{project_id}/design/ve/import", json={"path": "Design/AMP TA CALC.xlsx"})
    assert db_session.query(ProjectDesign).count() == 1
    assert client.delete(f"/projects/{project_id}").status_code == 204
    # The row itself, not just the route: the foreign key has no ON DELETE,
    # so this rests on the ORM cascade.
    assert db_session.query(ProjectDesign).count() == 0
    # The load-limit rule belongs to the platform, not the project.
    assert db_session.query(DesignRule).filter(DesignRule.category == "ve.limit").count() == 1
