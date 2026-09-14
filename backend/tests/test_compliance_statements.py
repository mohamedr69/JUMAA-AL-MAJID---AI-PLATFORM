"""Preparing and checking compliance statements.

The synthetic specification copies the shape of EP-30784's 283111: a running
header naming the project and plot on every page, the section number, PART /
article / paragraph / item labels each on a line of their own, a table of
contents. The submitted statement copies EP-16481's workbook: a title block,
then clause | answer columns with one answer merged down a list.

Filling a statement from the knowledge base and reviewing a clause with the
model are in test_compliance_knowledge.
"""

import io
from pathlib import Path

import openpyxl
import pymupdf
import pytest

import app.ai.provider as provider_module
from app.ai.provider import RecordingProvider
from app.compliance import statements
from app.compliance.spec_text import read_bytes
from app.compliance.verify import verify
from app.core.config import get_settings
from app.models import Project
from app.services.spec_finder import find_specs

from .conftest import login
from .test_compliance import _clear_cache, _project

settings = get_settings()

HEADER = "BINGHATTI SKYBLADE ON PLOT {plot}, BURJ KHALIFA DISTRICT"
BODY = [
    # page 1: cover
    ["SECTION 283111", "ADDRESSABLE FIRE DETECTION AND VOICE EVACUATION SYSTEM", "Revision History"],
    # page 2: table of contents
    ["INDEX", "PART 1 - GENERAL ........................ 2", "1.1 RELATED DOCUMENTS ............ 2",
     "1.2 SUMMARY ................. 2", "1.3 DEFINITIONS ............ 2", "PART 2 - PRODUCTS ........... 3"],
    # page 3
    ["PART 1 - GENERAL", "1.1", "RELATED DOCUMENTS", "A.",
     "Drawings and general provisions of the Contract apply to this Section.",
     "1.2", "SUMMARY", "A.", "This Section includes fire alarm systems with manual stations and detectors.",
     "B.", "Related Sections include the following:", "1.", "Division 26 Section \"Grounding and Bonding\".",
     "2.", "Division 26 Section \"Raceways and Boxes\".",
     "1.3", "DEFINITIONS", "A.", "FACP: Fire alarm control panel."],
    # page 4
    ["PART 2 - PRODUCTS", "2.1", "SMOKE DETECTORS", "A.", "General: Include the following features:",
     "1.", "Operating Voltage: 24-V dc, nominal.", "2.", "Self-Restoring: Detectors do not require resetting.",
     "B.", "Graphic annunciator mounted at the fire command centre.",
     "C.", "A strip printer mounted in the main FACP enclosure.",
     "D.", "Firefighters' telephone handsets connected to the FACP.",
     "END OF SECTION 283111"],
]


def spec_pdf(path: Path, plot: str = "3450398", body=BODY) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open()
    for number, lines in enumerate(body, start=1):
        page = doc.new_page()
        page.insert_text((40, 40), HEADER.format(plot=plot), fontsize=8)
        page.insert_text((40, 52), "Section 283111 - Addressable Fire Detection and Voice Evacuation System", fontsize=8)
        page.insert_text((40, 80), "\n".join(lines), fontsize=9)
        page.insert_text((40, 800), f"Page {number} of {len(body)}", fontsize=8)
    doc.save(path)
    doc.close()
    return path


def statement_xlsx(rows: list[tuple[str, str, str]], title: str = "PROJECT NAME : BINGHATTI SKYBLADE",
                   manufacturer: str = "EDWARDS", merge: tuple[int, int] | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["COMPLIANCE STATEMENT - FIRE ALARM SYSTEM"])
    ws.append([title])
    ws.append([f"MANUFACTURER: {manufacturer}"])
    ws.append(["SL.NO", "PROJECT SPECIFICATION - SECTION 283111", "AL ARABIA COMPLIANCE", "REMARKS"])
    for label, text, answer in rows:
        ws.append([label, text, answer, ""])
    if merge:
        ws.merge_cells(start_row=merge[0], start_column=3, end_row=merge[1], end_column=3)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def no_ai():
    """A provider that records every call, so a test can assert none was made."""
    provider = RecordingProvider()
    provider_module.set_provider(provider)
    yield provider
    provider_module.set_provider(None)


# --- reading the specification -----------------------------------------------------


def test_the_clauses_of_a_specification(tmp_path):
    spec = read_bytes(spec_pdf(tmp_path / "spec.pdf").read_bytes())
    refs = [(c.ref, c.heading) for c in spec.clauses]
    assert refs[:6] == [("PART 1", True), ("1.1", True), ("1.1.A", False), ("1.2", True), ("1.2.A", False), ("1.2.B", False)]
    # Items on a line of their own are clauses of their own, not text of B.
    assert ("1.2.B.1", False) in refs and ("1.2.B.2", False) in refs
    assert ("2.1.A.2", False) in refs and ("2.1.C", False) in refs
    by_ref = {c.ref: c for c in spec.clauses}
    assert by_ref["2.1.A.1"].text == "Operating Voltage: 24-V dc, nominal."
    assert by_ref["2.1.C"].page == 4
    # The table of contents is not read as clauses, and the running header
    # is kept apart as the specification's identity.
    assert sum(1 for c in spec.clauses if c.ref == "1.1") == 1
    assert any("PLOT 3450398" in line for line in spec.header_lines)
    assert not any("PLOT" in c.text for c in spec.clauses)
    assert spec.section_numbers == ["283111"]


def test_whose_specification_it_is(client, db_session, tmp_path):
    project = Project(ep_number="30784", project_name="BINGHATTI SKYBLADE", plot_number="3450398",
                      created_by_id=1)
    own = read_bytes(spec_pdf(tmp_path / "own.pdf").read_bytes())
    other = read_bytes(spec_pdf(tmp_path / "other.pdf", plot="3466814").read_bytes())

    assert verify(project, "FAS", own).project == "same"
    verdict = verify(project, "FAS", other)
    assert (verdict.project, verdict.system) == ("different", "same")
    assert "3466814" in verdict.evidence[0]
    # A lighting section is not the fire alarm's specification.
    assert verify(project, "EML", own).system == "different"
    # Nothing to go on: unknown, which is what the model is asked about.
    project.plot_number, project.project_name = None, "Sky Hills Astra"
    assert verify(project, "FAS", other).project == "unknown"


# --- reading a statement ------------------------------------------------------------------


def test_a_statement_is_read_by_what_its_columns_hold():
    rows = [("A", "Clause one text.", "Comply"), ("B", "List of devices:", ""), ("1", "Manual stations.", "Comply"),
            ("2", "Smoke detectors.", ""), ("C", "Door holders.", "By others"), ("D", "Graphic annunciator.", "N/A")]
    # One "Comply" merged down the list's two items (rows 7-8 of the sheet).
    content = statement_xlsx(rows, merge=(7, 8))
    statement = statements.read_statement(content, "statement.xlsx")
    answers = {r.text: r.answer for r in statement.rows}
    assert answers["Manual stations."] == "Comply" and answers["Smoke detectors."] == "Comply"
    assert answers["Door holders."] == "By others" and answers["Graphic annunciator."] == "Not applicable"
    assert any("BINGHATTI" in line for line in statement.title_lines)


def test_a_file_that_is_not_a_statement():
    wb = openpyxl.Workbook()
    wb.active.append(["BOQ", "qty"])
    buffer = io.BytesIO()
    wb.save(buffer)
    assert statements.read_statement(buffer.getvalue(), "boq.xlsx") is None
    with pytest.raises(ValueError):
        statements.read_statement(b"%PDF", "statement.pdf")


# --- finding the specification ------------------------------------------------------------


def test_a_specification_that_numbers_nothing_is_found_by_its_content(tmp_path):
    """EP-29495's "FA SPECS.pdf" and "CBS SPECS.pdf" open "( I ) FIRE
    DETECTION SYSTEM / PART 1 - GENERAL" and never carry a section number;
    the engineer's own upload, "Specification FAS <stamp>.pdf", is the same
    document. Both must be found."""
    doc = pymupdf.open()
    for text in ["( I ) FIRE DETECTION SYSTEM\nPART 1 - GENERAL\n1.1\nDESCRIPTION:\nA.\nThis section covers the fire alarm.",
                 "1.2\nSCOPE:\nA.\nBasic Performance:"]:
        doc.new_page().insert_text((50, 60), text, fontsize=9)
    (tmp_path / "named" / "MS" / "FA").mkdir(parents=True)
    doc.save(tmp_path / "named" / "MS" / "FA" / "FA SPECS.pdf")
    (tmp_path / "unnamed" / "Scan").mkdir(parents=True)
    doc.save(tmp_path / "unnamed" / "Scan" / "Scan0012.pdf")
    doc.close()

    # Named as a specification: read in the first pass.
    (match,) = find_specs(tmp_path / "named", {"FAS"})[0]
    assert (match.path, match.matched_on, match.heading) == (str(Path("MS") / "FA" / "FA SPECS.pdf"), "content",
                                                              "FIRE DETECTION SYSTEM")
    # Named for nothing: found by its first page, since nothing else was.
    (match,) = find_specs(tmp_path / "unnamed", {"FAS"})[0]
    assert (match.path, match.matched_on) == (str(Path("Scan") / "Scan0012.pdf"), "content")
    assert find_specs(tmp_path / "unnamed", {"FAS"}, read_contents=False)[0] == []
    # And it verifies as the fire alarm's, from its opening line.
    project = Project(ep_number="29495", project_name="IVY GARDEN 2", created_by_id=1)
    spec = read_bytes((tmp_path / "unnamed" / "Scan" / "Scan0012.pdf").read_bytes())
    assert verify(project, "FAS", spec, "Scan0012.pdf").system == "same"
    assert verify(project, "EML", spec, "Scan0012.pdf").system == "different"


def test_uploaded_specifications_are_listed_for_the_system_chosen(client, tmp_path, monkeypatch):
    """The upload used to vanish: the file's name said nothing the finder
    recognised and the content carried no section number."""
    monkeypatch.setattr(settings, "uploads_root", str(tmp_path / "uploads"))
    _clear_cache()
    folder = tmp_path / "EP-29495"
    folder.mkdir()
    login(client, settings.default_admin_email, settings.default_admin_password)
    project_id = _project(client, folder, systems=("Fire Alarm", "Central Battery System"))
    listed = client.get(f"/projects/{project_id}/compliance").json()
    # Project Info's systems are the tabs, even with no Design Sheet or BOQ.
    assert [s["code"] for s in listed["systems"]] == ["FAS", "CBS"]
    # The page carries the knowledge base's status, never a path.
    assert listed["knowledge"]["records"]["responses"] == 0 and "root" not in listed["knowledge"]

    doc = pymupdf.open()
    doc.new_page().insert_text((50, 60), "( I ) FIRE DETECTION SYSTEM\nPART 1 - GENERAL\n1.1\nDESCRIPTION:", fontsize=9)
    doc.new_page().insert_text((50, 60), "A.\nThe fire alarm shall be addressable.", fontsize=9)
    content = doc.tobytes()
    doc.close()
    resp = client.post(f"/projects/{project_id}/compliance/specs",
                       files={"file": ("whatever.pdf", content, "application/pdf")}, data={"system_code": "CBS"})
    assert resp.status_code == 200, resp.text
    cbs = next(s for s in resp.json()["systems"] if s["code"] == "CBS")
    (spec,) = cbs["specs"]
    assert spec["uploaded"] is True and spec["matched_on"] == "uploaded" and spec["pages"] == 2
    # Trusted for the system chosen, but honest about what it reads like.
    assert spec["verification"]["system"] == "different"
    assert client.get(f"/projects/{project_id}/compliance/file", params={"path": spec["path"]}).status_code == 200

    # And it can be taken away again.
    assert client.delete(f"/projects/{project_id}/compliance/specs", params={"path": spec["path"]}).status_code == 204
    _clear_cache()
    assert next(s for s in client.get(f"/projects/{project_id}/compliance").json()["systems"] if s["code"] == "CBS")["specs"] == []
    assert client.delete(f"/projects/{project_id}/compliance/specs", params={"path": str(folder / "x.pdf")}).status_code == 404


def test_a_specification_inside_a_past_submittal(tmp_path):
    """EP-30353 has no specification of its own, but its PA/BGM submittal
    carries section 27 51 16 between the cover and the datasheets."""
    doc = pymupdf.open()
    pages = ["MATERIAL SUBMITTAL - PA / BGM SYSTEM"] + [f"SECTION 27 51 16\nPUBLIC ADDRESS SYSTEM\nPART 1 - GENERAL\nclause {i}" for i in range(3)] + ["TOA DATASHEET PC-1860BS"] * 2
    for text in pages:
        doc.new_page().insert_text((50, 60), text, fontsize=9)
    (tmp_path / "MS").mkdir()
    doc.save(tmp_path / "MS" / "MAT-ELEC-023-R1-PA-BGM System.pdf")
    doc.close()
    spec_pdf(tmp_path / "MS" / "FA SCHEMATIC DIAGRAM.pdf")  # a drawing is never a specification

    (match,) = find_specs(tmp_path, {"PAVA"})[0]
    assert (match.system_code, match.matched_on, match.first_page, match.last_page) == ("PAVA", "submittal", 2, 4)
    assert find_specs(tmp_path, {"PAVA"}, look_in_submittals=False)[0] == []


# --- prepare and check through the API ---------------------------------------------------


def test_prepare_applies_the_rules_and_calls_no_model(client, tmp_path, monkeypatch, no_ai):
    monkeypatch.setattr(settings, "uploads_root", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "ai_enabled", True)
    _clear_cache()
    folder = tmp_path / "EP-30784"
    spec_pdf(folder / "Specification" / "283111 - FIRE DETECTION.pdf")
    login(client, settings.default_admin_email, settings.default_admin_password)
    project_id = _project(client, folder)

    listed = client.get(f"/projects/{project_id}/compliance").json()
    spec = listed["systems"][0]["specs"][0]
    assert spec["verification"]["system"] == "same"
    body = {"system_code": "FAS", "path": spec["path"], "member": spec["member"], "first_page": spec["first_page"],
            "last_page": spec["last_page"]}
    resp = client.post(f"/projects/{project_id}/compliance/prepare", json=body)
    assert resp.status_code == 200, resp.text
    statement = resp.json()
    rows = {r["ref"]: r for r in statement["rows"]}
    # Rules: related sections and definitions are noted without asking.
    assert (rows["1.2.B.1"]["response"], rows["1.2.B.1"]["source"], rows["1.2.B.1"]["workflow"]) == ("Noted", "rule", "autofilled")
    assert (rows["1.3.A"]["response"], rows["1.3.A"]["source"]) == ("Noted", "rule")
    # A lead-in is answered by its items.
    assert rows["2.1.A"]["source"] == "lead_in"
    # Everything else waits, unfilled, with both statuses open.
    assert rows["2.1.A.1"]["response"] == "" and rows["2.1.A.1"]["workflow"] == "unfilled"
    assert rows["2.1.A.1"]["technical"] == {"status": None, "origin": None, "verified": False}
    assert statement["summary"]["by_workflow"]["unfilled"] == 6 and statement["ai_calls"] == 0
    assert statement["summary"]["inputs"]["spec_sha256"] == statement["spec"]["sha256"]
    assert no_ai.calls == 0

    # The engineer's change replaces the proposal and is audited.
    edited = client.patch(f"/projects/{project_id}/compliance/statements/{statement['id']}",
                          json={"rows": [{"id": rows["2.1.C"]["id"], "response": "Not applicable", "remark": "No printer in BOQ",
                                          "technical_status": "not_applicable"}]})
    assert edited.status_code == 200
    changed = next(r for r in edited.json()["rows"] if r["ref"] == "2.1.C")
    assert (changed["response"], changed["source"], changed["origin"], changed["remark"]) == ("Not applicable", "engineer", "manual", "No printer in BOQ")
    assert changed["technical"]["status"] == "not_applicable" and changed["workflow"] == "autofilled"
    # A row cannot be marked reviewed without a response.
    empty = client.post(f"/projects/{project_id}/compliance/statements/{statement['id']}/rows/{rows['2.1.D']['id']}/reviewed",
                        json={"reviewed": True})
    assert empty.status_code == 400

    # And the workbook carries the change.
    export = client.get(f"/projects/{project_id}/compliance/statements/{statement['id']}/export")
    assert export.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(export.content)).active
    values = [[c for c in row] for row in sheet.iter_rows(values_only=True)]
    assert any(row[1] == "A strip printer mounted in the main FACP enclosure." and row[2] == "Not applicable" for row in values)
    assert any(str(row[0]).startswith("PROJECT NAME : BINGHATTI SKYBLADE") for row in values if row[0])
    assert no_ai.calls == 0


def test_the_audit_records_every_change(client, db_session, tmp_path, monkeypatch, no_ai):
    from app.models import ComplianceAudit

    monkeypatch.setattr(settings, "uploads_root", str(tmp_path / "uploads"))
    _clear_cache()
    folder = tmp_path / "EP-30784"
    spec_pdf(folder / "Specification" / "283111 - FIRE DETECTION.pdf")
    login(client, settings.default_admin_email, settings.default_admin_password)
    project_id = _project(client, folder)
    spec = client.get(f"/projects/{project_id}/compliance").json()["systems"][0]["specs"][0]
    statement = client.post(f"/projects/{project_id}/compliance/prepare", json={"system_code": "FAS", "path": spec["path"]}).json()
    clause = next(r["id"] for r in statement["rows"] if r["ref"] == "2.1.B")
    client.patch(f"/projects/{project_id}/compliance/statements/{statement['id']}", json={"rows": [{"id": clause, "response": "Comply"}]})
    client.post(f"/projects/{project_id}/compliance/statements/{statement['id']}/rows/{clause}/reviewed", json={"reviewed": True})
    trail = db_session.query(ComplianceAudit).filter_by(statement_id=statement["id"], clause_id=clause).order_by(ComplianceAudit.id).all()
    assert [(a.action, a.origin, a.previous_response, a.current_response, a.review_status) for a in trail] == [
        ("manual", "manual", "", "Comply", "autofilled"), ("reviewed", "manual", "Comply", "Comply", "reviewed")]
    assert trail[0].spec_sha256 == statement["spec"]["sha256"] and trail[0].boq_hash and trail[0].user_id == 1


def test_check_finds_what_a_statement_misses(client, db_session, tmp_path, monkeypatch, no_ai):
    monkeypatch.setattr(settings, "uploads_root", str(tmp_path / "uploads"))
    monkeypatch.setattr(settings, "ai_enabled", True)
    _clear_cache()
    folder = tmp_path / "EP-30784"
    spec_pdf(folder / "Specification" / "283111 - FIRE DETECTION.pdf", plot="3466814")
    login(client, settings.default_admin_email, settings.default_admin_password)
    project_id = _project(client, folder)
    db_session.get(Project, project_id).plot_number = "3450398"
    db_session.commit()
    spec = client.get(f"/projects/{project_id}/compliance").json()["systems"][0]["specs"][0]

    submitted = statement_xlsx(
        [
            ("A", "Drawings and general provisions of the Contract apply to this Section.", "Noted"),
            ("A", "This Section includes fire alarm systems with manual stations and detectors.", "Comply"),
            ("1", "Operating Voltage: 24-V dc, nominal.", "Comply"),
            ("B", "Graphic annunciator mounted at the fire command centre.", ""),
            ("C", "A strip printer mounted in the main FACP enclosure.", "Not applicable"),
            ("X", "Something the specification never asked.", "Comply"),
        ],
        title="PROJECT NAME : DISTRICT COOLING PLANT JVC",
        manufacturer="NOTIFIER",
    )
    resp = client.post(
        f"/projects/{project_id}/compliance/check",
        data={"system_code": "FAS", "path": spec["path"]},
        files={"file": ("FAS Compliance.xlsx", submitted, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    result = resp.json()
    codes = {f["code"] for f in result["summary"]["general"]}
    assert "spec_other_project" in codes and "statement_other_project" in codes
    rows = {r["ref"]: r for r in result["rows"]}
    finding = lambda ref: {f["code"] for f in rows[ref]["findings"]}  # noqa: E731
    assert "missing" in finding("2.1.A.2")
    assert "unanswered" in finding("2.1.B")
    assert "hard_requirement" in finding("2.1.A.1")
    assert finding("1.2.A") == set()
    assert rows["2.1.A"]["source"] == "lead_in"
    assert result["summary"]["rows_not_in_spec"] >= 1
    # Checking calls no model, whatever the settings say.
    assert no_ai.calls == 0 and result["ai_calls"] == 0

    # Not a statement at all: said plainly.
    bad = client.post(f"/projects/{project_id}/compliance/check",
                      data={"system_code": "FAS", "path": spec["path"]},
                      files={"file": ("notes.xlsx", statement_xlsx([]), "application/octet-stream")})
    assert bad.status_code == 400


def test_statements_are_for_editors(client, db_session, tmp_path, monkeypatch):
    from app.models import RoleEnum

    from .conftest import make_user

    monkeypatch.setattr(settings, "uploads_root", str(tmp_path / "uploads"))
    _clear_cache()
    folder = tmp_path / "EP-30784"
    spec_pdf(folder / "Specification" / "283111 - FIRE DETECTION.pdf")
    login(client, settings.default_admin_email, settings.default_admin_password)
    project_id = _project(client, folder)
    make_user(db_session, "viewer@ep-platform.com", RoleEnum.viewer)
    login(client, "viewer@ep-platform.com")
    body = {"system_code": "FAS", "path": "Specification/283111 - FIRE DETECTION.pdf"}
    assert client.post(f"/projects/{project_id}/compliance/prepare", json=body).status_code == 403
    assert client.get(f"/projects/{project_id}/compliance/statements").json() == []
