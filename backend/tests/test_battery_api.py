"""The battery calculation API and the datasheet catalogue behind it."""

from app.core.config import get_settings
from app.models import DesignRule, RoleEnum

from .conftest import login, make_user

settings = get_settings()

PANEL = "EST4 Main Fire Alarm Control Panel"


def _login_admin(client):
    return login(client, settings.default_admin_email, settings.default_admin_password)


def _project_with_boq(client) -> int:
    resp = client.post("/projects", json={"ep_number": "30784", "project_name": "Skyblade", "design_sheets": []})
    project_id = resp.json()["id"]
    lines = [
        {"system_code": "FAS", "group_heading": PANEL, "description": "EST4 panel. Includes:", "quantity": "1"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "4-CPU", "description": "CPU", "quantity": "1"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "3-SDDC2", "description": "Loop card", "quantity": "5"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "4-FIL", "description": "Filler plate", "quantity": "16"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "12V26A", "description": "Battery, 12 V @ 26 AH", "quantity": "2"},
        {"system_code": "FAS", "group_heading": "Booster Power Supply", "catalog_no": "BPS10A/230", "description": "BPS", "quantity": "1"},
        {"system_code": "FAS", "catalog_no": "SIGA-PS", "description": "Smoke detector", "quantity": "300"},
    ]
    assert client.put(f"/projects/{project_id}/boq", json=lines).status_code == 200
    return project_id


def _current(client, part, standby, alarm, source="EST4 datasheet p.3"):
    return client.post(
        "/design-rules/part-currents",
        json={"part_no": part, "standby_ma": standby, "alarm_ma": alarm, "source": source},
    )


def test_nothing_is_claimed_until_every_part_has_a_current(client):
    _login_admin(client)
    project_id = _project_with_boq(client)

    body = client.get(f"/projects/{project_id}/design/battery").json()
    assert body["rule"]["data"] == {"standby_hours": 24, "alarm_minutes": 30, "panel_voltage": 24}
    assert [(g["heading"], g["treatment"]) for g in body["groups"]] == [
        (PANEL, "panel"),
        ("Booster Power Supply", "bps"),
        (None, "ungrouped"),
    ]
    panel, bps = body["panels"]
    assert (bps["kind"], bps["name"]) == ("bps", "BPS")
    assert panel["missing_parts"] == ["4-CPU", "3-SDDC2", "4-FIL"]
    assert panel["status"] == "incomplete" and panel["selected"] is None
    assert panel["quoted_ah"] == 26
    assert body["unlisted_batteries"] == [{"part_no": "12V26A", "capacity_ah": 26, "voltage": 12}]

    assert _current(client, "4-CPU", 211, 211).status_code == 200
    assert _current(client, "3-SDDC2 ", 264, 336).status_code == 200
    assert _current(client, "4-fil", 0, 0, source="No electrical load: blank filler plate").status_code == 200

    panel = client.get(f"/projects/{project_id}/design/battery").json()["panels"][0]
    # (1531 x 24 + 1891 x 0.5) / 1000 = 37.69 Ah, against 26 Ah quoted.
    assert (panel["standby_ma"], panel["alarm_ma"]) == (1531, 1891)
    assert round(panel["required_ah"], 2) == 37.69
    assert not panel["lower_bound"] and panel["quoted_short"]
    # No ROCKET battery on file in the tests (no datasheet library): nothing to select.
    assert panel["status"] == "no_selection"
    # No battery units catalogued yet, so nothing to propose from.
    assert panel["selected"] is None


def test_battery_units_drive_the_proposal(client):
    _login_admin(client)
    project_id = _project_with_boq(client)
    for part, current in [("4-CPU", (211, 211)), ("3-SDDC2", (264, 336)), ("4-FIL", (0, 0))]:
        _current(client, part, *current)
    for part, ah in [("12V26A", 26), ("12V42A", 42), ("12V65A", 65)]:
        resp = client.post(
            "/design-rules/battery-units",
            json={"part_no": part, "capacity_ah": ah, "voltage": 12, "brand": "ROCKET", "source": "ROCKET datasheet"},
        )
        assert resp.status_code == 200

    body = client.get(f"/projects/{project_id}/design/battery").json()
    panel = body["panels"][0]
    # A 42 Ah pair carries it. With the 1.2 design factor this took 65 Ah:
    # dropping the factor is what moved it down a size.
    assert [(p["part_no"], p["units"], p["strings"]) for p in panel["selected"]] == [("12V42A", 2, 1)]
    assert panel["required_ah"] <= 42 and panel["required_ah"] > 26
    assert body["unlisted_batteries"] == []


def test_catalogue_entries_are_versioned_not_overwritten(client, db_session):
    _login_admin(client)
    first = _current(client, "4-CPU", 211, 211).json()
    assert first["version"] == 1
    # Same values and source: no new version.
    assert _current(client, "4-cpu", 211, 211).json()["id"] == first["id"]
    second = _current(client, "4-CPU", 220, 230, source="EST4 datasheet rev B").json()
    assert (second["version"], second["key"]) == (2, "4-CPU")

    listed = client.get("/design-rules/part-currents").json()
    assert [(r["key"], r["version"], r["data"]["standby_ma"]) for r in listed] == [("4-CPU", 2, 220)]
    old = db_session.get(DesignRule, first["id"])
    assert old.superseded_at is not None


def test_a_current_needs_a_source(client):
    _login_admin(client)
    resp = client.post("/design-rules/part-currents", json={"part_no": "4-CPU", "standby_ma": 1, "alarm_ma": 1, "source": ""})
    assert resp.status_code == 422
    resp = _current(client, " , ", 1, 1)
    assert resp.status_code == 422


def test_viewers_can_read_but_not_add_to_the_catalogue(client, db_session):
    _login_admin(client)
    project_id = _project_with_boq(client)
    make_user(db_session, "viewer@ep-platform.com", RoleEnum.viewer)
    login(client, "viewer@ep-platform.com")
    assert client.get(f"/projects/{project_id}/design/battery").status_code == 200
    assert client.get("/design-rules/part-currents").status_code == 200
    assert _current(client, "4-CPU", 211, 211).status_code == 403


# --- filling currents from the datasheet library ------------------------------


def _library(tmp_path):
    import pymupdf

    from app.services.datasheet_library import DatasheetLibrary

    root = tmp_path / "EST4"
    root.mkdir()
    for name, lines in {
        "01- 4-CPU.pdf": [(322, 64, "Standby"), (412, 64, "211 mA at 24 Vdc"), (322, 74, "Alarm"),
                          (412, 84, "See the 4-COMREL Common Relay Module")],
        "08- SDDC.pdf": [(53, 56, "Catalog Number", "hebo"), (298, 56, "3-SSDC2", "hebo"), (473, 56, "3-SDDC2", "hebo"),
                         (260, 101, "Standby 144 mA Alarm 204 mA"), (436, 101, "Standby 264 mA Alarm 336 mA")],
    }.items():
        doc = pymupdf.open()
        page = doc.new_page(width=612, height=792)
        for x, y, text, *font in lines:
            page.insert_text((x, y), text, fontsize=8, fontname=font[0] if font else "helv")
        doc.save(root / name)
        doc.close()
    return DatasheetLibrary("EDWARDS", root)


def _fill_project(client) -> int:
    resp = client.post("/projects", json={"ep_number": "30784", "project_name": "Skyblade", "design_sheets": []})
    project_id = resp.json()["id"]
    line = lambda part, desc, qty: {  # noqa: E731
        "system_code": "FAS", "group_heading": PANEL, "catalog_no": part, "description": desc,
        "quantity": qty, "manufacturer": "EDWARDS",
    }
    lines = [
        {"system_code": "FAS", "group_heading": PANEL, "description": "EST4 panel. Includes:", "quantity": "1"},
        line("4-CPU", "Central Processor Module", "1"),
        line("3-SDDC2", "Dual loop SIGA Data Controller", "2"),
        line("3-CHAS7", "Seven Space Chassis with Rails", "1"),
        line("4-COMREL", "Common Relay Module", "1"),
        line("12V65A", "Battery, 12 V @ 65 AH", "2"),
    ]
    client.put(f"/projects/{project_id}/boq", json=lines)
    return project_id


def test_currents_are_filled_from_the_datasheets(client, tmp_path, monkeypatch):
    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)

    filled = client.post(f"/projects/{project_id}/design/battery/fill-currents").json()["filled"]
    assert [(f["part_no"], f["standby_ma"], f["alarm_ma"]) for f in filled] == [
        ("4-CPU", 211, 211),
        ("3-SDDC2", 264, 336),
        ("3-CHAS7", 0, 0),
        # Built into the CPU: no current of its own (see INCLUDED_IN_MODULE).
        ("4-COMREL", 0, 0),
    ]
    assert filled[0]["source"].startswith("Edwards datasheet (01- 4-CPU.pdf), p.1: read automatically")
    # Both settled by the equipment current table, which knows them.
    assert filled[2]["source"].startswith("Equipment current table: no electrical load")
    assert filled[3]["source"].startswith("Equipment current table: no current of its own: built into 4-CPU")

    body = client.get(f"/projects/{project_id}/design/battery").json()
    (panel,) = body["panels"]
    assert panel["missing_parts"] == []
    assert body["unresolved"] == []
    loaded = {line["part_no"]: line for line in panel["lines"] if line["kind"] == "load" and not line["missing_current"]}
    assert loaded["3-SDDC2"]["total_standby_ma"] == 528
    assert "SDDC.pdf" in loaded["3-SDDC2"]["current_source"]

    # Only missing parts are touched: a second open fills nothing.
    assert client.post(f"/projects/{project_id}/design/battery/fill-currents").json()["filled"] == []


def test_filling_is_for_editors(client, db_session, tmp_path, monkeypatch):
    import app.routers.design as design_router

    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": _library(tmp_path)})
    _login_admin(client)
    project_id = _fill_project(client)
    make_user(db_session, "viewer@ep-platform.com", RoleEnum.viewer)
    login(client, "viewer@ep-platform.com")
    assert client.post(f"/projects/{project_id}/design/battery/fill-currents").status_code == 403


# --- per-panel cards, settings, added components, export -----------------------


def _two_panel_types(client) -> int:
    resp = client.post("/projects", json={"ep_number": "30784", "project_name": "Skyblade", "design_sheets": []})
    project_id = resp.json()["id"]
    sub = "EST4 Fire Alarm Panel"
    lines = [
        {"system_code": "FAS", "group_heading": PANEL, "description": "Main panel. Includes:", "quantity": "1"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "4-CPU", "description": "CPU", "quantity": "1"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "4-FIL", "description": "Blank Filler Plate", "quantity": "16"},
        {"system_code": "FAS", "group_heading": PANEL, "catalog_no": "12V65A", "description": "Battery, 12 V @ 65 AH", "quantity": "2"},
        {"system_code": "FAS", "group_heading": sub, "description": "Sub panel. Includes:", "quantity": "2"},
        {"system_code": "FAS", "group_heading": sub, "catalog_no": "4-CPU", "description": "CPU", "quantity": "1"},
        {"system_code": "FAS", "group_heading": sub, "catalog_no": "12V26A", "description": "Battery, 12 V @ 26 AH", "quantity": "2"},
    ]
    client.put(f"/projects/{project_id}/boq", json=lines)
    _current(client, "4-CPU", 211, 211)
    _current(client, "4-FIL", 0, 0, source="No electrical load: blank filler plate")
    return project_id


def test_every_physical_panel_gets_its_own_card(client):
    _login_admin(client)
    project_id = _two_panel_types(client)
    panels = client.get(f"/projects/{project_id}/design/battery").json()["panels"]
    assert [(p["name"], p["heading"], p["instance"]) for p in panels] == [
        ("FACP-01", PANEL, 1),
        ("FACP-02", "EST4 Fire Alarm Panel", 1),
        ("FACP-03", "EST4 Fire Alarm Panel", 2),
    ]
    assert panels[0]["settings"] == {"standby_hours": 24, "alarm_minutes": 30, "panel_voltage": 24}
    assert panels[0]["overridden"] == []


def test_panel_settings_and_added_components_are_saved_and_recalculated(client):
    _login_admin(client)
    project_id = _two_panel_types(client)
    panels = client.get(f"/projects/{project_id}/design/battery").json()["panels"]
    key = panels[2]["key"]
    design = {
        "panels": {
            key: {
                "name": "FACP-L20", "location": "Level 20 electrical room", "standby_hours": 72,
                "extra_components": [
                    {"description": "Door holder supply", "quantity": 2, "standby_ma": 0, "alarm_ma": 500, "source": "Door holder datasheet p.2"}
                ],
            }
        }
    }
    resp = client.put(f"/projects/{project_id}/design/battery", json=design)
    assert resp.status_code == 200, resp.text
    third = resp.json()["panels"][2]
    assert (third["name"], third["location"], third["overridden"]) == ("FACP-L20", "Level 20 electrical room", ["standby_hours"])
    # 211 mA x 72 h + (211 + 1000) mA x 0.5 h = 15797.5 mAh = 15.797 Ah
    assert (third["standby_ma"], third["alarm_ma"]) == (211, 1211)
    assert round(third["required_ah"], 3) == 15.797
    added = [line for line in third["lines"] if line["extra_index"] == 0]
    assert added and added[0]["total_alarm_ma"] == 1000
    # The identical panel beside it keeps the rule.
    second = resp.json()["panels"][1]
    assert second["overridden"] == [] and second["alarm_ma"] == 211
    # Stored with the project.
    assert client.get(f"/projects/{project_id}/design/battery").json()["design"]["panels"][key]["standby_hours"] == 72


def test_panel_settings_are_validated_and_editor_only(client, db_session):
    _login_admin(client)
    project_id = _two_panel_types(client)
    bad = {"panels": {"x": {"standby_hours": 0}}}
    assert client.put(f"/projects/{project_id}/design/battery", json=bad).status_code == 422
    make_user(db_session, "viewer@ep-platform.com", RoleEnum.viewer)
    login(client, "viewer@ep-platform.com")
    assert client.put(f"/projects/{project_id}/design/battery", json={"panels": {}}).status_code == 403


def test_export_lists_only_parts_that_draw_current_with_live_formulas(client):
    from io import BytesIO

    import openpyxl

    _login_admin(client)
    project_id = _two_panel_types(client)
    resp = client.get(f"/projects/{project_id}/design/battery/export.xlsx")
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "FACP-01", "FACP-02", "FACP-03"]
    ws = wb["FACP-01"]
    parts = [ws.cell(row=r, column=1).value for r in range(7, ws.max_row + 1) if ws.cell(row=r, column=1).value]
    assert parts == ["4-CPU"]  # not the filler plate, not the battery
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=C7*D7" in formulas and any(f.startswith("=SUM(F7:F7)") for f in formulas)
    assert any("Capacity sufficient" in f for f in formulas)

    one = client.get(f"/projects/{project_id}/design/battery/export.xlsx", params={"panel": "FAS|EST4 Fire Alarm Panel|2"})
    assert openpyxl.load_workbook(BytesIO(one.content)).sheetnames == ["FACP-03"]
    assert "FACP-03" in one.headers["content-disposition"]
    assert client.get(f"/projects/{project_id}/design/battery/export.xlsx", params={"panel": "nope"}).status_code == 404


# --- parts built into another module --------------------------------------


def test_a_part_built_into_another_module_draws_no_current(client, tmp_path, monkeypatch):
    """4-COMREL is on the EST4 CPU board. Its current is already inside
    4-CPU's figure, and no datasheet gives it one of its own -- so left as
    "missing" it held the panel at a lower bound for ever and no battery was
    ever proposed. Platform owner, 13 September 2026."""
    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)

    client.post(f"/projects/{project_id}/design/battery/fill-currents")
    body = client.get(f"/projects/{project_id}/design/battery").json()
    panel = body["panels"][0]

    comrel = next(line for line in panel["lines"] if line.get("part_no") == "4-COMREL")
    assert (comrel["standby_ma"], comrel["alarm_ma"]) == (0, 0)
    assert "built into 4-CPU" in comrel["current_source"]
    # The panel is now a firm figure, not a lower bound, so it can be sized.
    assert panel["lower_bound"] is False
    assert panel["status"] != "incomplete"


def test_a_built_in_part_is_not_given_its_host_s_current(client, tmp_path, monkeypatch):
    """The host module's datasheet is the one that mentions the part, so
    reading it would count 4-CPU's own current a second time."""
    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)

    client.post(f"/projects/{project_id}/design/battery/fill-currents")
    panel = client.get(f"/projects/{project_id}/design/battery").json()["panels"][0]

    cpu = next(line for line in panel["lines"] if line.get("part_no") == "4-CPU")
    comrel = next(line for line in panel["lines"] if line.get("part_no") == "4-COMREL")
    assert cpu["standby_ma"] == 211
    assert comrel["standby_ma"] == 0


def test_a_battery_is_selected_for_the_load_even_when_the_boq_quotes_less(client, tmp_path, monkeypatch):
    """The selection follows the load, not the BOQ. Where it comes out above
    what the BOQ quotes, `quoted_short` marks it so the page can say so --
    the capacity is right and the BOQ is the thing to change."""
    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)
    client.post(f"/projects/{project_id}/design/battery/fill-currents")
    for part, ah in [("12V26A", 26), ("12V42A", 42), ("12V65A", 65)]:
        assert client.post(
            "/design-rules/battery-units",
            json={"part_no": part, "capacity_ah": ah, "voltage": 12, "brand": "ROCKET", "source": "ROCKET datasheet"},
        ).status_code == 200

    # Quote a battery far below what this panel's load needs. Not one of the
    # catalogued parts, so its capacity is read from the BOQ description.
    boq = client.get(f"/projects/{project_id}/boq").json()
    for item in boq:
        if item["catalog_no"] == "12V65A":
            item["catalog_no"] = "12V7A"
            item["description"] = "Battery, 12 V @ 7 AH"
    client.put(f"/projects/{project_id}/boq", json=boq)

    panel, = client.get(f"/projects/{project_id}/design/battery").json()["panels"]
    assert panel["quoted_ah"] == 7
    assert panel["quoted_short"] is True
    # Still selected, and sized to the requirement rather than to the quote.
    assert panel["selected"], "a battery must still be proposed for the real load"
    assert panel["selected_ah"] >= panel["required_ah"]
    assert panel["selected_ah"] > panel["quoted_ah"]


# --- the calculation as the company's own sheet -----------------------------


def test_the_pdf_export_follows_the_company_template(client, tmp_path, monkeypatch):
    """Laid out to FACP_Battery_Calculation_Clean.pdf in the submittal
    builder: A3 landscape, one page per panel, the project block, the
    connected-equipment table and the sizing rows it names."""
    import pymupdf

    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)
    client.post(f"/projects/{project_id}/design/battery/fill-currents")

    resp = client.get(f"/projects/{project_id}/design/battery/export.pdf")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    doc = pymupdf.open(stream=resp.content, filetype="pdf")
    page = doc[0]
    # The template's page: A3 landscape.
    assert round(page.rect.width) == 1191 and round(page.rect.height) == 842
    text = page.get_text()
    doc.close()

    for heading in ("BATTERY CALCULATION", "Connected equipment", "Battery sizing", "Calculation notes"):
        assert heading in text
    for label in ("PROJECT", "CLIENT", "CONSULTANT", "MEP CONTRACTOR", "FAS SUPPLIER", "MANUFACTURER", "DATE"):
        assert label in text
    for column in ("Part no.", "Description", "Qty", "Standby / unit", "Alarm / unit", "Standby total", "Alarm total"):
        assert column in text
    for row in ("Standby current", "Alarm current", "Standby, 24 h", "Alarm, 30 min",
                "Required capacity", "SELECTED BATTERY"):
        assert row in text


def test_the_sheet_carries_this_project_and_its_numbers(client, tmp_path, monkeypatch):
    import pymupdf

    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)
    client.post(f"/projects/{project_id}/design/battery/fill-currents")
    panel = client.get(f"/projects/{project_id}/design/battery").json()["panels"][0]

    resp = client.get(f"/projects/{project_id}/design/battery/export.pdf")
    doc = pymupdf.open(stream=resp.content, filetype="pdf")
    text = doc[0].get_text()
    doc.close()

    assert "Skyblade".upper() in text.upper()
    assert "4-CPU" in text                                   # a connected part
    assert f"{panel['standby_ma']:,.3f}" in text             # its own totals
    assert f"{panel['required_ah']:,.3f}" in text


def test_one_page_per_panel_and_a_page_count(client, tmp_path, monkeypatch):
    import pymupdf

    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)
    panels = client.get(f"/projects/{project_id}/design/battery").json()["panels"]

    resp = client.get(f"/projects/{project_id}/design/battery/export.pdf")
    doc = pymupdf.open(stream=resp.content, filetype="pdf")
    pages = len(doc)
    first = doc[0].get_text()
    doc.close()

    assert pages == len(panels)
    assert f"01 / {pages:02d}" in first


def test_a_single_panel_can_be_exported(client, tmp_path, monkeypatch):
    import pymupdf

    import app.routers.design as design_router

    library = _library(tmp_path)
    monkeypatch.setattr(design_router, "_datasheet_libraries", lambda: {"EDWARDS": library})
    _login_admin(client)
    project_id = _fill_project(client)
    key = client.get(f"/projects/{project_id}/design/battery").json()["panels"][0]["key"]

    resp = client.get(f"/projects/{project_id}/design/battery/export.pdf?panel={key}")
    doc = pymupdf.open(stream=resp.content, filetype="pdf")
    assert len(doc) == 1
    doc.close()

    assert client.get(f"/projects/{project_id}/design/battery/export.pdf?panel=nope").status_code == 404


def test_a_panel_still_missing_a_current_is_shown_as_a_lower_bound(client, tmp_path, monkeypatch):
    """The sheet must not read as a finished figure while a part has no
    current: every sizing row carries the >= marker instead."""
    import pymupdf

    _login_admin(client)
    project_id = _fill_project(client)          # no datasheet library: nothing fills

    resp = client.get(f"/projects/{project_id}/design/battery/export.pdf")
    doc = pymupdf.open(stream=resp.content, filetype="pdf")
    text = doc[0].get_text()
    doc.close()

    # ">=" rather than the sign: the base PDF fonts have no U+2265 and
    # substitute a middle dot for it.
    assert ">=" in text
    assert "Pending: a part has no current yet" in text


def test_a_panel_is_recalculated_only_when_its_own_inputs_move(client, db_session, monkeypatch):
    """Each panel's calculation is kept under the hash of its inputs: an
    unchanged panel is read back, a panel whose part got a current is the
    one recalculated, and a recalculation that fails leaves the previous
    figures visible, marked stale."""
    from app.models import BatteryPanelResult
    from app.routers import design as design_router

    _login_admin(client)
    project_id = _project_with_boq(client)
    first = client.get(f"/projects/{project_id}/design/battery").json()
    assert (first["reused_panels"], first["recalculated_panels"]) == (0, 2)
    again = client.get(f"/projects/{project_id}/design/battery").json()
    assert (again["reused_panels"], again["recalculated_panels"]) == (2, 0)
    assert again["panels"] == first["panels"]

    # A current arrives for a part of the main panel: that panel alone is recalculated.
    assert _current(client, "4-CPU", 211, 211).status_code == 200
    third = client.get(f"/projects/{project_id}/design/battery").json()
    assert (third["reused_panels"], third["recalculated_panels"]) == (1, 1)
    main = next(p for p in third["panels"] if p["kind"] == "panel")
    assert "4-CPU" not in main["missing_parts"] and not main["stale"]

    # The next change reaches a calculation that falls over: the previous
    # figures stay, marked stale with the reason; the other panel is untouched.
    def broken(*args, **kwargs):
        raise RuntimeError("the sizing rule fell over")

    monkeypatch.setattr(design_router, "calculate_panel", broken)
    assert _current(client, "3-SDDC2", 264, 336).status_code == 200
    fourth = client.get(f"/projects/{project_id}/design/battery").json()
    stale = next(p for p in fourth["panels"] if p["kind"] == "panel")
    assert stale["stale"] and "fell over" in stale["error"]
    assert stale["standby_ma"] == main["standby_ma"] and stale["missing_parts"] == main["missing_parts"]
    assert (fourth["reused_panels"], fourth["recalculated_panels"]) == (1, 0)
    assert {r.state for r in db_session.query(BatteryPanelResult).filter(BatteryPanelResult.project_id == project_id)} == {"fresh", "stale"}

    # Recalculating works again: fresh figures replace the stale ones.
    monkeypatch.undo()
    fifth = client.get(f"/projects/{project_id}/design/battery").json()
    fixed = next(p for p in fifth["panels"] if p["kind"] == "panel")
    assert not fixed["stale"] and "3-SDDC2" not in fixed["missing_parts"]
