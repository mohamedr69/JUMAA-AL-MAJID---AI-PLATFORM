"""Each user's record: what they did is logged, and read back as their account."""

from io import BytesIO

from openpyxl import load_workbook

from app.models import ActivityEvent, RoleEnum

from .conftest import login, make_user
from .test_projects import _login_admin, _payload_without_documents, _valid_project_payload


def _actions(client, path="/auth/me/activity"):
    return [e["action"] for e in client.get(path).json()["events"]]


def _engineer_with_project(client, db_session):
    make_user(db_session, "eng@ep-platform.com", RoleEnum.fire_alarm_design_engineer)
    assert login(client, "eng@ep-platform.com").status_code == 200
    project = client.post("/projects", json=_payload_without_documents()).json()
    return project["id"]


def test_login_project_boq_and_submittal_work_is_recorded_for_the_user(client, db_session):
    project_id = _engineer_with_project(client, db_session)

    lines = [{"system_code": "FAS", "description": "Smoke detector", "quantity": "40"},
             {"system_code": "FAS", "description": "Sounder", "quantity": "12"}]
    assert client.put(f"/projects/{project_id}/boq", json=lines).status_code == 200
    assert client.post(f"/projects/{project_id}/boq/revisions", json={"note": "first issue"}).status_code == 201
    submittal = client.post(f"/projects/{project_id}/submittals",
                            json={"title": "Detectors", "revision": "R00", "status": "not_submitted"}).json()
    assert client.patch(f"/projects/{project_id}/submittals/{submittal['id']}",
                        json={"status": "approved"}).status_code == 200
    assert client.put(f"/projects/{project_id}", json={**_valid_project_payload(), "client": "Emaar"}).status_code == 200

    actions = _actions(client)
    for expected in ("auth.login", "project.created", "boq.saved", "boq.revision_issued",
                     "submittal.created", "submittal.updated", "project.updated"):
        assert expected in actions
    # Newest first.
    assert actions[0] == "project.updated"

    events = {e["action"]: e for e in client.get("/auth/me/activity").json()["events"]}
    assert events["boq.saved"]["detail"]["lines"] == 2
    assert events["submittal.updated"]["detail"]["status"] == "not_submitted -> approved"
    assert events["project.updated"]["detail"]["client"] == "Samana -> Emaar"
    assert events["project.created"]["project_label"] == "EP-29495 IVY Garden 2"


def test_opening_a_project_is_recorded_once_per_window(client, db_session):
    project_id = _engineer_with_project(client, db_session)
    for _ in range(3):
        assert client.get(f"/projects/{project_id}").status_code == 200
    assert _actions(client).count("project.opened") == 1


def test_a_refused_change_leaves_no_event(client, db_session):
    make_user(db_session, "viewer@ep-platform.com", RoleEnum.viewer)
    _login_admin(client)
    project_id = client.post("/projects", json=_valid_project_payload()).json()["id"]

    login(client, "viewer@ep-platform.com")
    assert client.put(f"/projects/{project_id}/boq", json=[]).status_code == 403
    assert "boq.saved" not in _actions(client)


def test_account_gathers_projects_submittals_revisions_and_activity(client, db_session):
    project_id = _engineer_with_project(client, db_session)
    client.get(f"/projects/{project_id}")
    client.put(f"/projects/{project_id}/boq", json=[{"system_code": "FAS", "description": "CPU", "quantity": "1"}])
    client.post(f"/projects/{project_id}/boq/revisions", json={})
    client.post(f"/projects/{project_id}/submittals", json={"title": "Panels", "revision": "R00", "status": "under_review"})

    account = client.get("/auth/me/account").json()
    assert account["user"]["email"] == "eng@ep-platform.com"
    [project] = account["projects"]
    assert project["created"] and project["assigned"] and project["opened_count"] == 1
    assert account["boq_revisions"][0]["label"] == "Rev 00"
    assert account["submittals"][0]["title"] == "Panels"
    assert account["counts"]["boq_saves"] == 1
    assert account["counts"]["logins"] == 1
    assert account["activity_total"] == len(account["activity"])


def test_deleted_project_stays_in_the_record_of_who_worked_on_it(client, db_session):
    _login_admin(client)
    project_id = client.post("/projects", json=_valid_project_payload()).json()["id"]
    assert client.delete(f"/projects/{project_id}").status_code == 204

    account = client.get("/auth/me/account").json()
    [project] = account["projects"]
    assert project["deleted"] and project["label"] == "EP-29495 IVY Garden 2"
    assert "project.deleted" in _actions(client)


def test_only_an_admin_reads_another_users_record(client, db_session):
    engineer = make_user(db_session, "eng@ep-platform.com", RoleEnum.fire_alarm_design_engineer)
    login(client, "eng@ep-platform.com")
    assert client.get(f"/users/{engineer.id}/account").status_code == 403

    _login_admin(client)
    account = client.get(f"/users/{engineer.id}/account")
    assert account.status_code == 200
    assert account.json()["counts"]["logins"] == 1
    assert client.get("/users/99999/account").status_code == 404


def test_activity_filters_by_project_and_action_and_pages(client, db_session):
    project_id = _engineer_with_project(client, db_session)
    for quantity in ("1", "2", "3"):
        client.put(f"/projects/{project_id}/boq", json=[{"system_code": "FAS", "description": "CPU", "quantity": quantity}])

    page = client.get(f"/auth/me/activity?action=boq&project_id={project_id}&limit=2").json()
    assert page["total"] == 3 and len(page["events"]) == 2
    assert all(e["action"] == "boq.saved" for e in page["events"])
    rest = client.get(f"/auth/me/activity?action=boq&limit=2&offset=2").json()
    assert len(rest["events"]) == 1


def test_record_exports_as_a_workbook(client, db_session):
    project_id = _engineer_with_project(client, db_session)
    client.put(f"/projects/{project_id}/boq", json=[{"system_code": "FAS", "description": "CPU", "quantity": "1"}])

    response = client.get("/auth/me/account/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Account", "Projects", "Activity", "Submittals", "BOQ revisions", "Compliance"]
    actions = [row[1] for row in workbook["Activity"].iter_rows(min_row=2, values_only=True)]
    assert "boq.saved" in actions and "auth.login" in actions


def test_admin_changes_to_an_account_are_recorded(client, db_session):
    engineer = make_user(db_session, "eng@ep-platform.com", RoleEnum.fire_alarm_design_engineer)
    _login_admin(client)
    assert client.patch(f"/users/{engineer.id}", json={"role": "design_manager"}).status_code == 200

    event = db_session.query(ActivityEvent).filter(ActivityEvent.action == "user.updated").one()
    assert event.entity_id == engineer.id
    assert event.detail == {"role": "fire_alarm_design_engineer -> design_manager"}


def test_logout_is_recorded(client, db_session):
    _engineer_with_project(client, db_session)
    assert client.post("/auth/logout").status_code == 200
    assert db_session.query(ActivityEvent).filter(ActivityEvent.action == "auth.logout").count() == 1
