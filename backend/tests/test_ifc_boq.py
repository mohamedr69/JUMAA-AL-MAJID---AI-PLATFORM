"""BOQ as per IFC drawings (app/ifc): the original BOQ Extraction tool's
end-to-end tests on synthetic drawings -- the same device inserted under
different, wrong or meaningless block names -- run against the platform.

The test bodies are the tool's, unchanged. They build one symbol library
across the file (a symbol verified in one test is recognised in the next),
so they share one database and one project, in file order. `_Api` carries
the tool's /api URLs over to the platform's routes: a drawing belongs to a
project here, and the library is shared by every project.
"""
from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

_TMP = Path(tempfile.mkdtemp(prefix="boq-test-"))

import ezdxf  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.config import get_settings  # noqa: E402
from app.main import app  # noqa: E402

from .conftest import login, reset_database  # noqa: E402


class _Api:
    """The tool's API, answered by the platform's routes."""

    def __init__(self, client: TestClient, project_id: int):
        self._client = client
        self._drawings = f"/projects/{project_id}/ifc-drawings"

    def _route(self, url: str, kwargs: dict) -> str:
        if url in ("/api/drawings/capabilities", "/api/drawings/reprocess"):
            return "/ifc/" + url.rsplit("/", 1)[1]
        if url == "/api/drawings":
            return self._drawings
        found = re.fullmatch(r"/api/drawings/(\d+)(/.*)?", url)
        if found:
            return f"{self._drawings}/{found.group(1)}{found.group(2) or ''}"
        if url.startswith("/api/device-types"):
            return "/ifc" + url[len("/api"):]
        found = re.fullmatch(r"/api/symbols/(verify|review)", url)
        if found:
            body = dict(kwargs.pop("json"))
            kwargs["json"] = body
            return f"{self._drawings}/{body.pop('drawing_id')}/{found.group(1)}"
        if url == "/api/symbols/unverify":
            return "/ifc/symbols/unverify"
        raise AssertionError(f"no platform route for {url}")

    def __getattr__(self, method):
        call = getattr(self._client, method)

        def send(url, **kwargs):
            return call(self._route(url, kwargs), **kwargs)

        return send


@pytest.fixture(scope="module")
def client():
    os.environ["BOQ_LIBRARY_PATH"] = str(_TMP / "symbol_library.json")
    reset_database()
    try:
        with TestClient(app) as c:
            settings = get_settings()
            assert login(c, settings.default_admin_email, settings.default_admin_password).status_code == 200
            project_id = c.post("/projects", json={"ep_number": "90001", "project_name": "IFC",
                                                   "design_sheets": []}).json()["id"]
            # Emergency lighting is not seeded yet; the tool's tests count an exit sign.
            c.post("/ifc/device-types", json={"code": "EXIT", "name": "Exit Sign (Single-sided)",
                                              "category": "emergency_light"})
            yield _Api(c, project_id)
    finally:
        os.environ.pop("BOQ_LIBRARY_PATH", None)


def smoke(block, letter="S"):
    block.add_circle((0, 0), 200)
    block.add_text(letter, height=150).set_placement((-50, -75))


def exit_sign(block):
    block.add_lwpolyline([(-300, -120), (300, -120), (300, 120), (-300, 120)], close=True)
    block.add_line((-300, 0), (300, 0))


def make_dxf(path: Path, placements: list[tuple[str, callable, float, float, float]], layer: str = "E-FIRE"):
    """placements: (block name, drawer, x, y, rotation)."""
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 4
    made = set()
    for name, draw, *_ in placements:
        if name not in made:
            draw(doc.blocks.new(name))
            made.add(name)
    msp = doc.modelspace()
    for name, _, x, y, rot in placements:
        msp.add_blockref(name, (x, y), dxfattribs={"rotation": rot, "layer": layer})
    doc.saveas(path)
    return path


def upload(client, path: Path) -> dict:
    with open(path, "rb") as f:
        r = client.post("/api/drawings", files={"file": (path.name, f, "application/dxf")})
    assert r.status_code == 201, r.text
    return r.json()


def type_id(client, code: str) -> int:
    return next(t["id"] for t in client.get("/api/device-types").json() if t["code"] == code)


def by_label(d: dict, label: str) -> list[dict]:
    return [g for g in d["groups"] if g["label"] == label]


def test_rejects_other_file_types(client):
    r = client.post("/api/drawings", files={"file": ("plan.pdf", b"%PDF", "application/pdf")})
    assert r.status_code == 415


def test_rejects_a_file_named_dwg_that_is_not_one(client):
    r = client.post("/api/drawings", files={"file": ("plan.dwg", b"not a drawing", "application/octet-stream")})
    assert r.status_code == 422
    assert "not a DWG" in r.json()["detail"]


def test_capabilities_reports_dwg_support(client):
    caps = client.get("/api/drawings/capabilities").json()
    assert caps["dxf"] is True
    assert isinstance(caps["dwg"], bool)


def test_dwg_upload_is_converted_then_extracted(client, monkeypatch):
    """The wiring, without AutoCAD: a stand-in converter writes the DXF."""
    import shutil

    from app.ifc.dxf import convert

    dxf = make_dxf(_TMP / "for_dwg.dxf", [("SD", smoke, 0, 0, 0), ("SD", smoke, 800, 0, 0)])
    seen = {}

    def fake_convert(dwg_path, dxf_path, converter=None):
        seen["dwg"] = Path(dwg_path).read_bytes()[:6]
        shutil.copyfile(dxf, dxf_path)
        return convert.ConversionResult(converter="Test converter", seconds=0.1, log_tail="")

    monkeypatch.setattr(convert, "convert_dwg_to_dxf", fake_convert)
    r = client.post("/api/drawings", files={"file": ("FA-999.dwg", b"AC1032" + bytes(64), "application/octet-stream")})
    assert r.status_code == 201, r.text
    d = r.json()
    assert seen["dwg"] == b"AC1032"
    assert d["filename"] == "FA-999.dwg"
    assert d["conversion"] == {"source_format": "dwg", "converter": "Test converter", "seconds": 0.1}
    assert by_label(d, "S")[0]["count"] == 2
    assert client.delete(f"/api/drawings/{d['id']}").status_code == 204


FA105_DWG = Path(r"C:/Users/moham/Desktop/FA-AI-Pipeline/00-INPUTS/1-IFC/257-Sheet - FA-105 - TYP(1ST TO 14TH) FLOORS FIRE ALARM LAYOUT.dwg")
FA105_DXF = FA105_DWG.with_suffix(".dxf")


def test_real_dwg_converts_to_the_same_symbols_as_a_hand_saved_dxf():
    """Only where AutoCAD (or ODA) and the FA-105 drawings are on the PC."""
    from app.ifc.dxf import convert
    from app.ifc.dxf.extract import extract

    if convert.find_converter() is None or not FA105_DWG.exists() or not FA105_DXF.exists():
        pytest.skip("needs a DWG converter and the FA-105 DWG/DXF pair")
    out = _TMP / "fa105_converted.dxf"
    res = convert.convert_dwg_to_dxf(FA105_DWG, out)
    assert out.stat().st_size > 0 and res.seconds < convert.TIMEOUT_S
    a = {g.signature: g.count for g in extract(str(FA105_DXF)).groups}
    b = {g.signature: g.count for g in extract(str(out)).groups}
    assert a == b


def test_same_symbol_under_many_names_is_one_group(client):
    p = make_dxf(_TMP / "a.dxf", [
        ("SMOKE DETECTOR-1001", smoke, 0, 0, 0),
        ("SMOKE DETECTOR-1002", smoke, 1000, 0, 90),
        ("SD", smoke, 2000, 0, 0),
        ("*U12", smoke, 3000, 0, 0),
        ("EXIT", exit_sign, 0, 2000, 0),
        ("EXIT", exit_sign, 1000, 2000, 0),
    ])
    d = upload(client, p)
    s = by_label(d, "S")
    assert len(s) == 1, "four block names, one symbol"
    assert s[0]["count"] == 4
    assert set(s[0]["block_names"]) == {"SMOKE DETECTOR-1001", "SMOKE DETECTOR-1002", "SD", "*U12"}
    assert d["totals"]["unknown_instances"] == 6


def test_verify_counts_into_the_right_tab(client):
    d = upload(client, _TMP / "a.dxf")
    sd = by_label(d, "S")[0]["signature"]
    ex = next(g for g in d["groups"] if g["label"] == "")["signature"]
    d = client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [sd], "device_type_id": type_id(client, "SD")}).json()
    d = client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [ex], "device_type_id": type_id(client, "EXIT")}).json()
    assert d["totals"]["fire_alarm"] == 4
    assert d["totals"]["emergency_light"] == 2
    assert d["totals"]["unknown_symbols"] == 0


def test_library_recognises_symbol_on_a_new_drawing_with_a_wrong_name(client):
    # A different drawing: the smoke detector is now called "HEAT DET" (wrong)
    p = make_dxf(_TMP / "b.dxf", [
        ("HEAT DET", smoke, 0, 0, 0),
        ("HEAT DET", smoke, 500, 0, 0),
    ])
    d = upload(client, p)
    g = by_label(d, "S")[0]
    assert g["status"] == "verified", "matched by drawing, not name"
    assert g["device_type"]["code"] == "SD"
    assert d["totals"]["fire_alarm"] == 2


def test_same_circle_different_letters_is_not_suggested(client):
    p = make_dxf(_TMP / "c.dxf", [
        ("SD", lambda b: smoke(b, "H"), 0, 0, 0),  # right name, wrong symbol
    ])
    d = upload(client, p)
    g = by_label(d, "H")[0]
    # The block name "SD" is a known alias, so the name suggests SD; but the
    # symbol is not verified and nothing is counted until a person decides.
    assert g["status"] in ("suggested", "unknown")
    assert d["totals"]["fire_alarm"] == 0


def smoke_as_arcs(block):
    """The same smoke detector drawn another way: two arcs, not a circle.
    Different entities, different signature, same look and letters."""
    block.add_arc((0, 0), 200, 0, 180)
    block.add_arc((0, 0), 200, 180, 360)
    block.add_text("S", height=150).set_placement((-50, -75))


def test_close_match_with_same_letters_is_counted_from_the_library(client):
    d = upload(client, make_dxf(_TMP / "arcs.dxf", [
        ("SMOKE DETECTOR-V94", smoke_as_arcs, 0, 0, 0),
        ("SMOKE DETECTOR-V94", smoke_as_arcs, 900, 0, 0),
    ]))
    g = by_label(d, "S")[0]
    assert g["status"] == "verified"
    assert g["device_type"]["code"] == "SD"
    assert g["match"]["kind"] == "library" and g["match"]["score"] >= 0.9
    assert d["totals"]["fire_alarm"] == 2


def test_close_match_to_not_a_device_is_only_suggested(client):
    def box(b):
        b.add_lwpolyline([(0, 0), (1000, 0), (1000, 600), (0, 600)], close=True)

    def box_as_lines(b):
        for a, c in [((0, 0), (1000, 0)), ((1000, 0), (1000, 600)), ((1000, 600), (0, 600)), ((0, 600), (0, 0))]:
            b.add_line(a, c)

    d = upload(client, make_dxf(_TMP / "box1.dxf", [("TABLE", box, 0, 0, 0)], layer="0"))
    client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [d["groups"][0]["signature"]], "ignore": True})
    d = upload(client, make_dxf(_TMP / "box2.dxf", [("SOMETHING", box_as_lines, 0, 0, 0)], layer="0"))
    g = d["groups"][0]
    assert g["status"] == "suggested", "never hidden automatically"
    assert g["suggestion"]["is_ignored"]


def test_library_file_is_written_and_restores_a_lost_symbol(client):
    import json

    from app.database import SessionLocal
    from app.ifc.library_file import import_library, library_path
    from app.models import IfcSymbol as Symbol

    data = json.loads(library_path().read_text(encoding="utf-8"))
    sd = [s for s in data["symbols"] if s["device_code"] == "SD"]
    assert sd and sd[0]["label"] == "S" and sd[0]["raster_hex"]
    assert any(t["code"] == "SD" for t in data["device_types"])

    # lose the symbol from the database only, as a fresh PC would
    with SessionLocal() as db:
        db.query(Symbol).filter(Symbol.signature == sd[0]["signature"]).delete()
        db.commit()
        assert import_library(db)["symbols"] == 1
        restored = db.query(Symbol).filter(Symbol.signature == sd[0]["signature"]).one()
        assert restored.device_type.code == "SD"


def test_a_new_device_type_is_kept_for_good(client):
    """A device missing from the list is added from the dropdown. It is
    saved in the database and in the library file, so a fresh database gets
    it back, together with the symbols verified as it."""
    import json

    from app.database import SessionLocal
    from app.ifc.library_file import import_library, library_path
    from app.models import IfcBlockAlias as BlockAlias, IfcDeviceType as DeviceType, IfcSymbol as Symbol

    r = client.post("/api/device-types", json={"code": " lhd ", "name": " Linear Heat Detector ", "category": "fire_alarm"})
    assert r.status_code == 201, r.text
    lhd = r.json()
    assert (lhd["code"], lhd["name"], lhd["unit"], lhd["is_active"]) == ("LHD", "Linear Heat Detector", "Nos", True)
    assert client.post("/api/device-types", json={"code": "Lhd", "name": "Again", "category": "fire_alarm"}).status_code == 409
    assert client.post("/api/device-types", json={"code": "  ", "name": "No code", "category": "fire_alarm"}).status_code == 422
    assert any(t["code"] == "LHD" for t in client.get("/api/device-types").json())

    def lhd_symbol(block):
        block.add_lwpolyline([(-300, -150), (300, -150), (300, 150), (-300, 150)], close=True)
        block.add_text("LHD", height=120).set_placement((-150, -60))

    d = upload(client, make_dxf(_TMP / "lhd.dxf", [("L1", lhd_symbol, 0, 0, 0), ("L1", lhd_symbol, 5000, 0, 0)]))
    g = by_label(d, "LHD")[0]
    d = client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [g["signature"]], "device_type_id": lhd["id"]}).json()
    assert by_label(d, "LHD")[0]["device_type"]["code"] == "LHD"
    assert d["totals"]["fire_alarm"] == 2

    data = json.loads(library_path().read_text(encoding="utf-8"))
    saved = next(t for t in data["device_types"] if t["code"] == "LHD")
    assert (saved["name"], saved["category"], saved["unit"]) == ("Linear Heat Detector", "fire_alarm", "Nos")
    assert any(s["device_code"] == "LHD" for s in data["symbols"])

    # a fresh database: the device type and its symbol come back from the file
    with SessionLocal() as db:
        sym = db.query(Symbol).filter(Symbol.signature == g["signature"]).one()
        db.query(BlockAlias).filter(BlockAlias.symbol_id == sym.id).delete()
        db.delete(sym)
        db.flush()
        db.query(DeviceType).filter(DeviceType.id == lhd["id"]).delete()
        db.commit()
        added = import_library(db)
        assert added["device_types"] == 1 and added["symbols"] == 1
        assert db.query(Symbol).filter(Symbol.signature == g["signature"]).one().device_type.name == "Linear Heat Detector"


def test_unlettered_circle_is_not_suggested_as_a_lettered_detector(client):
    # An architectural door tag is a circle with no letters; the verified
    # smoke detector is a circle with "S". Shape alone must not match them.
    def door_tag(b):
        b.add_arc((0, 0), 200, 0, 180)
        b.add_arc((0, 0), 200, 180, 360)

    d = upload(client, make_dxf(_TMP / "tag.dxf", [("M_Door Tag-1", door_tag, 0, 0, 0)]))
    g = d["groups"][0]
    assert g["status"] == "unknown"
    assert g["suggestion"] is None


def call_point(b):
    b.add_lwpolyline([(-150, -150), (150, -150), (150, 150), (-150, 150)], close=True)
    b.add_circle((0, 0), 60)


def test_wp_beside_a_symbol_is_read_and_goes_to_the_nearest_symbol(client):
    """As on FA-102: 'wp' sits 60-300 mm beside a 300 mm call point, outside
    its outline. It belongs to that call point, not to the one next to it,
    and a cable note ("TO FACP") or riser arrow ("F/B") is never letters."""
    doc = ezdxf.new("R2018")
    call_point(doc.blocks.new("MCP"))
    msp = doc.modelspace()
    msp.add_blockref("MCP", (0, 0), dxfattribs={"layer": "E-FIRE"})
    msp.add_text("wp", height=200).set_placement((230, -100))  # ~80 mm right of the outline
    msp.add_blockref("MCP", (1500, 0), dxfattribs={"layer": "E-FIRE"})  # indoor, no wp
    msp.add_text("TO FACP", height=100).set_placement((1500, 250))
    msp.add_text("F/B", height=100).set_placement((1700, -60))
    p = _TMP / "wp.dxf"
    doc.saveas(p)
    d = upload(client, p)
    labels = sorted((g["label"], g["count"]) for g in d["groups"])
    assert labels == [("", 1), ("WP", 1)]


def test_reprocess_moves_wp_call_points_to_the_weatherproof_type(client, monkeypatch):
    """A drawing read before 'wp' was understood keeps the user's decision
    and, now that the mark is read, lands on the weatherproof type."""
    from app.ifc.dxf import geometry as G

    doc = ezdxf.new("R2018")
    call_point(doc.blocks.new("MCP-OLD"))
    msp = doc.modelspace()
    for x in (0, 3000):
        msp.add_blockref("MCP-OLD", (x, 0), dxfattribs={"layer": "E-FIRE"})
        msp.add_text("wp", height=200).set_placement((x + 230, -100))
    p = _TMP / "reprocess.dxf"
    doc.saveas(p)

    # read it the old way: 'wp' not recognised as letters
    real = G.is_label_like
    monkeypatch.setattr(G, "is_label_like", lambda s: real(s) and s.strip().upper() != "WP")
    d = upload(client, p)
    g = d["groups"][0]
    assert g["label"] == ""
    client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [g["signature"]], "device_type_id": type_id(client, "MCP")})
    monkeypatch.setattr(G, "is_label_like", real)

    rep = client.post("/api/drawings/reprocess").json()
    assert any("MCP -> MCP-WP" in w for w in rep["weatherproof"])
    d = client.get(f"/api/drawings/{d['id']}").json()
    g = d["groups"][0]
    assert g["label"] == "WP" and g["status"] == "verified" and g["device_type"]["code"] == "MCP-WP"
    assert g["count"] == 2


def call_point_clipped(b):
    """FA-104's call point: the bottom edge of the box lost in export."""
    b.add_line((-150, 150), (150, 150))
    b.add_line((-150, 150), (-150, -100))
    b.add_line((150, 150), (150, -100))
    b.add_circle((0, 0), 60)


def test_clipped_symbol_of_a_verified_revit_family_is_counted(client):
    fam = "INTELLIGENT MANUAL CALL POINT - MANUAL CALL POINT-2016599-{}"
    d = upload(client, make_dxf(_TMP / "fam1.dxf", [(fam.format("FA-105-TYP FLOOR"), call_point, 0, 0, 0)]))
    client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [d["groups"][0]["signature"]], "device_type_id": type_id(client, "MCP")})

    d = upload(client, make_dxf(_TMP / "fam2.dxf", [
        (fam.format("FA-104-HEALTH COMPLEX"), call_point_clipped, 0, 0, 0),
        (fam.format("FA-104-HEALTH COMPLEX"), call_point_clipped, 2000, 0, 0),
    ]))
    g = d["groups"][0]
    assert g["status"] == "verified" and g["device_type"]["code"] == "MCP"
    assert g["match"]["kind"] == "family"


def test_clipped_symbol_without_the_family_name_is_not_counted(client):
    """Shape alone never counts a partial drawing: an unrelated block with
    the same clipped shape waits for a person."""
    d = upload(client, make_dxf(_TMP / "fam3.dxf", [("SOMETHING ELSE", call_point_clipped, 0, 0, 0)]))
    assert d["groups"][0]["status"] != "verified"


def test_device_layer_symbol_is_never_suggested_as_not_a_device(client):
    def panel(b):
        b.add_lwpolyline([(0, 0), (600, 0), (600, 900), (0, 900)], close=True)

    # a plain rectangle on the architecture layer, marked not a device
    doc = ezdxf.new("R2018")
    panel(doc.blocks.new("CUPBOARD"))
    doc.modelspace().add_blockref("CUPBOARD", (0, 0), dxfattribs={"layer": "0"})
    doc.saveas(_TMP / "cup.dxf")
    d = upload(client, _TMP / "cup.dxf")
    client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [d["groups"][0]["signature"]], "ignore": True})

    # the same rectangle drawn with lines, on an electrical layer: a panel
    doc = ezdxf.new("R2018")
    b = doc.blocks.new("VOICE EVACUATION CONTROL PANEL")
    for a, c in [((0, 0), (600, 0)), ((600, 0), (600, 900)), ((600, 900), (0, 900)), ((0, 900), (0, 0))]:
        b.add_line(a, c)
    doc.modelspace().add_blockref("VOICE EVACUATION CONTROL PANEL", (0, 0), dxfattribs={"layer": "E-ELEC-EQPM"})
    doc.saveas(_TMP / "vep.dxf")
    g = upload(client, _TMP / "vep.dxf")["groups"][0]
    assert g["status"] == "unknown", "never offered as not a device"


def building_dxf(filename: str = "building.dxf") -> Path:
    """A full-building set: one sheet per floor with a title block, a
    typical floor standing for three floors, and a riser diagram repeating
    the devices as symbols."""
    doc = ezdxf.new("R2018")
    smoke(doc.blocks.new("SD"))
    msp = doc.modelspace()
    areas = {"ground": (0, 0), "typical": (100_000, 0), "riser": (200_000, 0)}
    for x in range(4):                                     # 4 on the ground floor
        msp.add_blockref("SD", (areas["ground"][0] + x * 2000, 5000))
    for x in range(5):                                     # 5 on the typical floor
        msp.add_blockref("SD", (areas["typical"][0] + x * 2000, 5000))
    for x in range(20):                                    # riser symbols
        msp.add_blockref("SD", (areas["riser"][0] + x * 1000, 5000))
    msp.add_blockref("SD", (0, 900_000))                   # stray copy, on no sheet

    title_block = doc.blocks.new("TITLE")
    title_block.add_attdef("TITLE", (0, 0), dxfattribs={"height": 5})
    for name, area, title in (("FA-01", "ground", "GROUND FLOOR PLAN"),
                              ("FA-02", "typical", "TYP 2-4 FLOOR PLAN"),
                              ("FA-03", "riser", "FIRE ALARM SCHEMATIC RISER DIAGRAM")):
        lay = doc.layouts.new(name)
        cx, cy = areas[area][0] + 20_000, areas[area][1] + 5000
        lay.add_viewport(center=(200, 150), size=(400, 300), view_center_point=(cx, cy), view_height=30_000)
        lay.add_blockref("TITLE", (10, 10)).add_auto_attribs({"TITLE": title})
    p = _TMP / filename
    doc.saveas(p)
    return p


def test_building_quantities_follow_the_sheets_and_typical_floors(client):
    """Quantities = ground + 3 x typical; the riser and a stray copy
    outside every sheet are not counted."""
    d = upload(client, building_dxf())
    sheets = {s["name"]: s for s in d["sheets"]}
    assert sheets["FA-01"]["kind"] == "plan" and sheets["FA-01"]["multiplier"] == 1
    assert sheets["FA-02"]["kind"] == "plan" and sheets["FA-02"]["multiplier"] == 3
    assert sheets["FA-03"]["kind"] == "diagram"
    g = by_label(d, "S")[0]
    assert g["status"] == "verified"                       # SD is in the library by now
    assert g["count"] == 30
    assert g["by_sheet"] == {"FA-01": 4, "FA-02": 5, "FA-03": 20, "(outside sheets)": 1}
    assert g["boq_qty"] == 4 + 3 * 5
    assert d["totals"]["fire_alarm"] == 19
    assert d["totals"]["not_counted_instances"] == 21

    # the typical plan actually stands for 4 floors: override, then reset
    d = client.put(f"/api/drawings/{d['id']}/floors", json={"sheet": "FA-02", "multiplier": 4}).json()
    assert d["totals"]["fire_alarm"] == 4 + 4 * 5
    d = client.put(f"/api/drawings/{d['id']}/floors", json={"sheet": "FA-02", "multiplier": None}).json()
    assert d["totals"]["fire_alarm"] == 19


def test_quantities_are_listed_floor_by_floor(client):
    """Each floor plan with its devices, the floors it stands for and the
    product, in sheet order; the same list in the Excel export."""
    from io import BytesIO

    from openpyxl import load_workbook

    d = upload(client, building_dxf("building-floors.dxf"))
    fa = d["floor_boq"]["fire_alarm"]
    assert [(f["sheet"], f["floor_name"], f["multiplier"]) for f in fa["floors"]] == [("FA-01", "GROUND FLOOR", 1), ("FA-02", "TYP 2-4 FLOOR", 3)]
    assert [(r["device_type"]["code"], r["per_floor"], r["qty"]) for r in fa["floors"][0]["rows"]] == [("SD", 4, 4)]
    assert [(r["device_type"]["code"], r["per_floor"], r["qty"]) for r in fa["floors"][1]["rows"]] == [("SD", 5, 15)]
    assert [(b["device_type"]["code"], b["qty"]) for b in fa["building"]] == [("SD", 19)] and fa["qty"] == 19
    assert d["floor_boq"]["emergency_light"]["qty"] == 0

    wb = load_workbook(BytesIO(client.get(f"/api/drawings/{d['id']}/export").content))
    rows = list(wb["Fire Alarm"].iter_rows(values_only=True))
    assert rows[0][:8] == ("S.No", "Floor", "Sheet", "Device", "Code", "Qty per floor", "No. of floors", "Total qty")
    assert rows[1][:2] == (1, "GROUND FLOOR") and rows[1][4:8] == ("SD", 4, 1, 4)
    assert rows[2][3] == "Floor total" and rows[2][7] == 4
    assert rows[3][:2] == (2, "TYP 2-4 FLOOR") and rows[3][4:8] == ("SD", 5, 3, 15)
    assert rows[4][3] == "Floor total" and rows[4][7] == 15
    assert rows[6][1].startswith("BUILDING TOTAL") and rows[6][4] == "SD" and rows[6][7] == 19
    assert rows[-1][3] == "TOTAL" and rows[-1][7] == 19


def test_one_plan_for_several_floors_is_multiplied(client):
    d = upload(client, make_dxf(_TMP / "257-Sheet - FA-105 - TYP(1ST TO 14TH) FLOORS FIRE ALARM LAYOUT.dxf", [("SD", smoke, 0, 0, 0), ("SD", smoke, 900, 0, 0)]))
    assert d["floor_info"]["mode"] == "multiple"
    assert d["floor_info"]["floors"] == 14 and d["floor_info"]["floor_name"] == "TYP (1ST TO 14TH) FLOORS"
    assert d["totals"]["fire_alarm"] == 2 * 14


def architecture_block(doc, name="ARCH-RVT-1-FA-104"):
    """An architectural plan exported as one block: walls, and 12 doors."""
    door = doc.blocks.new(f"{name}-DOOR")
    door.add_arc((0, 0), 900, 0, 90)
    door.add_line((0, 0), (900, 0))
    arch = doc.blocks.new(name)
    arch.add_lwpolyline([(0, 0), (40_000, 0), (40_000, 30_000), (0, 30_000)], close=True)
    for i in range(12):
        arch.add_blockref(f"{name}-DOOR", (2000 + i * 3000, 0))
    return name


def test_single_floor_counts_only_devices_on_the_architecture(client):
    """One floor: named from its title block, and the legend symbols and a
    stray copy beside the building are not counted."""
    doc = ezdxf.new("R2018")
    smoke(doc.blocks.new("SD"))
    msp = doc.modelspace()
    msp.add_blockref(architecture_block(doc), (0, 0), dxfattribs={"layer": "0"})
    for i in range(6):                                  # on the plan
        msp.add_blockref("SD", (5000 + i * 5000, 15_000), dxfattribs={"layer": "E-FIRE"})
    msp.add_blockref("SD", (60_000, 10_000), dxfattribs={"layer": "E-FIRE"})   # legend, beside the plan
    msp.add_text("LEGEND", height=500).set_placement((58_000, 12_000))
    msp.add_blockref("SD", (0, 500_000), dxfattribs={"layer": "E-FIRE"})       # stray copy
    lay = doc.layouts.get("Layout1")
    lay.add_viewport(center=(420, 297), size=(800, 560), view_center_point=(35_000, 250_000), view_height=600_000)
    lay.add_text("HC FLOOR FIRE ALARM LAYOUT", height=4.5).set_placement((20, 20))
    p = _TMP / "25C43C~1.dxf"
    doc.saveas(p)

    d = upload(client, p)
    fi = d["floor_info"]
    assert fi["mode"] == "single" and fi["floor_name"] == "HC FLOOR" and fi["architecture_found"]
    g = by_label(d, "S")[0]
    assert g["count"] == 8 and g["boq_qty"] == 6
    assert g["by_sheet"]["(not on the architecture)"] == 2
    assert d["totals"]["fire_alarm"] == 6

    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(client.get(f"/api/drawings/{d['id']}/export").content))
    rows = list(wb["Fire Alarm"].iter_rows(values_only=True))
    assert rows[0][1] == "Floor" and rows[1][1] == "HC FLOOR" and rows[1][4] == 6


def test_reprocess_changes_nothing_when_nothing_changed(client):
    """Re-reading with unchanged rules must not decide anything: two
    unrelated items at the same point inside an underlay (no handles) once
    let one inherit the other's 'not a device'."""
    doc = ezdxf.new("R2018")
    a = doc.blocks.new("TABLE")
    a.add_lwpolyline([(0, 0), (900, 0), (900, 600), (0, 600)], close=True)
    b = doc.blocks.new("CHAIR")
    b.add_circle((0, 0), 250)
    under = doc.blocks.new("UNDERLAY")
    for i in range(6):  # a container: >= 10 inserts
        under.add_blockref("TABLE", (i * 3000, 0))
        under.add_blockref("CHAIR", (i * 3000, 0))  # same insertion point as the table
    doc.modelspace().add_blockref("UNDERLAY", (0, 0), dxfattribs={"layer": "0"})
    p = _TMP / "underlay.dxf"
    doc.saveas(p)
    d = upload(client, p)
    table = next(g for g in d["groups"] if "TABLE" in g["block_names"])
    chair = next(g for g in d["groups"] if "CHAIR" in g["block_names"])
    client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [table["signature"]], "ignore": True})
    before = client.get(f"/api/drawings/{d['id']}").json()
    status_before = next(g for g in before["groups"] if g["signature"] == chair["signature"])["status"]

    client.post("/api/drawings/reprocess")
    after = client.get(f"/api/drawings/{d['id']}").json()
    assert next(g for g in after["groups"] if g["signature"] == chair["signature"])["status"] == status_before != "ignored"


def test_letters_laid_over_the_block_count_as_its_label(client):
    doc = ezdxf.new("R2018")
    b = doc.blocks.new("CIRCLE ONLY")
    b.add_circle((0, 0), 200)
    msp = doc.modelspace()
    msp.add_blockref("CIRCLE ONLY", (0, 0))
    msp.add_text("MS", height=150).set_placement((-80, -60))
    msp.add_blockref("CIRCLE ONLY", (5000, 0))
    msp.add_text("H", height=150).set_placement((4950, -60))
    p = _TMP / "d.dxf"
    doc.saveas(p)
    d = upload(client, p)
    assert {g["label"] for g in d["groups"]} == {"MS", "H"}


def test_nested_container_is_opened(client):
    doc = ezdxf.new("R2018")
    smoke(doc.blocks.new("DET"))
    plan = doc.blocks.new("FLOOR PLAN")
    for i in range(12):
        plan.add_blockref("DET", (i * 1000, 0))
    doc.modelspace().add_blockref("FLOOR PLAN", (0, 0))
    p = _TMP / "e.dxf"
    doc.saveas(p)
    d = upload(client, p)
    s = by_label(d, "S")[0]
    assert s["count"] == 12
    assert s["status"] == "verified"
    assert "FLOOR PLAN" in d["containers"]


def test_ignore_and_unverify(client):
    d = upload(client, _TMP / "a.dxf")
    ex = next(g for g in d["groups"] if g["label"] == "")["signature"]
    d = client.post("/api/symbols/verify", json={"drawing_id": d["id"], "signatures": [ex], "ignore": True}).json()
    assert d["totals"]["emergency_light"] == 0
    assert d["totals"]["ignored_symbols"] == 1
    client.post("/api/symbols/unverify", json={"signatures": [ex]})
    d = client.get(f"/api/drawings/{d['id']}").json()
    assert next(g for g in d["groups"] if g["signature"] == ex)["status"] in ("unknown", "suggested")


def test_export_has_both_tabs(client):
    from io import BytesIO

    from openpyxl import load_workbook

    d = upload(client, _TMP / "a.dxf")
    # the exit sign was sent back for verification by an earlier test: no quantities until it is answered
    ask = [g for g in d["groups"] if g["review"] in ("answer", "suggested", "architecture", "confirm")]
    assert [g["label"] for g in ask] == [""] and not d["review"]["ready"]
    assert client.get(f"/api/drawings/{d['id']}/export").status_code == 409
    d = client.post("/api/symbols/review", json={"drawing_id": d["id"], "answers": [
        {"signature": ask[0]["signature"], "device_type_id": type_id(client, "EXIT")}]}).json()
    assert d["review"]["ready"]
    r = client.get(f"/api/drawings/{d['id']}/export")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert {"Info", "Fire Alarm", "Emergency Lighting", "Pending Verification", "Occurrences"} <= set(wb.sheetnames)
    rows = list(wb["Fire Alarm"].iter_rows(values_only=True))
    assert rows[0][1] == "Floor"  # a single-floor drawing names its floor
    assert rows[1][3] == "SD" and rows[1][4] == 4


def test_quantities_wait_for_the_symbols_to_be_answered(client):
    """No quantities until every symbol on the plan that the library does not
    know is answered: a new device (no guess) and furniture from the
    architect's xref (offered as not a device). A door without letters is
    not asked. One save puts the answers in the library and gives the
    quantities."""
    import json

    from app.ifc.library_file import library_path

    doc = ezdxf.new("R2018")
    smoke(doc.blocks.new("SD"))
    q = doc.blocks.new("NEW DEVICE")
    q.add_circle((0, 0), 200)
    q.add_text("Q", height=150).set_placement((-50, -75))
    bed = doc.blocks.new("XR-ARCH$0$BED")
    bed.add_lwpolyline([(0, 0), (2000, 0), (2000, 1500), (0, 1500)], close=True)
    bed.add_text("BED", height=200).set_placement((600, 600))
    door = doc.blocks.new("DOOR-90")
    door.add_arc((0, 0), 900, 0, 90)
    door.add_line((0, 0), (900, 0))
    msp = doc.modelspace()
    for i in range(4):
        msp.add_blockref("SD", (i * 3000, 0), dxfattribs={"layer": "E-FIRE"})
    for i in range(3):
        msp.add_blockref("NEW DEVICE", (i * 3000, 5000), dxfattribs={"layer": "E-FIRE"})
    for i in range(2):
        msp.add_blockref("XR-ARCH$0$BED", (i * 5000, 10_000), dxfattribs={"layer": "XR-ARCH$0$A-FURN"})
    msp.add_blockref("DOOR-90", (0, 15_000), dxfattribs={"layer": "A-DOOR"})
    p = _TMP / "review.dxf"
    doc.saveas(p)

    d = upload(client, p)
    new, furniture = by_label(d, "Q")[0], by_label(d, "BED")[0]
    door_group = next(g for g in d["groups"] if "DOOR-90" in g["block_names"])
    assert (new["review"], furniture["review"], door_group["review"], by_label(d, "S")[0]["review"]) == ("answer", "architecture", "optional", None)
    assert d["review"]["required"] == 2 and not d["review"]["ready"]
    assert client.get(f"/api/drawings/{d['id']}/export").status_code == 409
    assert next(x for x in client.get("/api/drawings").json() if x["id"] == d["id"])["review_required"] == 2

    # one answer per symbol only
    both = {"signature": new["signature"], "device_type_id": type_id(client, "HD"), "ignore": True}
    assert client.post("/api/symbols/review", json={"drawing_id": d["id"], "answers": [both]}).status_code == 422

    d = client.post("/api/symbols/review", json={"drawing_id": d["id"], "answers": [
        {"signature": new["signature"], "device_type_id": type_id(client, "HD")},
        {"signature": furniture["signature"], "ignore": True},
    ]}).json()
    assert d["review"]["ready"] and d["review"]["required"] == 0
    assert d["totals"]["fire_alarm"] == 4 + 3
    assert client.get(f"/api/drawings/{d['id']}/export").status_code == 200
    saved = {x["signature"]: x for x in json.loads(library_path().read_text(encoding="utf-8"))["symbols"]}
    assert saved[new["signature"]]["device_code"] == "HD" and saved[furniture["signature"]]["is_ignored"]


def test_only_guesses_the_letters_confirm_are_filled_in():
    """Furniture from a Revit export is architecture by its family name; a
    device named like one (a door holder) never is. A guess from the shape
    alone is left for the user; one the letters agree with is filled in."""
    from app.ifc.resolve import architecture_like, review_kind

    sofa = {"layers": {"0": 5}, "block_names": {"Sofa - Double - Sofa-10297482-FA-105-_1ST TO14TH_TYP  FLOOR": 5}}
    holder = {"layers": {"0": 2}, "block_names": {"Door Holder - Magnetic-2044123-FA-105-_1ST TO14TH_TYP  FLOOR": 2}}
    lobby = {"layers": {"E-FIRE": 2}, "block_names": {"LIFT LOBBY": 2}}
    xref = {"layers": {"1014-FLR-1$0$Flat number": 3}, "block_names": {"*U3283": 3}}
    bound = {"layers": {"0": 9}, "block_names": {"1014-FLR-1 Landscape$0$lovesacplan": 9}}
    xref_device = {"layers": {"0": 3}, "block_names": {"FA-DEVICES$0$SMOKE DETECTOR": 3}}
    assert architecture_like(sofa) and architecture_like(xref) and architecture_like(bound)
    assert not architecture_like(holder) and not architecture_like(lobby) and not architecture_like(xref_device)

    guess = {"device_type": {"code": "SPKW"}, "is_ignored": False}
    base = {"signature": "x", "status": "suggested", "on_plans": 4, "hint": 3, "suggestion": guess, "layers": {"E-FIRE": 4}, "block_names": {"WALL SPEAKER": 4}}
    assert review_kind({**base, "label": ""}, set()) == "answer"
    assert review_kind({**base, "label": "E"}, set()) == "suggested"


def test_an_unlettered_look_alike_is_suggested_not_counted(client):
    """A box without letters that looks like a verified exit sign (a
    wardrobe, a treadmill) is not counted by its shape alone: it is offered
    as a guess for the user to answer."""
    def wardrobe(block):
        block.add_lwpolyline([(-300, -120), (300, -120), (300, 120), (-300, 120)], close=True)
        block.add_line((-300, 0), (300, 0))
        block.add_line((0, -120), (0, -60))

    d = upload(client, make_dxf(_TMP / "wardrobe.dxf", [("Wardrobe 600 mm2 - 1800x2200-14005969-FA-105-TYP", wardrobe, 0, 0, 0)], layer="0"))
    g = d["groups"][0]
    assert g["status"] == "suggested" and g["suggestion"]["device_type"]["code"] == "EXIT"
    assert d["totals"]["emergency_light"] == 0
    assert g["review"] in ("answer", "architecture")


def test_indoor_is_not_a_door():
    from app.ifc.resolve import _ARCH_NAME

    assert _ARCH_NAME.search("Aluminum Glass Door1 - Double Glass Door 1350x2000 mm")
    assert not _ARCH_NAME.search("INDOOR SPEAKER WITH FLASHER") and not _ARCH_NAME.search("OUTDOOR SOUNDER")
    for name in ("Rectangular Mullion - 50 x 150mm", "M_Concrete-Rectangular-Column - FC1", "Exercise_Bike_4766", "tire_family_11232", "Laundry_Chute_Intake_3944"):
        assert _ARCH_NAME.search(name), name


def test_a_skipped_symbol_is_left_out_without_teaching_the_library(client):
    import json

    from app.ifc.library_file import library_path

    def zed(block):
        block.add_circle((0, 0), 200)
        block.add_text("Z", height=150).set_placement((-50, -75))

    d = upload(client, make_dxf(_TMP / "skip.dxf", [("ZED", zed, 0, 0, 0), ("ZED", zed, 3000, 0, 0)]))
    z = by_label(d, "Z")[0]
    assert z["review"] == "answer" and not d["review"]["ready"]

    d = client.post("/api/symbols/review", json={"drawing_id": d["id"], "answers": [{"signature": z["signature"], "skip": True}]}).json()
    assert d["review"]["ready"] and d["review"]["skipped"] == 1 and by_label(d, "Z")[0]["review"] == "skipped"
    assert d["totals"]["fire_alarm"] == 0
    assert z["signature"] not in {x["signature"] for x in json.loads(library_path().read_text(encoding="utf-8"))["symbols"]}

    # taking the skip back asks again
    d = client.post("/api/symbols/review", json={"drawing_id": d["id"], "answers": [{"signature": z["signature"]}]}).json()
    assert not d["review"]["ready"] and by_label(d, "Z")[0]["review"] == "answer"


def test_device_type_in_use_cannot_be_deleted(client):
    assert client.delete(f"/api/device-types/{type_id(client, 'SD')}").status_code == 409
