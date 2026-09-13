r"""Reproduce review findings using synthetic data, without touching the app database.

Run from the workspace root:
    backend\venv\Scripts\python.exe docs\review_checks.py

These checks describe current defects; they are not acceptance tests that should
be kept passing after the defects are fixed. No archive or document files are read.
"""

import io
import os
import sys
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["DATASHEET_LIBRARIES"] = "{}"
os.environ["SECRET_KEY"] = "synthetic-review-key-unused-for-authentication"

from app.core.config import Settings
from app.schemas_project import ProjectBoqItemIn, ProjectCreate
from app.services.battery_calculation import MECHANICAL_RE, _quantity
from app.services.boq_export import parse_quantity
from app.routers.compliance import _specs
import openpyxl
import pymupdf


project = ProjectCreate(
    ep_number="x/../../escaped",
    source_folder_path="C:/review-only",
    drf_document_path="C:/review-only/example.pdf",
)
uploads = ROOT / "synthetic-review-uploads"
target = (uploads / f"EP-{project.ep_number}").resolve()
print("EP traversal accepted; computed folder escapes uploads:", not target.is_relative_to(uploads))
print("Caller-supplied source folder accepted:", project.source_folder_path is not None)
print("Blank JWT secret accepted:", Settings(_env_file=None, secret_key="").secret_key == "")

line = ProjectBoqItemIn(description="", quantity="-5", unit_price=-1, total_price=99)
print("Empty description and negative quantity/price accepted:", line.quantity == "-5")
print("Quantity '1,000': export =", parse_quantity("1,000"), "; battery =", _quantity("1,000"))
print("Powered cabinet matches mechanical rule:", bool(MECHANICAL_RE.search("Powered cabinet with monitoring electronics")))

matches, warnings = _specs(SimpleNamespace(
    id=-999, source_folder_path=None, design_sheets=[], boq_items=[], systems=[]
))
print("No-archive project returns before scanning uploaded specifications:", not matches, warnings)

stream = io.BytesIO()
workbook = openpyxl.Workbook()
workbook.active.append(["Catalog", "Description", "Qty"])
workbook.active.append(["TEST-1", "Synthetic review item", 3])
workbook.save(stream)
workbook.close()
try:
    with pymupdf.open(stream=stream.getvalue()) as document:
        print("XLSX opened by PyMuPDF; page count:", document.page_count)
        print("XLSX rendered first page:", bool(document[0].get_pixmap().samples))
except Exception as exc:
    print("XLSX open/render exception:", type(exc).__name__)

cell = openpyxl.Workbook().active.cell(1, 1, "=1+1")
print("Unescaped exported description becomes Excel formula:", cell.data_type == "f")
